from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from pptx import Presentation

from coursepilot.domain.ppt import (
    PPTArtifact,
    SlideArchitecture,
    SlideAssetPlaceholder,
    SlideContent,
    SlidePlan,
)
from coursepilot.exporters.pptx import PPTXVersionedExporter
from evaluation.p16_critical_repair import _recover_structured_response, preflight
from evaluation.p16_ppt_eval import _BudgetCheckpoint


def test_p16_critical_repair_scope_is_exact_and_network_free() -> None:
    result = preflight()
    assert result["external_calls_made"] == 0
    assert result["critical_page_count"] == 17
    assert result["scope_exact"] is True
    assert result["hard_limits"] == {
        "max_cost_cny": "0.45",
        "max_input_tokens": 130_000,
        "max_output_tokens": 60_000,
        "max_requests": 20,
        "max_repairs_per_slide": 1,
    }


def test_p16_targeted_export_hint_materializes_editable_table(tmp_path: Path) -> None:
    plan = SlidePlan(
        slide_id="slide-1",
        slide_index=1,
        slide_type="title",
        layout_role="title",
        title_intent="职业风险表",
        asset_kind="editable_table",
    )
    references = SlidePlan(
        slide_id="slide-2",
        slide_index=2,
        slide_type="concept",
        layout_role="concept",
        title_intent="说明",
    )
    last = SlidePlan(
        slide_id="slide-3",
        slide_index=3,
        slide_type="references",
        layout_role="references",
        title_intent="来源",
    )
    architecture = SlideArchitecture(
        course_id="course",
        lesson_artifact_id="lesson",
        template_id="template",
        template_snapshot_id="snapshot",
        slide_count=3,
        plans=[plan, references, last],
    )
    table_asset = SlideAssetPlaceholder(
        asset_id="table",
        kind="editable_table",
        description="职业概率表",
        alt_text="职业概率表",
    )
    artifact = PPTArtifact(
        course_id="course",
        lesson_artifact_id="lesson",
        template_id="template",
        template_snapshot_id="snapshot",
        architecture=architecture,
        slides=[
            SlideContent(slide_id="slide-1", title="职业风险表", assets=[table_asset]),
            SlideContent(slide_id="slide-2", title="说明"),
            SlideContent(slide_id="slide-3", title="来源"),
        ],
    )
    output = PPTXVersionedExporter().export(
        artifact,
        tmp_path / "table.pptx",
        slide_render_hints={
            "slide-1": {
                "title_top": True,
                "editable_table_tsv": "职业\t概率\n电话推销员\t99%",
            }
        },
    )
    presentation = Presentation(output)
    assert any(shape.has_table for shape in presentation.slides[0].shapes)
    assert "Asset placeholder" not in "\n".join(
        getattr(shape, "text", "") for shape in presentation.slides[0].shapes
    )


def test_p16_recovers_billed_response_with_only_invalid_model_hash(tmp_path: Path) -> None:
    checkpoint = _BudgetCheckpoint(
        additional_max_input=130_000,
        additional_max_output=60_000,
        additional_max_requests=20,
        additional_max_cost_cny=Decimal("0.45"),
        console_baseline_cny=Decimal("0"),
        console_total_cap_cny=Decimal("0.45"),
        path=tmp_path / "checkpoint.json",
        enforce_off_peak=False,
    )
    raw = {
        "slide_id": "slide-1",
        "title": "修复后",
        "bullets": ["内容"],
        "body_text": "",
        "speaker_notes": "说明",
        "citations": [],
        "assets": [],
        "content_sha256": "0" * 64,
    }
    checkpoint.failed_invocations = [
        {
            "request_sha256": "a" * 64,
            "error_category": "structured_parse_error",
            "attempts": [
                {
                    "error_message": "Failed to parse SlideContent from completion "
                    + __import__("json").dumps(raw)
                    + ". Got: mismatch",
                    "usage": {"input_tokens": 10, "output_tokens": 20},
                }
            ],
        }
    ]
    recovered = _recover_structured_response(checkpoint)
    assert recovered.slide_id == "slide-1"
    assert recovered.content_sha256 != "0" * 64
    assert checkpoint.cache["a" * 64]["recovered_without_provider_retry"] is True
