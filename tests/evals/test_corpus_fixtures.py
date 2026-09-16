import hashlib
from pathlib import Path

import fitz
from docx import Document
from docx.oxml.ns import qn

from courserag.evals.schemas import CorpusFixtureManifest, DS0CorpusDataset, DS1SamplingPlan
from evaluation.corpus_fixtures import (
    CLEAN_SOURCE_PAGES,
    DOCX_ANCHORS,
    DOCX_PRIMARY,
    PDF_PRIMARY,
    PDF_SCAN_CLEAN,
    TWO_COLUMN_WIDTH_DXA,
    build_docx_structure_stress,
    build_raster_subset_pdf,
    sha256_file,
)


def test_raster_subset_is_deterministic_and_has_no_text_layer(tmp_path):
    source_path = tmp_path / "source.pdf"
    with fitz.open() as source:
        for number in (1, 2):
            page = source.new_page()
            page.insert_text((72, 72), f"source page {number}")
        source.save(source_path, no_new_id=True)

    first = tmp_path / "first.pdf"
    second = tmp_path / "second.pdf"
    first_map = build_raster_subset_pdf(
        source_path,
        first,
        source_pages=[1, 2],
        dpi=72,
        grayscale=False,
        jpeg_quality=None,
    )
    second_map = build_raster_subset_pdf(
        source_path,
        second,
        source_pages=[1, 2],
        dpi=72,
        grayscale=False,
        jpeg_quality=None,
    )

    assert first_map == second_map
    assert sha256_file(first) == sha256_file(second)
    with fitz.open(first) as document:
        assert document.page_count == 2
        assert all(not page.get_text("text").strip() for page in document)


def test_docx_stress_copy_is_deterministic_and_source_hashed(tmp_path):
    source_path = tmp_path / "source.docx"
    source = Document()
    source.add_paragraph("Source title")
    source.add_paragraph("Source paragraph")
    for number in range(6):
        source.add_paragraph(f"Filler paragraph {number}")
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "left"
    table.cell(0, 1).text = "right"
    source.save(str(source_path))

    first = tmp_path / "first.docx"
    second = tmp_path / "second.docx"
    first_map = build_docx_structure_stress(source_path, first)
    second_map = build_docx_structure_stress(source_path, second)

    assert first_map == second_map
    assert sha256_file(first) == sha256_file(second)
    assert {item.text_sha256 for item in first_map} >= {
        hashlib.sha256(text.encode("utf-8")).hexdigest()
        for text in ("Source title", "Source paragraph", "left", "right")
    }
    rendered = Document(str(first))
    table_width = rendered.tables[0]._tbl.tblPr.first_child_found_in("w:tblW")
    assert table_width is not None
    assert table_width.get(qn("w:type")) == "dxa"
    assert int(table_width.get(qn("w:w"))) == TWO_COLUMN_WIDTH_DXA


def test_tracked_pre_p03_candidates_and_sampling_plan_are_unapproved_and_exact():
    assert CLEAN_SOURCE_PAGES == [21, 23, 24, 26, 30]
    root = Path("datasets/courserag_eval/v1")
    manifest_path = root / "provenance" / "corpus_fixture_manifest.json"
    manifest = CorpusFixtureManifest.model_validate_json(manifest_path.read_text(encoding="utf-8"))
    candidates = DS0CorpusDataset.model_validate_json(
        (root / "candidates" / "ds0" / "pilot.json").read_text(encoding="utf-8")
    )
    sampling_plan = DS1SamplingPlan.model_validate_json(
        (root / "plans" / "ds1_sampling_plan.json").read_text(encoding="utf-8")
    )

    assert manifest.semantic_source_count == 2
    assert manifest.evaluation_document_count == 6
    assert len(candidates.documents) == 6
    assert all(document.review_status == "candidate" for document in candidates.documents)
    assert all(document.approval is None for document in candidates.documents)
    assert all(
        document.document_version == f"eval-v1-{document.sha256[:16]}"
        for document in candidates.documents
    )
    manifest_sha256 = sha256_file(manifest_path)
    assert all(
        document.derivation_manifest_sha256 == manifest_sha256
        for document in candidates.documents
        if document.document_role == "derived_fixture"
    )
    assert {
        artifact.document_id for artifact in manifest.artifacts if artifact.retrieval_eligible
    } == {DOCX_PRIMARY, PDF_PRIMARY}

    selections = {selection.selection_id: selection for selection in sampling_plan.selections}
    assert selections["ds1-pilot-native-pdf-pages"].source_page_numbers == [5, 10, 12, 20, 32]
    assert selections["ds1-pilot-docx-structure-anchors"].section_anchors == DOCX_ANCHORS
    assert selections["ds1-pilot-ocr-clean-pages"].document_id == PDF_SCAN_CLEAN
    assert selections["ds1-pilot-ocr-clean-pages"].source_page_numbers == CLEAN_SOURCE_PAGES
