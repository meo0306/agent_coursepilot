from __future__ import annotations

import io
import zipfile
from base64 import b64decode

import fitz
import pytest
from docx import Document

from courserag.domain.document import sha256_bytes
from courserag.parsers.base import ParserLimits, UnsafeDocumentError, validate_docx
from courserag.parsers.docx import StructuredDOCXParser
from courserag.parsers.pdf import StructuredPDFParser


def _pdf_bytes(*, include_text: bool = True) -> bytes:
    document = fitz.open()
    page = document.new_page(width=400, height=600)
    if include_text:
        page.insert_text((40, 60), "Structured PDF paragraph with enough native text for parsing.")
        page.insert_text((40, 100), "Second line keeps span and bounding-box information.")
    result = document.tobytes()
    document.close()
    return result


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Course heading", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("Bold run").bold = True
    paragraph.add_run(" and normal run")
    paragraph.add_run().add_picture(
        io.BytesIO(
            b64decode(
                "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
            )
        )
    )
    document.add_paragraph("First item", style="List Number")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "C"
    table.cell(1, 1).text = "D"
    document.add_page_break()
    document.add_paragraph("After page break")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def test_pdf_parser_keeps_rawdict_structure_and_routes_blank_page_to_ocr() -> None:
    native = _pdf_bytes()
    parsed = StructuredPDFParser().parse(
        native,
        document_id="pdf-doc",
        document_version_id="pdf-v1",
        document_sha256=sha256_bytes(native),
    )
    page = parsed.document.pages[0]
    assert page.parse_decision is not None
    assert page.parse_decision.mode == "native"
    assert page.blocks[0].bbox is not None
    assert page.blocks[0].lines[0].spans[0].source_span.bboxes

    blank = _pdf_bytes(include_text=False)
    blank_result = StructuredPDFParser().parse(
        blank,
        document_id="blank-doc",
        document_version_id="blank-v1",
        document_sha256=sha256_bytes(blank),
    )
    assert blank_result.document.pages[0].parse_decision is not None
    assert blank_result.document.pages[0].parse_decision.mode == "ocr"
    assert {warning.code for warning in blank_result.document.warnings} == {"OCR_REQUIRED"}


def test_docx_parser_keeps_order_headings_runs_numbering_table_and_breaks() -> None:
    content = _docx_bytes()
    parsed = StructuredDOCXParser().parse(
        content,
        document_id="docx-doc",
        document_version_id="docx-v1",
        document_sha256=sha256_bytes(content),
    )
    page = parsed.document.pages[0]
    assert page.source_mode == "logical_docx"
    assert page.physical_page_index is None
    assert page.blocks[0].block_type == "heading"
    assert page.blocks[1].lines[0].spans[0].bold is True
    assert any(block.block_type == "list_item" for block in page.blocks)
    assert any("page" in block.style.get("breaks", []) for block in page.blocks)
    assert len(parsed.document.tables) == 1
    assert [cell.text for cell in parsed.document.tables[0].cells] == ["A", "B", "C", "D"]
    assert len(parsed.document.images) == 1
    assert next(iter(parsed.binary_assets)).startswith("images/")


def test_pdf_table_detector_preserves_row_column_cells() -> None:
    document = fitz.open()
    page = document.new_page(width=300, height=300)
    for x in (40, 140, 240):
        page.draw_line((x, 40), (x, 140), color=(0, 0, 0))
    for y in (40, 90, 140):
        page.draw_line((40, y), (240, y), color=(0, 0, 0))
    for x, y, value in ((60, 70, "A"), (160, 70, "B"), (60, 120, "C"), (160, 120, "D")):
        page.insert_text((x, y), value)
    content = document.tobytes()
    document.close()
    parsed = StructuredPDFParser().parse(
        content,
        document_id="table-doc",
        document_version_id="table-v1",
        document_sha256=sha256_bytes(content),
    )
    assert len(parsed.document.tables) == 1
    table = parsed.document.tables[0]
    assert (table.row_count, table.column_count) == (2, 2)
    assert [cell.text for cell in table.cells] == ["A", "B", "C", "D"]


def test_docx_validator_rejects_path_escape_and_compression_bomb() -> None:
    escaped = io.BytesIO()
    with zipfile.ZipFile(escaped, "w") as archive:
        archive.writestr("word/document.xml", "<document/>")
        archive.writestr("../escape", "bad")
    with pytest.raises(UnsafeDocumentError, match="escapes"):
        validate_docx(escaped.getvalue(), ParserLimits())

    compressed = io.BytesIO()
    with zipfile.ZipFile(compressed, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * 10_000)
    with pytest.raises(UnsafeDocumentError, match="compression-ratio"):
        validate_docx(compressed.getvalue(), ParserLimits(max_docx_compression_ratio=2.0))
