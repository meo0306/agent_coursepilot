import hashlib
import json
import shutil
import zipfile
from pathlib import Path

import pytest
from docx import Document

from coursepilot.templates import TemplateRegistry

ROOT = Path(__file__).resolve().parents[3]
REGISTRY = ROOT / "resources/templates/registry_v1.yaml"


def test_registry_contains_nine_templates_and_five_resource_roles() -> None:
    registry = TemplateRegistry(REGISTRY)

    assert len(registry.template_ids) == 9
    assert len(registry.resource_roles) == 5
    assert {registry.get(item).artifact_type for item in registry.template_ids} == {
        "lesson",
        "exam",
        "ppt",
    }


def test_snapshot_is_stable_and_pins_resource_hashes() -> None:
    registry = TemplateRegistry(REGISTRY)
    kwargs = {
        "prompt_hashes": {"planner": "a" * 64, "generator": "b" * 64},
        "model_profile_sha256": "c" * 64,
        "schema_versions": {"input": "v1", "output": "v1"},
    }

    first = registry.snapshot("lesson_standard_university_v1", **kwargs)
    second = registry.snapshot("lesson_standard_university_v1", **kwargs)

    assert first.snapshot_sha256 == second.snapshot_sha256
    assert (
        first.resource_hashes["lesson_default_docx"]
        == hashlib.sha256(registry.resource_path("lesson_default_docx").read_bytes()).hexdigest()
    )


def test_registry_rejects_definition_hash_drift(tmp_path: Path) -> None:
    candidate_root = tmp_path / "templates"
    shutil.copytree(REGISTRY.parent, candidate_root)
    candidate = candidate_root / "registry_v1.yaml"
    payload = json.loads(candidate.read_text(encoding="utf-8"))
    payload["definitions"][0]["sha256"] = "0" * 64
    candidate.write_text(json.dumps(payload), encoding="utf-8")

    with pytest.raises(ValueError, match="hash mismatch"):
        TemplateRegistry(candidate)


def test_docx_resources_are_openable_and_editable() -> None:
    registry = TemplateRegistry(REGISTRY)
    for role in ("lesson_default_docx", "exam_default_docx"):
        path = registry.resource_path(role)
        document = Document(path)
        assert document.paragraphs
        assert any("{{" in paragraph.text for paragraph in document.paragraphs)
        with zipfile.ZipFile(path) as archive:
            assert "word/styles.xml" in archive.namelist()


def test_pptx_resources_preserve_editable_master_layout_hierarchy() -> None:
    registry = TemplateRegistry(REGISTRY)
    for role in (
        "ppt_standard_lecture_pptx",
        "ppt_concept_explanation_pptx",
        "ppt_case_seminar_pptx",
    ):
        path = registry.resource_path(role)
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            assert "ppt/presentation.xml" in names
            assert any(name.startswith("ppt/slideMasters/slideMaster") for name in names)
            assert any(name.startswith("ppt/slideLayouts/slideLayout") for name in names)
            assert any(name.startswith("ppt/slides/slide") for name in names)
