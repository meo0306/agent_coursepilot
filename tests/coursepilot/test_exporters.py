from docx import Document
from pptx import Presentation

from coursepilot.exporters import ExamDocxExporter, LessonDocxExporter, PPTXExporter
from coursepilot.schemas.exam_schema import ExamBlueprintContent, QuestionGroupPlan
from coursepilot.schemas.lesson_schema import (
    LessonDesignContent,
    LessonSession,
    Reference,
    SessionPlan,
    TeachingProcessItem,
    TimeAllocation,
)
from coursepilot.schemas.ppt_schema import SlideItem, SlideOutlineContent
from coursepilot.schemas.question_schema import QuestionItem


def test_lesson_docx_exporter_writes_file(tmp_path):
    design = LessonDesignContent(
        course_name="AI",
        chapter="Search",
        total_sessions=1,
        session_duration=45,
        knowledge_points=["state space"],
        session_plan=[
            SessionPlan(
                session_index=1,
                session_title="Search",
                duration=45,
                knowledge_points=["state space"],
                teaching_focus="state space",
                time_allocation=[
                    TimeAllocation(activity="Intro", minutes=5),
                    TimeAllocation(activity="Lecture", minutes=30),
                    TimeAllocation(activity="Practice", minutes=5),
                    TimeAllocation(activity="Summary", minutes=5),
                ],
            )
        ],
        sessions=[
            LessonSession(
                session_index=1,
                session_title="Search",
                teaching_objectives=["Explain state space"],
                key_points=["state space"],
                teaching_process=[
                    TeachingProcessItem(stage="Intro", minutes=5, content="case")
                ],
                references=[Reference(chunk_id="chunk-1")],
            )
        ],
    )

    output = LessonDocxExporter().export(design, tmp_path / "lesson.docx")

    assert output.exists()
    doc = Document(output)
    assert "AI - Search" in doc.paragraphs[0].text


def test_exam_docx_exporter_writes_four_files_without_student_answers(tmp_path):
    blueprint = ExamBlueprintContent(
        course_name="AI",
        chapter_range="Search",
        generation_type="exam",
        total_score=2,
        question_groups=[
            QuestionGroupPlan(
                question_type="single_choice",
                count=1,
                score_each=2,
                total_score=2,
                knowledge_points=["state space"],
                difficulty="medium",
            )
        ],
        knowledge_points=["state space"],
    )
    question = QuestionItem(
        question_type="single_choice",
        knowledge_point="state space",
        difficulty="medium",
        score=2,
        question_text="Which option describes state space?",
        options={"A": "State representation", "B": "Noise"},
        correct_answer="A",
        explanation="State space represents search states.",
        references=[Reference(chunk_id="chunk-1")],
    )
    exporter = ExamDocxExporter()

    student = exporter.export_student_exam(blueprint, [question], tmp_path / "student.docx")
    answer = exporter.export_teacher_answer(blueprint, [question], tmp_path / "answer.docx")
    explanation = exporter.export_explanation(blueprint, [question], tmp_path / "explanation.docx")
    answer_sheet = exporter.export_answer_sheet(blueprint, [question], tmp_path / "sheet.docx")

    for path in [student, answer, explanation, answer_sheet]:
        assert path.exists()
        assert Document(path).paragraphs

    student_text = "\n".join(paragraph.text for paragraph in Document(student).paragraphs)
    assert "Answer: __________________" in student_text
    assert "Answer: A" not in student_text
    assert "Explanation:" not in student_text


def test_pptx_exporter_writes_editable_presentation(tmp_path):
    outline = SlideOutlineContent(
        course_name="AI",
        chapter="Search",
        lesson_id="lesson-1",
        slides=[
            SlideItem(
                slide_index=1,
                slide_type="title",
                title="AI Search",
                bullet_points=["2 sessions"],
            ),
            SlideItem(
                slide_index=2,
                slide_type="content",
                title="State Space",
                bullet_points=["Define state", "Define action"],
                references=[Reference(chunk_id="chunk-1")],
            ),
            SlideItem(
                slide_index=3,
                slide_type="references",
                title="References",
                bullet_points=["chunk-1"],
                references=[Reference(chunk_id="chunk-1")],
            ),
        ],
    )

    output = PPTXExporter().export(outline, tmp_path / "slides.pptx")

    assert output.exists()
    presentation = Presentation(output)
    assert len(presentation.slides) == 3
    assert presentation.slides[0].shapes.title.text == "AI Search"
