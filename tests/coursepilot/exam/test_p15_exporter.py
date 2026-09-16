from pathlib import Path

import pytest
from docx import Document

from coursepilot.domain.exam import ExamArtifact
from coursepilot.exporters.exam_versioned_exporter import ExamVersionedExporter

from .test_p15_exam_workflow import blueprint, question


def artifact(*, approved: bool) -> ExamArtifact:
    bp = blueprint()
    return ExamArtifact(
        task_id="task-1",
        course_id="course-1",
        blueprint=bp,
        questions=[question(bp.slots[0]), question(bp.slots[1])],
        blueprint_approved=True,
        global_review_approved=approved,
    )


def test_export_requires_global_review(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="GLOBAL_REVIEW_REQUIRED"):
        ExamVersionedExporter().export(artifact(approved=False), tmp_path)


def test_export_publishes_four_roles(tmp_path: Path) -> None:
    paths = ExamVersionedExporter().export(artifact(approved=True), tmp_path)
    assert set(paths) == {"student_exam", "answer_key", "detailed_explanation", "answer_sheet"}
    assert all(path.exists() for path in paths.values())
    rendered = Document(str(paths["student_exam"]))
    assert rendered.paragraphs
    assert any("Template:" in paragraph.text for paragraph in rendered.paragraphs)


def test_export_does_not_overwrite_existing_operation_directory(tmp_path: Path) -> None:
    destination = tmp_path / "operation"
    ExamVersionedExporter().export(artifact(approved=True), destination)
    with pytest.raises(FileExistsError, match="EXPORT_TARGET_EXISTS"):
        ExamVersionedExporter().export(artifact(approved=True), destination)
