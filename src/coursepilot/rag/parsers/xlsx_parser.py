from pathlib import Path

from openpyxl import load_workbook

from coursepilot.rag.parsers.base import BaseParser
from coursepilot.rag.types import ParsedDocument, ParsedSection


class XLSXParser(BaseParser):
    supported_suffixes = {".xlsx"}

    def parse(self, path: str | Path) -> ParsedDocument:
        path = Path(path)
        workbook = load_workbook(path, read_only=True, data_only=True)
        sections: list[ParsedSection] = []
        try:
            for sheet in workbook.worksheets:
                rows: list[str] = []
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value).strip() for value in row if value is not None]
                    if values:
                        rows.append(" | ".join(values))
                if rows:
                    sections.append(
                        ParsedSection(
                            content="\n".join(rows),
                            title=sheet.title,
                            metadata={"sheet": sheet.title},
                        )
                    )
        finally:
            workbook.close()
        return ParsedDocument(source_path=str(path), sections=sections)

