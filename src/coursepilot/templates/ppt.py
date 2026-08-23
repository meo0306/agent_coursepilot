from __future__ import annotations

import hashlib
import mimetypes
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from pptx import Presentation
from pydantic import Field

from coursepilot.domain.common import DomainModel, canonical_sha256


class PPTLayoutProfile(DomainModel):
    master_index: int = Field(ge=0)
    layout_index: int = Field(ge=0)
    name: str = Field(min_length=1)
    placeholder_roles: tuple[str, ...] = ()


class PPTTemplateProfile(DomainModel):
    template_id: str = Field(min_length=1)
    version: str = Field(min_length=1)
    source_sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
    master_count: int = Field(ge=1)
    layout_count: int = Field(ge=1)
    layouts: list[PPTLayoutProfile] = Field(min_length=1)
    slide_type_layout_map: dict[str, str] = Field(default_factory=dict)
    required_placeholders: tuple[str, ...] = ()
    license: str = "project-mit-attribution"

    @property
    def profile_sha256(self) -> str:
        return canonical_sha256(self.model_dump(mode="json"))


class PPTTemplateSnapshot(DomainModel):
    template_id: str = Field(min_length=1)
    version: str = Field(min_length=1)
    profile_sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
    source_sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
    resource_uri: str = Field(min_length=1)

    @property
    def snapshot_sha256(self) -> str:
        return canonical_sha256(self.model_dump(mode="json"))


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _placeholder_role(shape: Any) -> str | None:
    if not getattr(shape, "is_placeholder", False):
        return None
    name = str(getattr(shape, "name", "")).lower()
    placeholder_type = str(getattr(shape.placeholder_format, "type", "")).lower()
    if "picture" in name or "picture" in placeholder_type:
        return "picture"
    if "slide number" in name or "slide number" in placeholder_type:
        return "slide_number"
    if "footer" in name or "footer" in placeholder_type:
        return "footer"
    if "title" in name or "title" in placeholder_type:
        return "title"
    if "body" in name or "content" in name or "body" in placeholder_type:
        return "body"
    return "placeholder"


def inspect_pptx(
    path: str | Path, *, template_id: str, version: str = "1.0.0"
) -> PPTTemplateProfile:
    """Inspect every Master/Layout and fail closed for malformed archives."""
    source = Path(path)
    if source.suffix.lower() not in {".pptx", ".potx"}:
        raise ValueError("PPT_TEMPLATE_TYPE_UNSUPPORTED")
    mime = mimetypes.guess_type(source.name)[0]
    if mime not in {
        None,
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "application/vnd.openxmlformats-officedocument.presentationml.template",
    }:
        raise ValueError("PPT_TEMPLATE_MIME_MISMATCH")
    try:
        with ZipFile(source) as archive:
            if archive.testzip() is not None:
                raise ValueError("PPT_TEMPLATE_ARCHIVE_INVALID")
            if len(archive.infolist()) > 20000:
                raise ValueError("PPT_TEMPLATE_ARCHIVE_TOO_LARGE")
    except (BadZipFile, OSError) as exc:
        raise ValueError("PPT_TEMPLATE_ARCHIVE_INVALID") from exc
    try:
        presentation = Presentation(str(source))
    except Exception as exc:
        raise ValueError("PPT_TEMPLATE_OPEN_FAILED") from exc
    layouts: list[PPTLayoutProfile] = []
    for master_index, master in enumerate(presentation.slide_masters):
        for layout_index, layout in enumerate(master.slide_layouts):
            roles = tuple(
                sorted({role for role in (_placeholder_role(s) for s in layout.shapes) if role})
            )
            layouts.append(
                PPTLayoutProfile(
                    master_index=master_index,
                    layout_index=layout_index,
                    name=str(layout.name),
                    placeholder_roles=roles,
                )
            )
    if not layouts:
        raise ValueError("PPT_TEMPLATE_NO_LAYOUTS")
    return PPTTemplateProfile(
        template_id=template_id,
        version=version,
        source_sha256=sha256_file(source),
        master_count=len(presentation.slide_masters),
        layout_count=len(layouts),
        layouts=layouts,
        slide_type_layout_map={},
        required_placeholders=("title", "body"),
    )


def apply_mapping(
    profile: PPTTemplateProfile,
    *,
    slide_type_layout_map: dict[str, str],
    required_placeholders: tuple[str, ...] = ("title", "body"),
) -> PPTTemplateProfile:
    names = {layout.name for layout in profile.layouts}
    missing = set(slide_type_layout_map.values()) - names
    if missing:
        raise ValueError(f"PPT_TEMPLATE_LAYOUT_MISSING:{','.join(sorted(missing))}")
    mapped = [
        layout for layout in profile.layouts if layout.name in set(slide_type_layout_map.values())
    ]
    available = (
        set().union(*(set(layout.placeholder_roles) for layout in mapped)) if mapped else set()
    )
    missing_roles = set(required_placeholders) - available
    if missing_roles:
        raise ValueError(f"PPT_TEMPLATE_PLACEHOLDER_MISSING:{','.join(sorted(missing_roles))}")
    return profile.model_copy(
        update={
            "slide_type_layout_map": dict(slide_type_layout_map),
            "required_placeholders": required_placeholders,
        }
    )
