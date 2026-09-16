import argparse
import hashlib
import json
import os
import sqlite3
import zipfile
from collections import Counter
from collections.abc import Sequence
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, TypedDict
from xml.etree import ElementTree as ET

from dotenv import load_dotenv
from sqlalchemy import text

load_dotenv(Path(os.getenv("COURSEPILOT_EVAL_ENV_FILE", ".env.eval")), override=True)

from client.coursepilot_client import CoursePilotClient  # noqa: E402
from core.settings import settings  # noqa: E402
from coursepilot.db.session import get_coursepilot_engine  # noqa: E402
from coursepilot.evals.checkpoint import EvalCheckpointRunner  # noqa: E402
from coursepilot.evals.metrics import (  # noqa: E402
    AnswerCompletenessCase,
    CitationCase,
    CoursePilotEvalInput,
    CoursePilotEvaluator,
    ExportCase,
    QuestionCountCase,
    RetrievalCase,
    SchemaCase,
)
from coursepilot.rag.vector_store import collection_name_for_course  # noqa: E402
from coursepilot.token_usage import get_deepseek_tokenizer  # noqa: E402


class RetrievalProbe(TypedDict):
    query: str
    expected_terms: list[str]


DEFAULT_RETRIEVAL_CASES: list[RetrievalProbe] = [
    {"query": "状态空间搜索", "expected_terms": ["状态空间", "搜索"]},
    {"query": "启发式搜索", "expected_terms": ["启发式", "搜索"]},
    {"query": "知识图谱", "expected_terms": ["知识图谱"]},
    {"query": "机器学习", "expected_terms": ["机器学习"]},
    {"query": "智能体", "expected_terms": ["智能体"]},
]

EVAL_GENERATION_TASK_TYPES = (
    "lesson_design",
    "exam_blueprint",
    "exam_questions",
    "ppt_outline",
)


def main() -> None:
    runtime: dict[str, Any] = {}
    try:
        _main(runtime)
    except BaseException as exc:
        runner = runtime.get("runner")
        if runner is not None and not runtime.get("finished"):
            runner.fail(exc)
        raise


def _main(runtime: dict[str, Any]) -> None:
    parser = argparse.ArgumentParser(description="Run a real CoursePilot evaluation workflow.")
    parser.add_argument("--base-url", default=os.getenv("AGENT_URL", "http://localhost:8080"))
    parser.add_argument("--sample-dir", default="data/coursepilot_sample")
    parser.add_argument("--output", default="storage_eval/coursepilot_real_eval_report.json")
    parser.add_argument(
        "--chroma-dir", default=os.getenv("COURSEPILOT_CHROMA_DIR", "chroma_db_eval")
    )
    parser.add_argument("--auth-token", default=os.getenv("AUTH_SECRET", ""))
    parser.add_argument("--top-k", type=int, default=8)
    parser.add_argument(
        "--build-kb-timeout",
        type=float,
        default=float(os.getenv("COURSEPILOT_EVAL_BUILD_KB_TIMEOUT_SECONDS", "7200")),
        help="Seconds to wait for each document knowledge-base build.",
    )
    parser.add_argument(
        "--generation-timeout",
        type=float,
        default=float(os.getenv("COURSEPILOT_EVAL_GENERATION_TIMEOUT_SECONDS", "1800")),
        help="Seconds to poll each lesson, exam, or PPT generation task.",
    )
    parser.add_argument("--skip-generation", action="store_true")
    parser.add_argument("--allow-fallback", action="store_true")
    parser.add_argument(
        "--checkpoint",
        default=None,
        help="Checkpoint path. Defaults to <output-stem>.checkpoint.json.",
    )
    parser.add_argument(
        "--resume",
        action="store_true",
        help="Resume from the checkpoint and skip succeeded steps.",
    )
    parser.add_argument(
        "--force-resume",
        action="store_true",
        help="Resume even when configuration or sample file fingerprints differ.",
    )
    parser.add_argument(
        "--overwrite-checkpoint",
        action="store_true",
        help="Discard an existing checkpoint and start a new evaluation run.",
    )
    args = parser.parse_args()
    if args.resume and args.overwrite_checkpoint:
        parser.error("--resume and --overwrite-checkpoint cannot be used together")
    if args.force_resume and not args.resume:
        parser.error("--force-resume requires --resume")

    output_path = Path(args.output)
    checkpoint_path = (
        Path(args.checkpoint)
        if args.checkpoint
        else output_path.with_name(f"{output_path.stem}.checkpoint.json")
    )
    if checkpoint_path.resolve() == output_path.resolve():
        parser.error("--checkpoint and --output must use different paths")
    sample_dir = Path(args.sample_dir)
    identity = evaluation_identity(args, sample_dir)
    runner = EvalCheckpointRunner(
        checkpoint_path=checkpoint_path,
        output_path=output_path,
        fingerprint=evaluation_fingerprint(identity),
        metadata=identity,
        resume=args.resume,
        force_resume=args.force_resume,
        overwrite=args.overwrite_checkpoint,
    )
    runtime["runner"] = runner
    headers = {"Authorization": f"Bearer {args.auth_token}"} if args.auth_token else None
    client = CoursePilotClient(base_url=args.base_url, headers=headers, timeout=300)

    course_step_id = "course.create"
    course = runner.call(
        course_step_id,
        "create_course",
        lambda: client.create_course(
            {
                "course_name": "人工智能：从算法到系统",
                "course_type": "demo_eval",
                "student_level": "本科",
                "student_background": "具备基础编程和数学基础",
                "description": "CoursePilot real evaluation sample course.",
            },
            idempotency_key=runner.idempotency_key(course_step_id),
        ),
    )
    course_id = course["result"]["id"]
    runner.update_context(course_id=course_id)

    upload_results = []
    build_results = []
    for path in sorted(sample_dir.iterdir()):
        if not path.is_file():
            continue
        source_type = source_type_for_file(path)
        file_digest = file_sha256(path)
        step_suffix = f"{path.name}:{file_digest[:16]}"

        def upload_current_document() -> dict[str, Any]:
            return client.upload_document(
                course_id,
                filename=path.name,
                content=path.read_bytes(),
                source_type=source_type,
            )

        upload = runner.call(
            f"document.upload:{step_suffix}",
            f"upload:{path.name}",
            upload_current_document,
        )
        upload_results.append(upload)
        document_id = str(upload["result"]["id"])
        documents = dict(runner.state["context"].get("documents", {}))
        documents[path.name] = {
            "document_id": document_id,
            "sha256": file_digest,
        }
        runner.update_context(documents=documents)
        build_step_id = f"document.build_kb:{step_suffix}"

        def build_current_document() -> dict[str, Any]:
            return client.build_kb(
                document_id,
                timeout=args.build_kb_timeout,
                idempotency_key=runner.idempotency_key(build_step_id),
            )

        build_results.append(
            runner.call(
                build_step_id,
                f"build_kb:{path.name}",
                build_current_document,
            )
        )

    retrieval_cases, retrieval_details = run_retrieval_cases(
        client,
        course_id,
        args.top_k,
        runner,
    )

    workflow_results: dict[str, Any] = {}
    task_ids: list[str] = []
    if not args.skip_generation:
        workflow_results, task_ids = run_generation_workflows(
            client,
            course_id,
            runner,
            generation_timeout=args.generation_timeout,
        )

    db_stats = read_db_stats(course_id, task_ids)
    validate_task_coverage(task_ids, db_stats["tasks"])
    chroma_stats = read_chroma_stats(Path(args.chroma_dir), course_id)
    chunk_token_stats = read_chroma_chunk_token_stats(Path(args.chroma_dir), course_id)

    eval_payload = build_eval_payload(
        retrieval_cases=retrieval_cases,
        workflow_results=workflow_results,
    )
    eval_report = CoursePilotEvaluator().evaluate(eval_payload).model_dump()
    fallback_violations = fallback_violations_from_tasks(db_stats.get("tasks", []))

    report = {
        "generated_at": datetime.now(UTC).isoformat(),
        "elapsed_ms": runner.total_step_latency_ms,
        "strict_fallback_expected": not args.allow_fallback,
        "strict_fallback_passed": not fallback_violations,
        "fallback_violations": fallback_violations,
        "environment": {
            "base_url": args.base_url,
            "sample_dir": str(sample_dir),
            "chroma_dir": str(args.chroma_dir),
            "generation_mode": settings.COURSEPILOT_GENERATION_MODE,
            "disable_deterministic_fallback": settings.COURSEPILOT_DISABLE_DETERMINISTIC_FALLBACK,
            "embedding_provider": settings.COURSEPILOT_EMBEDDING_PROVIDER,
            "embedding_model": settings.COURSEPILOT_EMBEDDING_MODEL,
            "knowledge_points_mode": settings.COURSEPILOT_RAG_KNOWLEDGE_POINTS_MODE,
            "build_kb_timeout_seconds": args.build_kb_timeout,
            "generation_timeout_seconds": args.generation_timeout,
        },
        "data_card": {
            **sample_data_card(sample_dir),
            "course_id": course_id,
            "db_chunk_count": db_stats.get("chunk_count"),
            "db_chunk_count_by_source_type": db_stats.get("chunk_count_by_source_type"),
            "vector_entries": chroma_stats.get("entries"),
            "vector_dimension": chroma_stats.get("dimension"),
            "avg_chunk_tokens": chunk_token_stats.get("avg_chunk_tokens"),
            "avg_chunk_chars": chunk_token_stats.get("avg_chunk_chars"),
        },
        "rag_config": {
            "parsers": {
                ".docx": "docx2txt",
                ".xlsx": "openpyxl read_only data_only",
                ".pdf": "PyMuPDF text extraction",
                ".md": "Markdown parser",
                ".txt": "TXT parser",
            },
            "chunk_size_chars": 1000,
            "overlap_chars": 150,
            "vector_store": "Chroma",
            "persist_directory": str(args.chroma_dir),
            "top_k_eval": args.top_k,
            "api_default_top_k": 5,
            "workflow_top_k": 8,
            "metadata_filters": ["course_id", "chapter", "source_type", "verified_only"],
            "rerank": False,
        },
        "ingestion": {
            "uploads": upload_results,
            "builds": build_results,
            "db_chroma_count_match": db_stats.get("chunk_count") == chroma_stats.get("entries"),
        },
        "retrieval": {
            "cases": retrieval_details,
            "metrics": {
                key: eval_report[key]
                for key in (
                    "rag_recall_at_k",
                    "rag_hit_at_k",
                    "rag_mrr",
                    "rag_ndcg",
                    "context_precision",
                )
            },
        },
        "generation": {
            "workflow_results": workflow_results,
            "tasks": db_stats.get("tasks", []),
            "metrics": {
                key: eval_report[key]
                for key in (
                    "citation_coverage",
                    "schema_pass_rate",
                    "question_count_accuracy",
                    "duplicate_rate",
                    "answer_completeness",
                    "export_success_rate",
                )
            },
        },
    }

    completion_status = (
        "completed_with_fallback_violations"
        if fallback_violations and not args.allow_fallback
        else "completed"
    )
    runner.complete(report, status=completion_status)
    runtime["finished"] = True
    print(json.dumps(report, ensure_ascii=False, indent=2))

    enforce_strict_fallback(fallback_violations, allow_fallback=args.allow_fallback)


def source_type_for_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return "knowledge_graph"
    if suffix == ".docx":
        return "textbook"
    return "unknown"


def evaluation_identity(args: argparse.Namespace, sample_dir: Path) -> dict[str, Any]:
    sample_files = [
        {
            "name": path.name,
            "size_bytes": path.stat().st_size,
            "sha256": file_sha256(path),
        }
        for path in sorted(sample_dir.iterdir())
        if path.is_file()
    ]
    return {
        "base_url": args.base_url,
        "sample_dir": str(sample_dir.resolve()),
        "sample_files": sample_files,
        "chroma_dir": str(Path(args.chroma_dir).resolve()),
        "top_k": args.top_k,
        "skip_generation": args.skip_generation,
        "allow_fallback": args.allow_fallback,
        "build_kb_timeout_seconds": args.build_kb_timeout,
        "generation_timeout_seconds": args.generation_timeout,
        "generation_mode": settings.COURSEPILOT_GENERATION_MODE,
        "disable_deterministic_fallback": settings.COURSEPILOT_DISABLE_DETERMINISTIC_FALLBACK,
        "compatible_model": settings.COMPATIBLE_MODEL,
        "embedding_provider": settings.COURSEPILOT_EMBEDDING_PROVIDER,
        "embedding_model": settings.COURSEPILOT_EMBEDDING_MODEL,
        "embedding_base_url": settings.COURSEPILOT_EMBEDDING_BASE_URL,
        "knowledge_points_mode": settings.COURSEPILOT_RAG_KNOWLEDGE_POINTS_MODE,
    }


def payload_fingerprint(payload: dict[str, Any]) -> str:
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def evaluation_fingerprint(identity: dict[str, Any]) -> str:
    comparable_identity = dict(identity)
    comparable_identity.pop("build_kb_timeout_seconds", None)
    comparable_identity.pop("generation_timeout_seconds", None)
    return payload_fingerprint(comparable_identity)


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def run_retrieval_cases(
    client: CoursePilotClient,
    course_id: str,
    top_k: int,
    runner: EvalCheckpointRunner,
) -> tuple[list[RetrievalCase], list[dict[str, Any]]]:
    cases: list[RetrievalCase] = []
    details: list[dict[str, Any]] = []
    for index, case in enumerate(DEFAULT_RETRIEVAL_CASES, start=1):
        query = case["query"]

        def search_current_case() -> dict[str, Any]:
            return client.search_kb(
                course_id,
                {"query": query, "top_k": top_k},
            )

        response = runner.call(
            f"retrieval.case:{index}",
            f"search:{query}",
            search_current_case,
        )
        results = response["result"].get("results", [])
        retrieved_ids = [item["chunk_id"] for item in results]
        relevant_ids = [
            item["chunk_id"]
            for item in results
            if content_matches_terms(item.get("content", ""), case["expected_terms"])
        ]
        expected_ids = relevant_ids or [f"__missing__:{case['query']}"]
        cases.append(
            RetrievalCase(
                query=case["query"],
                expected_chunk_ids=expected_ids,
                retrieved_chunk_ids=retrieved_ids,
            )
        )
        details.append(
            {
                "query": case["query"],
                "expected_terms": case["expected_terms"],
                "latency_ms": response["latency_ms"],
                "retrieved_chunk_ids": retrieved_ids,
                "relevant_retrieved_chunk_ids": relevant_ids,
                "top_scores": [item.get("score") for item in results],
            }
        )
    return cases, details


def content_matches_terms(content: str, terms: Sequence[str]) -> bool:
    normalized = content.lower()
    return all(term.lower() in normalized for term in terms)


def run_generation_workflows(
    client: CoursePilotClient,
    course_id: str,
    runner: EvalCheckpointRunner,
    *,
    generation_timeout: float,
) -> tuple[dict[str, Any], list[str]]:
    workflows: dict[str, Any] = {}
    task_ids: list[str] = []

    lesson_payload = {
        "chapter_range": "状态空间搜索、知识表示与知识图谱",
        "total_sessions": 2,
        "session_duration": 45,
        "student_level": "本科",
        "student_background": "具备基础编程和数学基础",
        "teaching_focus": "状态空间搜索、启发式搜索、知识图谱",
        "include_interaction": True,
        "include_homework": True,
        "additional_requirements": "强调课程材料引用，输出适合课堂使用的教学设计。",
    }
    lesson_step_id = "generation.lesson"
    lesson = runner.call(
        lesson_step_id,
        "generate_lesson",
        lambda: client.generate_lesson(
            course_id,
            lesson_payload,
            timeout=generation_timeout,
            idempotency_key=runner.idempotency_key(lesson_step_id),
        ),
    )
    workflows["lesson"] = lesson
    task_ids.append(lesson["result"]["task_id"])

    lesson_id = lesson["result"]["lesson_id"]
    runner.update_context(lesson_id=lesson_id)
    workflows["lesson_export"] = runner.call(
        "generation.lesson_export",
        "export_lesson",
        lambda: client.export_lesson(lesson_id),
    )

    exam_payload = {
        "chapter_range": "状态空间搜索、知识表示与知识图谱",
        "generation_type": "exam",
        "question_counts": {
            "single_choice": 2,
            "multiple_choice": 1,
            "judgement": 1,
            "short_answer": 1,
        },
        "score_per_question": {
            "single_choice": 2,
            "multiple_choice": 4,
            "judgement": 2,
            "short_answer": 10,
        },
        "difficulty_distribution": {"easy": 0.3, "medium": 0.5, "hard": 0.2},
        "include_answer": True,
        "include_explanation": True,
        "include_answer_sheet": True,
        "additional_requirements": "题目必须基于上传课程材料，并附引用。",
    }
    blueprint_step_id = "generation.exam_blueprint"
    blueprint = runner.call(
        blueprint_step_id,
        "create_exam_blueprint",
        lambda: client.create_exam_blueprint(
            course_id,
            exam_payload,
            timeout=generation_timeout,
            idempotency_key=runner.idempotency_key(blueprint_step_id),
        ),
    )
    workflows["exam_blueprint"] = blueprint
    task_ids.append(blueprint["result"]["task_id"])
    blueprint_id = blueprint["result"]["blueprint_id"]
    runner.update_context(blueprint_id=blueprint_id)
    workflows["exam_confirm"] = runner.call(
        "generation.exam_confirm",
        "confirm_exam_blueprint",
        lambda: client.confirm_exam_blueprint(blueprint_id),
    )
    questions_enqueue_step_id = "generation.exam_questions.enqueue"
    questions_accepted = runner.call(
        questions_enqueue_step_id,
        "enqueue_questions",
        lambda: client.enqueue_questions(
            blueprint_id,
            idempotency_key=runner.idempotency_key(questions_enqueue_step_id),
        ),
    )
    question_task_id = str(questions_accepted["result"]["task_id"])
    task_ids.append(question_task_id)
    questions_step_id = "generation.exam_questions.wait"

    def wait_for_questions() -> dict[str, Any]:
        result = client.wait_for_task(question_task_id, timeout=generation_timeout)
        return {**result, "task_id": question_task_id}

    questions = runner.call(
        questions_step_id,
        "wait_for_task:exam_questions",
        wait_for_questions,
    )
    workflows["exam_questions"] = questions
    workflows["exam_export"] = runner.call(
        "generation.exam_export",
        "export_exam",
        lambda: client.export_exam(blueprint_id),
    )

    ppt_payload = {
        "slide_count": 8,
        "style_template": "standard",
        "include_references": True,
        "additional_requirements": "PPT 大纲要覆盖每个课时的目标、关键知识点和课堂活动。",
    }
    ppt_step_id = "generation.ppt"
    ppt = runner.call(
        ppt_step_id,
        "generate_ppt_outline",
        lambda: client.generate_ppt_outline(
            lesson_id,
            ppt_payload,
            timeout=generation_timeout,
            idempotency_key=runner.idempotency_key(ppt_step_id),
        ),
    )
    workflows["ppt"] = ppt
    task_ids.append(ppt["result"]["task_id"])
    outline_id = ppt["result"]["outline_id"]
    runner.update_context(outline_id=outline_id)
    workflows["ppt_export"] = runner.call(
        "generation.ppt_export",
        "export_ppt",
        lambda: client.export_ppt(outline_id),
    )

    return workflows, task_ids


def build_eval_payload(
    *,
    retrieval_cases: list[RetrievalCase],
    workflow_results: dict[str, Any],
) -> CoursePilotEvalInput:
    citation_cases: list[CitationCase] = []
    schema_cases: list[SchemaCase] = []
    question_count_cases: list[QuestionCountCase] = []
    answer_cases: list[AnswerCompletenessCase] = []
    export_cases: list[ExportCase] = []
    duplicate_rate = 0.0

    lesson = workflow_results.get("lesson", {}).get("result")
    if lesson:
        report = lesson.get("validation_report", {})
        schema_cases.append(
            SchemaCase(item_id="lesson", schema_valid=bool(report.get("schema_valid", False)))
        )
        for session in lesson.get("lesson_design", {}).get("sessions", []):
            citation_cases.append(
                CitationCase(
                    item_id=f"lesson-session-{session.get('session_index')}",
                    citation_chunk_ids=[
                        ref["chunk_id"]
                        for ref in session.get("references", [])
                        if ref.get("chunk_id")
                    ],
                )
            )

    questions = workflow_results.get("exam_questions", {}).get("result")
    if questions:
        report = questions.get("validation_report", {})
        schema_cases.append(
            SchemaCase(
                item_id="exam-questions", schema_valid=bool(report.get("schema_valid", False))
            )
        )
        duplicate_rate = float(report.get("duplicate_rate") or 0.0)
        actual_counts: dict[str, int] = {}
        for question in questions.get("questions", []):
            question_type = str(question.get("question_type"))
            actual_counts[question_type] = actual_counts.get(question_type, 0) + 1
            citation_cases.append(
                CitationCase(
                    item_id=f"question-{len(answer_cases) + 1}",
                    citation_chunk_ids=[
                        ref["chunk_id"]
                        for ref in question.get("references", [])
                        if ref.get("chunk_id")
                    ],
                )
            )
            answer_cases.append(
                AnswerCompletenessCase(
                    item_id=f"question-{len(answer_cases) + 1}",
                    has_answer=bool(question.get("correct_answer")),
                    has_explanation=bool(question.get("explanation")),
                )
            )
        question_count_cases.append(
            QuestionCountCase(
                expected_counts={
                    "single_choice": 2,
                    "multiple_choice": 1,
                    "judgement": 1,
                    "short_answer": 1,
                },
                actual_counts=actual_counts,
            )
        )

    ppt = workflow_results.get("ppt", {}).get("result")
    if ppt:
        report = ppt.get("validation_report", {})
        schema_cases.append(
            SchemaCase(item_id="ppt", schema_valid=bool(report.get("schema_valid", False)))
        )
        for slide in ppt.get("outline", {}).get("slides", []):
            if slide.get("slide_type") == "title":
                continue
            citation_cases.append(
                CitationCase(
                    item_id=f"slide-{slide.get('slide_index')}",
                    citation_chunk_ids=[
                        ref["chunk_id"]
                        for ref in slide.get("references", [])
                        if ref.get("chunk_id")
                    ],
                )
            )

    for key in ("lesson_export", "exam_export", "ppt_export"):
        result = workflow_results.get(key, {}).get("result")
        if not result:
            continue
        if "files" in result:
            export_cases.extend(
                ExportCase(file_path=item["file_path"], file_role=item["file_role"])
                for item in result["files"]
            )
        elif "file_path" in result:
            export_cases.append(
                ExportCase(
                    file_path=result["file_path"],
                    file_role=result.get("file_role", key),
                )
            )

    return CoursePilotEvalInput(
        retrieval_cases=retrieval_cases,
        citation_cases=citation_cases,
        schema_cases=schema_cases,
        question_count_cases=question_count_cases,
        duplicate_rate=duplicate_rate,
        answer_cases=answer_cases,
        export_cases=export_cases,
    )


def read_db_stats(course_id: str, task_ids: list[str]) -> dict[str, Any]:
    engine = get_coursepilot_engine()
    with engine.connect() as conn:
        chunk_count = conn.execute(
            text("select count(*) from coursepilot_chunks where course_id = :course_id"),
            {"course_id": course_id},
        ).scalar()
        source_rows = conn.execute(
            text(
                """
                select source_type, count(*) as count
                from coursepilot_chunks
                where course_id = :course_id
                group by source_type
                order by source_type
                """
            ),
            {"course_id": course_id},
        ).mappings()
        tasks = []
        if task_ids:
            task_rows = conn.execute(
                text(
                    """
                    select id, task_type, status, created_at, updated_at,
                           input_params_json, intermediate_outputs_json,
                           validation_report_json, error_message
                    from coursepilot_generation_tasks
                    where course_id = :course_id
                      and task_type = any(:task_types)
                    order by created_at
                    """
                ),
                {
                    "course_id": course_id,
                    "task_types": list(EVAL_GENERATION_TASK_TYPES),
                },
            ).mappings()
            for row in task_rows:
                item = dict(row)
                created_at = item.pop("created_at")
                updated_at = item.pop("updated_at")
                item["created_at"] = created_at.isoformat()
                item["updated_at"] = updated_at.isoformat()
                item["latency_ms"] = int((updated_at - created_at).total_seconds() * 1000)
                outputs = item.get("intermediate_outputs_json") or {}
                item["llm_usage_summary"] = outputs.get("llm_usage_summary")
                item["repair_attempts"] = (item.get("validation_report_json") or {}).get(
                    "repair_attempts"
                )
                tasks.append(item)
        return {
            "chunk_count": int(chunk_count or 0),
            "chunk_count_by_source_type": {
                row["source_type"]: int(row["count"]) for row in source_rows
            },
            "tasks": tasks,
        }


def validate_task_coverage(
    expected_task_ids: Sequence[str],
    tasks: Sequence[dict[str, Any]],
) -> None:
    expected = [str(task_id) for task_id in expected_task_ids]
    observed = [str(task.get("id")) for task in tasks]
    expected_counts = Counter(expected)
    observed_counts = Counter(observed)

    duplicate_expected = sorted(task_id for task_id, count in expected_counts.items() if count > 1)
    duplicate_observed = sorted(task_id for task_id, count in observed_counts.items() if count > 1)
    missing = sorted(set(expected_counts) - set(observed_counts))
    unexpected = sorted(set(observed_counts) - set(expected_counts))
    if duplicate_expected or duplicate_observed or missing or unexpected:
        raise ValueError(
            "Generation task coverage audit failed: "
            f"missing={missing}, unexpected={unexpected}, "
            f"duplicate_expected={duplicate_expected}, "
            f"duplicate_observed={duplicate_observed}"
        )


def read_chroma_stats(chroma_dir: Path, course_id: str) -> dict[str, Any]:
    db_path = chroma_dir / "chroma.sqlite3"
    if not db_path.exists():
        return {"error": f"Chroma sqlite file not found: {db_path}"}
    collection_name = collection_name_for_course(course_id)
    with sqlite3.connect(db_path) as conn:
        conn.row_factory = sqlite3.Row
        row = conn.execute(
            """
            select c.dimension, count(e.id) as entries
            from collections c
            left join segments s on s.collection = c.id and s.scope = 'METADATA'
            left join embeddings e on e.segment_id = s.id
            where c.name = ?
            group by c.dimension
            """,
            (collection_name,),
        ).fetchone()
    if row is None:
        return {"collection": collection_name, "entries": 0, "dimension": None}
    return {
        "collection": collection_name,
        "entries": int(row["entries"] or 0),
        "dimension": row["dimension"],
    }


def read_chroma_chunk_token_stats(chroma_dir: Path, course_id: str) -> dict[str, Any]:
    db_path = chroma_dir / "chroma.sqlite3"
    if not db_path.exists():
        return {}
    tokenizer = get_deepseek_tokenizer()
    texts: list[str] = []
    with sqlite3.connect(db_path) as conn:
        conn.row_factory = sqlite3.Row
        rows = conn.execute(
            """
            select text_meta.string_value as content
            from embeddings e
            join embedding_metadata course_meta
              on course_meta.id = e.id and course_meta.key = 'course_id'
            join embedding_metadata text_meta
              on text_meta.id = e.id and text_meta.key = 'chroma:document'
            where course_meta.string_value = ?
            """,
            (course_id,),
        ).fetchall()
        texts = [row["content"] for row in rows if row["content"]]
    if not texts:
        return {}
    token_counts = [
        len(tokenizer.encode(text).ids) if tokenizer is not None else max(1, len(text) // 2)
        for text in texts
    ]
    char_counts = [len(text) for text in texts]
    return {
        "avg_chunk_tokens": round(sum(token_counts) / len(token_counts), 2),
        "avg_chunk_chars": round(sum(char_counts) / len(char_counts), 2),
    }


def fallback_violations_from_tasks(tasks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    violations = []
    for task in tasks:
        violation = {
            "task_id": task.get("id"),
            "task_type": task.get("task_type"),
        }
        usage = task.get("llm_usage_summary")
        if not isinstance(usage, dict):
            violations.append(
                {
                    **violation,
                    "fallback_count": None,
                    "reason": "missing_or_invalid_llm_usage_summary",
                }
            )
            continue
        if "fallback_count" not in usage:
            violations.append(
                {
                    **violation,
                    "fallback_count": None,
                    "reason": "missing_fallback_count",
                }
            )
            continue
        fallback_count = usage["fallback_count"]
        if (
            isinstance(fallback_count, bool)
            or not isinstance(fallback_count, int)
            or fallback_count < 0
        ):
            violations.append(
                {
                    **violation,
                    "fallback_count": fallback_count,
                    "reason": "invalid_fallback_count",
                }
            )
            continue
        if fallback_count > 0:
            violations.append(
                {
                    **violation,
                    "fallback_count": fallback_count,
                    "reason": "fallback_used",
                }
            )
    return violations


def enforce_strict_fallback(
    fallback_violations: Sequence[dict[str, Any]],
    *,
    allow_fallback: bool,
) -> None:
    if fallback_violations and not allow_fallback:
        raise SystemExit("Fallback audit failed during real evaluation.")


def sample_data_card(sample_dir: Path) -> dict[str, Any]:
    files = [path for path in sorted(sample_dir.iterdir()) if path.is_file()]
    by_extension: dict[str, int] = {}
    file_cards = []
    for path in files:
        by_extension[path.suffix.lower()] = by_extension.get(path.suffix.lower(), 0) + 1
        file_cards.append(file_card(path))
    return {
        "course_count": 1,
        "file_count": len(files),
        "file_types": by_extension,
        "total_mb": round(sum(path.stat().st_size for path in files) / 1024 / 1024, 2),
        "files": file_cards,
    }


def file_card(path: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()
    card: dict[str, Any] = {
        "name": path.name,
        "extension": suffix,
        "size_mb": round(path.stat().st_size / 1024 / 1024, 2),
        "source_type": source_type_for_file(path),
    }
    if suffix == ".docx":
        card["pages"] = docx_app_property(path, "Pages")
        card["words"] = docx_app_property(path, "Words")
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                card["sheets"] = len(workbook.worksheets)
                card["rows"] = sum(sheet.max_row for sheet in workbook.worksheets)
            finally:
                workbook.close()
        except Exception as exc:
            card["xlsx_error"] = str(exc)
    return card


def docx_app_property(path: Path, name: str) -> int | None:
    try:
        with zipfile.ZipFile(path) as archive:
            data = archive.read("docProps/app.xml")
        root = ET.fromstring(data)
        namespace = {
            "ep": "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
        }
        node = root.find(f"ep:{name}", namespace)
        if node is not None and node.text and node.text.isdigit():
            return int(node.text)
    except Exception:
        return None
    return None


if __name__ == "__main__":
    main()
