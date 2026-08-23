from __future__ import annotations

import hashlib
import shutil
import subprocess
from pathlib import Path
from typing import Protocol

from pptx import Presentation

from coursepilot.domain.ppt import PPTRenderReport
from coursepilot.templates.ppt import sha256_file


class PPTXRendererPort(Protocol):
    def render(
        self, path: str | Path, output_dir: str | Path, *, expected_slide_count: int | None = None
    ) -> PPTRenderReport: ...


def inspect_renderable_pptx(
    path: str | Path,
    *,
    expected_slide_count: int | None = None,
    renderer: str = "python-pptx-static",
) -> PPTRenderReport:
    source = Path(path)
    presentation = Presentation(str(source))
    out_of_bounds = 0
    empty = 0
    editable = 0
    for slide in presentation.slides:
        for shape in slide.shapes:
            if (
                shape.left < 0
                or shape.top < 0
                or shape.left + shape.width > presentation.slide_width
                or shape.top + shape.height > presentation.slide_height
            ):
                out_of_bounds += 1
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                editable += 1
        if not any(
            getattr(shape, "has_text_frame", False) and shape.text.strip() for shape in slide.shapes
        ):
            empty += 1
    count = len(presentation.slides)
    warnings: list[str] = []
    if expected_slide_count is not None and count != expected_slide_count:
        warnings.append(f"slide_count_expected:{expected_slide_count}:actual:{count}")
    return PPTRenderReport(
        renderer=renderer,
        renderer_version="1",
        source_pptx_sha256=sha256_file(source),
        slide_count=count,
        opened=True,
        # python-pptx inspection proves package structure only; it is not a renderer.
        rendered=False,
        out_of_bounds_count=out_of_bounds,
        empty_required_placeholder_count=empty,
        editable_object_count=editable,
        warnings=warnings,
    )


def render_with_libreoffice(
    path: str | Path,
    output_dir: str | Path,
    *,
    expected_slide_count: int | None = None,
    renderer_version: str = "libreoffice-7.4.7.2-poppler-22.12",
) -> PPTRenderReport:
    """Render using the fixed QA toolchain; never labels static inspection as rendered."""
    source = Path(path)
    destination = Path(output_dir)
    destination.mkdir(parents=True, exist_ok=True)
    static = inspect_renderable_pptx(source, expected_slide_count=expected_slide_count)
    if not shutil.which("soffice"):
        return static.model_copy(
            update={
                "renderer": "libreoffice",
                "renderer_version": renderer_version,
                "rendered": False,
                "warnings": ["SOFFICE_UNAVAILABLE"],
            }
        )
    try:
        subprocess.run(
            [
                "soffice",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(destination),
                str(source),
            ],
            check=True,
            capture_output=True,
            text=True,
            timeout=180,
        )
        pdf = destination / f"{source.stem}.pdf"
        if not pdf.exists() or not shutil.which("pdftoppm"):
            raise RuntimeError("PDF_RENDER_OUTPUT_MISSING")
        first = destination / "pass1"
        second = destination / "pass2"
        first.mkdir(exist_ok=True)
        second.mkdir(exist_ok=True)
        for target in (first, second):
            subprocess.run(
                ["pdftoppm", "-png", "-r", "144", str(pdf), str(target / "slide")],
                check=True,
                capture_output=True,
                timeout=180,
            )
        first_files = sorted(first.glob("slide-*.png"))
        second_files = sorted(second.glob("slide-*.png"))
        hashes = [hashlib.sha256(item.read_bytes()).hexdigest() for item in first_files]
        repeat = [hashlib.sha256(item.read_bytes()).hexdigest() for item in second_files]
        warnings = list(static.warnings)
        if hashes != repeat:
            warnings.append("RENDER_NON_REPRODUCIBLE")
        if expected_slide_count is not None and len(first_files) != expected_slide_count:
            warnings.append(
                f"slide_count_expected:{expected_slide_count}:actual:{len(first_files)}"
            )
        return static.model_copy(
            update={
                "renderer": "libreoffice",
                "renderer_version": renderer_version,
                "rendered_pdf_sha256": sha256_file(pdf),
                "slide_png_sha256s": hashes,
                "slide_count": len(first_files),
                "rendered": bool(first_files) and hashes == repeat,
                "warnings": warnings,
            }
        )
    except (OSError, subprocess.SubprocessError, RuntimeError) as exc:
        return static.model_copy(
            update={
                "renderer": "libreoffice",
                "renderer_version": renderer_version,
                "rendered": False,
                "warnings": [*static.warnings, f"RENDER_FAILED:{type(exc).__name__}"],
            }
        )
