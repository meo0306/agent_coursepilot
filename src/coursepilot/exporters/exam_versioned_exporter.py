from __future__ import annotations

import os
import shutil
import tempfile
from pathlib import Path
from typing import Literal

from docx import Document
from docx.document import Document as DocumentObject

from coursepilot.domain.exam import ExamArtifact, ExamQuestion

ExportRole = Literal["student_exam", "answer_key", "detailed_explanation", "answer_sheet"]


class ExamVersionedExporter:
    """Render the reviewed ExamArtifact into four independently inspectable DOCX files."""

    def __init__(self, *, template_path: str | Path | None = None) -> None:
        self.template_path = Path(
            template_path or "resources/templates/exporters/exam_default_v1.docx"
        )
        if not self.template_path.exists():
            raise FileNotFoundError(f"Exam DOCX template not found: {self.template_path}")

    def export(self, artifact: ExamArtifact, output_dir: str | Path) -> dict[str, Path]:
        if not artifact.global_review_approved:
            raise ValueError("GLOBAL_REVIEW_REQUIRED")
        root = Path(output_dir)
        if root.exists() and any(root.iterdir()):
            raise FileExistsError("EXPORT_TARGET_EXISTS")
        root.parent.mkdir(parents=True, exist_ok=True)
        staging = Path(tempfile.mkdtemp(prefix=".exam-export-", dir=str(root.parent)))
        published: list[Path] = []
        try:
            paths: dict[str, Path] = {}
            for role in ("student_exam", "answer_key", "detailed_explanation", "answer_sheet"):
                path = staging / f"{role}.docx"
                self._render(artifact, role, path)
                # Re-open before publishing the path to catch malformed OOXML.
                Document(str(path))
                paths[role] = root / path.name
            # Windows cannot atomically rename a directory over an existing
            # directory.  Rendering and OOXML validation already happened in
            # staging; publish the four validated files as one guarded phase.
            root.mkdir(parents=True, exist_ok=True)
            for staged_path in paths:
                destination = paths[staged_path]
                os.replace(staging / destination.name, destination)
                published.append(destination)
            shutil.rmtree(staging, ignore_errors=True)
            return paths
        except Exception:
            shutil.rmtree(staging, ignore_errors=True)
            for path in published:
                path.unlink(missing_ok=True)
            raise

    def _render(self, artifact: ExamArtifact, role: ExportRole, path: Path) -> None:
        doc = Document(str(self.template_path))
        self._fill_template(doc, artifact, role)
        for display_number, question in enumerate(
            sorted(artifact.questions, key=lambda item: item.question_number), start=1
        ):
            if role == "answer_sheet":
                doc.add_paragraph(f"{display_number}. ______________________________")
                continue
            self._write_question(doc, question, role, display_number=display_number)
        doc.save(str(path))

    @staticmethod
    def _remove_paragraph(paragraph: object) -> None:
        element = paragraph._element  # type: ignore[attr-defined]
        element.getparent().remove(element)

    def _fill_template(self, doc: DocumentObject, artifact: ExamArtifact, role: ExportRole) -> None:
        total_score = sum(question.score for question in artifact.questions)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    cell.text = cell.text.replace("{{ course_name }}", artifact.course_id)
                    cell.text = cell.text.replace("{{ scope }}", artifact.blueprint.chapter_range)
                    cell.text = cell.text.replace("{{ total_score }}", str(total_score))
        for paragraph in list(doc.paragraphs):
            text = paragraph.text
            if text == "课程试卷 / 作业":
                paragraph.text = f"{artifact.course_id} - {role.replace('_', ' ').title()}"
            elif "{{ blueprint_summary }}" in text:
                paragraph.text = f"Template: {artifact.blueprint.template_id}"
            elif "{{ question_blocks }}" in text:
                self._remove_paragraph(paragraph)
            elif "{{ answer_key_and_rubric }}" in text:
                if role in {"student_exam", "answer_sheet"}:
                    self._remove_paragraph(paragraph)
                else:
                    paragraph.text = "答案、解析与 Evidence 将在下方列出。"
            elif text == "参考答案与评分标准" and role in {"student_exam", "answer_sheet"}:
                self._remove_paragraph(paragraph)
            elif not text.strip():
                # The foundation template uses an empty paragraph carrying a
                # page break before its optional answer section.  Generated
                # content must not inherit that blank-page artifact.
                self._remove_paragraph(paragraph)

    @staticmethod
    def _write_question(
        doc: object,
        question: ExamQuestion,
        role: ExportRole,
        *,
        display_number: int,
    ) -> None:
        if question.stimulus:
            doc.add_paragraph(f"Material: {question.stimulus}")  # type: ignore[attr-defined]
        doc.add_paragraph(  # type: ignore[attr-defined]
            f"{display_number}. ({question.question_type}, {question.score} pts) {question.stem}"
        )
        if question.options:
            for key, value in question.options.items():
                doc.add_paragraph(f"{key}. {value}", style="List Bullet")  # type: ignore[attr-defined]
        if role in {"answer_key", "detailed_explanation"}:
            doc.add_paragraph(f"Answer: {question.answer}")  # type: ignore[attr-defined]
        if role == "detailed_explanation":
            doc.add_paragraph(f"Explanation: {question.explanation}")  # type: ignore[attr-defined]
            doc.add_paragraph(f"Evidence: {', '.join(question.evidence_ids)}")  # type: ignore[attr-defined]
