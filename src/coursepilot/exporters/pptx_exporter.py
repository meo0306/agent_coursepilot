from pathlib import Path
from typing import Any

from pptx import Presentation

from coursepilot.schemas.ppt_schema import SlideItem, SlideOutlineContent


class PPTXExporter:
    def export(self, outline: SlideOutlineContent, output_path: str | Path) -> Path:
        presentation: Any = Presentation()
        for slide in outline.slides:
            self._add_slide(presentation, slide)
        return self._save(presentation, output_path)

    def _add_slide(self, presentation: Any, slide_data: SlideItem) -> None:
        if slide_data.slide_type == "title":
            layout = presentation.slide_layouts[0]
            slide = presentation.slides.add_slide(layout)
            slide.shapes.title.text = slide_data.title
            if len(slide.placeholders) > 1:
                slide.placeholders[1].text = "\n".join(slide_data.bullet_points)
        else:
            layout = presentation.slide_layouts[1]
            slide = presentation.slides.add_slide(layout)
            slide.shapes.title.text = slide_data.title
            body = slide.placeholders[1].text_frame
            body.clear()
            for index, bullet in enumerate(slide_data.bullet_points):
                paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
                paragraph.text = bullet
                paragraph.level = 0

        notes_parts = []
        if slide_data.speaker_notes:
            notes_parts.append(slide_data.speaker_notes)
        for ref in slide_data.references:
            notes_parts.append(
                f"Reference: {ref.source_type or 'source'} | {ref.chapter or '-'} | "
                f"page={ref.page or '-'} | chunk={ref.chunk_id}"
            )
        if notes_parts:
            notes = slide.notes_slide.notes_text_frame
            notes.text = "\n".join(notes_parts)

    def _save(self, presentation: Any, output_path: str | Path) -> Path:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        presentation.save(str(output_path))
        return output_path
