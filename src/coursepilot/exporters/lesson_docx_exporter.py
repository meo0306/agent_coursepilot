"""
DOCX导出器
"""
from pathlib import Path

from docx import Document

from coursepilot.schemas.lesson_schema import LessonDesignContent


class LessonDocxExporter:
    def export(self, lesson_design: LessonDesignContent, output_path: str | Path) -> Path:
        # 导出路径
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        # 文档主体构建
        doc = Document()
        # 基本信息
        # 第一个大标题是课程名和章节
        doc.add_heading(f"{lesson_design.course_name} - {lesson_design.chapter}", level=0)
        doc.add_paragraph(f"Total sessions: {lesson_design.total_sessions}")
        doc.add_paragraph(f"Session duration: {lesson_design.session_duration} minutes")
        # 把知识点写成项目符号列表
        doc.add_heading("Knowledge Points", level=1)
        for point in lesson_design.knowledge_points:
            doc.add_paragraph(point, style="List Bullet")
        # 每个课时单独一个一级标题
        for session in lesson_design.sessions:
            doc.add_heading(f"Session {session.session_index}: {session.session_title}", level=1)
            # 课时内部再写 objectives、key points、difficult points、teaching process、interaction、homework、references。
            doc.add_heading("Objectives", level=2)
            for objective in session.teaching_objectives:
                doc.add_paragraph(objective, style="List Bullet")

            doc.add_heading("Key Points", level=2)
            for point in session.key_points:
                doc.add_paragraph(point, style="List Bullet")

            doc.add_heading("Difficult Points", level=2)
            for point in session.difficult_points:
                doc.add_paragraph(point, style="List Bullet")

            doc.add_heading("Teaching Process", level=2)
            for item in session.teaching_process:
                doc.add_paragraph(f"{item.stage} ({item.minutes} min): {item.content}")

            if session.interaction_design:
                doc.add_heading("Interaction Design", level=2)
                for item in session.interaction_design:
                    doc.add_paragraph(item, style="List Bullet")

            if session.homework_suggestion:
                doc.add_heading("Homework", level=2)
                for item in session.homework_suggestion:
                    doc.add_paragraph(item, style="List Bullet")

            doc.add_heading("References", level=2)
            for ref in session.references:
                doc.add_paragraph(
                    f"{ref.source_type or 'source'} | {ref.chapter or '-'} | "
                    f"page={ref.page or '-'} | chunk={ref.chunk_id}"
                )

        doc.save(output_path)
        return output_path

