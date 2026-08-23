from __future__ import annotations

from pathlib import Path

from docx import Document

from coursepilot.domain.lesson import LessonArtifact


class VersionedLessonDocxExporter:
    """Export a reviewed Lesson artifact from the pinned neutral DOCX template."""

    def __init__(
        self, template_path: str | Path = "resources/templates/exporters/lesson_default_v1.docx"
    ):
        self.template_path = Path(template_path)

    def export(self, artifact: LessonArtifact, output_path: str | Path) -> Path:
        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)
        document = Document(str(self.template_path)) if self.template_path.exists() else Document()
        document.add_heading(f"{artifact.chapter_scope}", level=0)
        document.add_paragraph(f"Course: {artifact.course_id}")
        document.add_paragraph(f"Sessions: {artifact.blueprint.total_sessions}")
        for session in artifact.sessions:
            document.add_heading(f"Session {session.session_index}: {session.title}", level=1)
            document.add_heading("Objectives", level=2)
            for objective in session.objectives:
                document.add_paragraph(objective, style="List Bullet")
            document.add_heading("Activities", level=2)
            for activity in session.activities:
                document.add_paragraph(f"{activity.title} ({activity.minutes} min)")
            document.add_heading("Key Points", level=2)
            for fact in session.key_points:
                document.add_paragraph(
                    f"{fact.text} [Evidence: {', '.join(fact.binding.evidence_ids)}]"
                )
            if session.homework:
                document.add_heading("Homework", level=2)
                for item in session.homework:
                    document.add_paragraph(item, style="List Bullet")
        document.save(str(output))
        return output
