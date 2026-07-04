from pathlib import Path

from docx import Document

from coursepilot.schemas.exam_schema import ExamBlueprintContent
from coursepilot.schemas.question_schema import QuestionItem


class ExamDocxExporter:
    def export_student_exam(
        self,
        blueprint: ExamBlueprintContent,
        questions: list[QuestionItem],
        output_path: str | Path,
    ) -> Path:
        doc = self._new_doc(blueprint, "Student Exam")
        for index, question in enumerate(questions, start=1):
            self._write_question(doc, index, question, include_answer=False, include_explanation=False)
            doc.add_paragraph("Answer: ______________________________")
        return self._save(doc, output_path)

    def export_teacher_answer(
        self,
        blueprint: ExamBlueprintContent,
        questions: list[QuestionItem],
        output_path: str | Path,
    ) -> Path:
        doc = self._new_doc(blueprint, "Teacher Answer Key")
        for index, question in enumerate(questions, start=1):
            self._write_question(doc, index, question, include_answer=True, include_explanation=False)
        return self._save(doc, output_path)

    def export_explanation(
        self,
        blueprint: ExamBlueprintContent,
        questions: list[QuestionItem],
        output_path: str | Path,
    ) -> Path:
        doc = self._new_doc(blueprint, "Detailed Explanation")
        for index, question in enumerate(questions, start=1):
            self._write_question(doc, index, question, include_answer=True, include_explanation=True)
            for ref in question.references:
                doc.add_paragraph(
                    f"Reference: {ref.source_type or 'source'} | {ref.chapter or '-'} | "
                    f"page={ref.page or '-'} | chunk={ref.chunk_id}"
                )
        return self._save(doc, output_path)

    def export_answer_sheet(
        self,
        blueprint: ExamBlueprintContent,
        questions: list[QuestionItem],
        output_path: str | Path,
    ) -> Path:
        doc = self._new_doc(blueprint, "Answer Sheet")
        for index, question in enumerate(questions, start=1):
            doc.add_paragraph(f"{index}. [{question.question_type}] ______________________________")
        return self._save(doc, output_path)

    def _new_doc(self, blueprint: ExamBlueprintContent, title: str) -> Document:
        doc = Document()
        doc.add_heading(f"{blueprint.course_name} - {title}", level=0)
        doc.add_paragraph(f"Chapter range: {blueprint.chapter_range}")
        doc.add_paragraph(f"Total score: {blueprint.total_score}")
        return doc

    def _write_question(
        self,
        doc: Document,
        index: int,
        question: QuestionItem,
        *,
        include_answer: bool,
        include_explanation: bool,
    ) -> None:
        doc.add_paragraph(
            f"{index}. ({question.question_type}, {question.score} pts) {question.question_text}"
        )
        if question.options:
            for key, value in question.options.items():
                doc.add_paragraph(f"{key}. {value}", style="List Bullet")
        if include_answer:
            doc.add_paragraph(f"Answer: {question.correct_answer}")
        if include_explanation:
            doc.add_paragraph(f"Explanation: {question.explanation}")

    def _save(self, doc: Document, output_path: str | Path) -> Path:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output_path)
        return output_path

