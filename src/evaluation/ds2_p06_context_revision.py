"""Build DS2 r7 after the course owner's all-record necessary-neighbor directive.

The r7 revision changes contextual dependency metadata only.  Gold text, source
units, BBoxes, semantic labels and stable Evidence IDs remain byte-for-byte
equivalent to r6.  Every record receives an explicit audit outcome: either a
smallest source-grounded neighbor set or a documented standalone decision.
"""

from __future__ import annotations

import argparse
import hashlib
import html
import json
from collections import Counter
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
from typing import Literal, TypedDict, cast

import fitz
from docx import Document
from pydantic import JsonValue

from courserag.evals.schemas import (
    DS2CandidateManifest,
    DS2EvidenceDataset,
    DS2ReviewDecisions,
    EvidenceNeighbor,
    EvidenceRecord,
)
from evaluation.contracts import CandidateRevisionArtifact, CandidateRevisionHistory
from evaluation.corpus_fixtures import sha256_file
from evaluation.datasets import record_digest
from evaluation.ds2_p06_data import (
    DATASET_ROOT,
    DOCX_RENDER_RUNS,
    PRIMARY_DOCX,
    _artifact,
    _build_render_index,
    _locate_rendered_text,
    _match_text,
    _poppler_page,
)
from evaluation.io import atomic_write_json, atomic_write_text

R6_PATH = DATASET_ROOT / "candidates/ds2/p06_evidence_r6.json"
R7_PATH = DATASET_ROOT / "candidates/ds2/p06_evidence_r7.json"
MANIFEST_PATH = DATASET_ROOT / "provenance/ds2_p06_candidate_manifest.json"
HISTORY_PATH = DATASET_ROOT / "provenance/ds2_p06_candidate_revision_history.json"
R6_DECISIONS_PATH = DATASET_ROOT / "reviews/ds2_p06_review_decisions_r6.json"
R7_DECISIONS_PATH = DATASET_ROOT / "reviews/ds2_p06_review_decisions_r7.json"
POLICY_PATH = DATASET_ROOT / "provenance/ds2_necessary_neighbor_policy_v1.json"
AUDIT_PATH = DATASET_ROOT / "provenance/ds2_p06_necessary_neighbor_audit_r7.json"
PHASE_REPORT_PATH = Path("docs/refactor/phase_reports/ED_PRE_P06_DS2_candidate_review.md")

R6_SHA256 = "f35292d55f6848e18f6e8067a8ca0e72cf0bedae49e92d60cd05cc1513c1b35c"
FORMULA_PROFILE_SHA256 = "8d84183ec617a8b816ee92cc91b7e2bae102e5dd59a7cdb5e72084de595c7bdd"

DependencyClass = Literal[
    "anaphora",
    "list_scope",
    "table_scope",
    "formula_symbols",
    "procedure_sequence",
    "causal_or_transition",
    "code_scope",
]

POLICY: JsonValue = {
    "schema_version": "courserag.ds2-necessary-neighbor-policy.v1",
    "policy_id": "ds2-necessary-neighbor-minimal-source-context-v1",
    "scope": "all 120 formal DS2 Evidence records",
    "decision_rule": (
        "Attach context only when the Evidence cannot be interpreted accurately on its own. "
        "Choose the nearest source passage that resolves the dependency; if a nearer passage "
        "is verbose, choose a smaller exact source passage that names the same antecedent or scope."
    ),
    "dependency_classes": {
        "anaphora": "Resolve external pronouns and deictic references such as 这、其中、上述、前述、它.",
        "list_scope": "Supply the exact list lead-in or category statement; do not add sibling items unless required.",
        "table_scope": "Supply the table caption; a row also receives the exact column header.",
        "formula_symbols": "Supply only the adjacent formula introduction or symbol definitions missing from the Evidence.",
        "procedure_sequence": "Resolve 首先、最后、而后、最终 with the minimal step or process statement.",
        "causal_or_transition": "Resolve 因为、与此同时、也 and similar links only when their other side is outside the Evidence.",
        "code_scope": "Supply the nearest variable definition and control-flow line needed to interpret the code fragment.",
    },
    "exclusions": [
        "A Section heading is not context merely because a record has semantic type other.",
        "Helpful background that does not resolve a dependency is excluded.",
        "P06 predictions, generated paraphrases and inferred course facts are forbidden.",
        "No neighbor may change the Evidence stable ID or become part of gold_text.",
    ],
    "course_owner_directive": (
        "Audit necessary neighbors for all records and add the nearest/minimum source-grounded "
        "context needed to remove pronoun, list and semantic dependencies; item-level re-review "
        "of this delegated context-only revision is waived, but exact batch Hash approval remains required."
    ),
}


def _sha256_text(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


@dataclass(frozen=True)
class ContextSpec:
    dependency_classes: tuple[DependencyClass, ...]
    rationale: str
    build: Callable[[EvidenceRecord, SourceContext], list[EvidenceNeighbor]]


class AuditEntry(TypedDict):
    evidence_id: str
    dependency_classes: list[DependencyClass]
    outcome: Literal["standalone", "neighbors_required"]
    rationale: str
    neighbor_text_sha256: list[str]


class SourceContext:
    def __init__(self, repository_root: Path) -> None:
        document = Document(str(repository_root / PRIMARY_DOCX))
        self.paragraphs = [paragraph.text for paragraph in document.paragraphs]
        self.tables = [
            [[cell.text for cell in row.cells] for row in table.rows] for table in document.tables
        ]
        del document
        self.repository_root = repository_root
        self.rendered_pdf_sha256 = (
            "4642e5c653b265b71bdbf9878218ca058188a9420b8b5930317cf4442c09fb95"
        )
        with fitz.open(repository_root / DOCX_RENDER_RUNS[0]) as rendered:
            self.render_index = _build_render_index(rendered)

    def _docx_page(self, text: str, preferred_page: int) -> int:
        located = _locate_rendered_text(
            text,
            self.render_index,
            rendered_pdf_sha256=self.rendered_pdf_sha256,
            preferred_page=preferred_page,
        )
        if located is None:
            # A few fixed-render lines contain equation/object spacing that is
            # absent from OOXML paragraph text.  Every such context is taken
            # from the same source paragraph as its Evidence, so the already
            # frozen Evidence physical page is the authoritative fallback.
            return preferred_page
        return located[0].page_number

    def docx_text(
        self,
        record: EvidenceRecord,
        paragraph_index: int,
        text: str,
        *,
        relation: Literal["parent_heading", "previous", "next"] = "previous",
    ) -> EvidenceNeighbor:
        paragraph = self.paragraphs[paragraph_index]
        if paragraph.count(text) != 1:
            raise ValueError(f"DOCX neighbor is not unique in paragraph {paragraph_index}: {text}")
        start = paragraph.index(text)
        page_number = self._docx_page(text, record.source_span.page_start or 1)
        return EvidenceNeighbor(
            relation=relation,
            source_span=record.source_span.model_copy(
                update={
                    "page_start": page_number,
                    "page_end": page_number,
                    "block_start": f"paragraph:{paragraph_index}",
                    "block_end": f"paragraph:{paragraph_index}",
                    "char_start": start,
                    "char_end": start + len(text),
                }
            ),
            text=text,
            text_sha256=_sha256_text(text),
        )

    def docx_structural(
        self,
        record: EvidenceRecord,
        source_unit_id: str,
        text: str,
        page_number: int,
        *,
        relation: Literal["parent_heading", "previous", "next"] = "previous",
    ) -> EvidenceNeighbor:
        return EvidenceNeighbor(
            relation=relation,
            source_span=record.source_span.model_copy(
                update={
                    "page_start": page_number,
                    "page_end": page_number,
                    "block_start": source_unit_id,
                    "block_end": source_unit_id,
                    "char_start": None,
                    "char_end": None,
                }
            ),
            text=text,
            text_sha256=_sha256_text(text),
        )

    def pdf_text(
        self,
        record: EvidenceRecord,
        page_number: int,
        first_block: int,
        last_block: int,
        text: str,
        *,
        relation: Literal["parent_heading", "previous", "next"] = "previous",
        visual_order: bool = False,
    ) -> EvidenceNeighbor:
        blocks = {
            block.block_number: block
            for block in _poppler_page(self.repository_root, page_number)[2]
        }
        selected = [blocks[index].text for index in range(first_block, last_block + 1)]
        if not visual_order and _match_text(text) not in _match_text("\n".join(selected)):
            raise ValueError(
                f"PDF neighbor is not grounded in p{page_number:03d} "
                f"b{first_block:04d}-b{last_block:04d}: {text}"
            )
        return EvidenceNeighbor(
            relation=relation,
            source_span=record.source_span.model_copy(
                update={
                    "page_start": page_number,
                    "page_end": page_number,
                    "block_start": f"poppler:p{page_number:03d}:b{first_block:04d}",
                    "block_end": f"poppler:p{page_number:03d}:b{last_block:04d}",
                    "char_start": None,
                    "char_end": None,
                }
            ),
            text=text,
            text_sha256=_sha256_text(text),
        )


def _paragraph(
    index: int,
    text: str,
    *,
    relation: Literal["parent_heading", "previous", "next"] = "previous",
) -> Callable[[EvidenceRecord, SourceContext], list[EvidenceNeighbor]]:
    return lambda record, source: [source.docx_text(record, index, text, relation=relation)]


def _paragraphs(
    items: list[tuple[int, str, Literal["parent_heading", "previous", "next"]]],
) -> Callable[[EvidenceRecord, SourceContext], list[EvidenceNeighbor]]:
    return lambda record, source: [
        source.docx_text(record, index, text, relation=relation) for index, text, relation in items
    ]


def _table_caption(
    paragraph_index: int,
    caption: str,
) -> Callable[[EvidenceRecord, SourceContext], list[EvidenceNeighbor]]:
    return _paragraph(paragraph_index, caption)


def _table_row_context(record: EvidenceRecord, source: SourceContext) -> list[EvidenceNeighbor]:
    caption = source.docx_text(record, 441, "表3.1 深度学习应用场景")
    page = record.source_span.page_start or caption.source_span.page_start or 1
    header = source.docx_structural(
        record,
        "table:1:header",
        "应用领域\t具体应用方向",
        page,
    )
    return [caption, header]


def _pdf(
    page: int,
    first_block: int,
    last_block: int,
    text: str,
    *,
    relation: Literal["parent_heading", "previous", "next"] = "previous",
    visual_order: bool = False,
) -> Callable[[EvidenceRecord, SourceContext], list[EvidenceNeighbor]]:
    return lambda record, source: [
        source.pdf_text(
            record,
            page,
            first_block,
            last_block,
            text,
            relation=relation,
            visual_order=visual_order,
        )
    ]


def _formula_distance(record: EvidenceRecord, source: SourceContext) -> list[EvidenceNeighbor]:
    return [
        source.docx_structural(
            record,
            "paragraph:163:rendered-equation",
            "1/||ω|| |ω·X_0+b|",
            record.source_span.page_start or 44,
            relation="next",
        )
    ]


def _alignment_method_context(
    record: EvidenceRecord, source: SourceContext
) -> list[EvidenceNeighbor]:
    items = [
        (1668, "1. 数据清理和预处理", 264),
        (1671, "2. 人类反馈", 265),
        (1674, "3. 对抗性训练", 265),
        (1677, "4. 多样化模型训练", 266),
        (1680, "5. 伦理审查和法规遵从", 266),
        (1683, "6. 持续监控和更新", 267),
    ]
    neighbors: list[EvidenceNeighbor] = []
    for paragraph_index, text, page_number in items:
        neighbor = source.docx_text(record, paragraph_index, text)
        neighbors.append(
            neighbor.model_copy(
                update={
                    "source_span": neighbor.source_span.model_copy(
                        update={"page_start": page_number, "page_end": page_number}
                    )
                }
            )
        )
    return neighbors


def _algorithm_context(record: EvidenceRecord, source: SourceContext) -> list[EvidenceNeighbor]:
    algorithm = next(
        item.gold_text
        for item in DS2EvidenceDataset.model_validate_json(
            R6_PATH.read_text(encoding="utf-8")
        ).evidence
        if item.source_units[0].source_unit_id == "table:0:rendered-visible-algorithm"
    )
    return [
        source.docx_structural(
            record,
            "table:0:rendered-visible-algorithm",
            algorithm,
            50,
        )
    ]


def _llama_table(record: EvidenceRecord, source: SourceContext) -> list[EvidenceNeighbor]:
    table_text = "\n".join("\t".join(row) for row in source.tables[7])
    return [
        source.docx_structural(
            record,
            "table:7",
            table_text,
            240,
            relation="next",
        )
    ]


def _formula_symbol_contexts(
    record: EvidenceRecord, source: SourceContext
) -> list[EvidenceNeighbor]:
    source_unit = record.source_units[0].source_unit_id
    if source_unit.startswith("paragraph:2420"):
        text = (
            "式中，σ(r(t))和c(r(t),d)代表沿摄像机射线观察方向d，r(t)的体积密度和颜色，"
            "dt代表射线在每个积分步长上的移动距离差。"
        )
        return [
            source.docx_structural(
                record,
                "paragraph:2419:rendered-inline-equations",
                text,
                341,
            )
        ]
    if source_unit.startswith("paragraph:2427"):
        return [
            source.docx_text(
                record,
                2426,
                source.paragraphs[2426],
            )
        ]
    if source_unit.startswith("paragraph:2444"):
        text = (
            "式中，T_cw代表相机的外参数矩阵，P代表相机空间到正则化空间的投影矩阵，"
            "w和h代表输出图像的宽和高，(c_x,c_y)代表图像中心点坐标。"
        )
        return [
            source.docx_structural(
                record,
                "paragraph:2446:rendered-inline-equations",
                text,
                346,
                relation="next",
            )
        ]
    if source_unit.startswith("paragraph:2448"):
        text = (
            "然而，三维高斯的透视投影不会产生二维高斯。因此，协方差矩阵Σ在像素空间的投影"
            "需通过在摄像机帧中t处的一阶泰勒展开近似获得，且该仿射变换由雅可比矩阵"
            "J∈R^{2×3}计算得出："
        )
        return [
            source.docx_structural(
                record,
                "paragraph:2447:rendered-inline-equations",
                text,
                347,
            )
        ]
    raise ValueError(f"formula context is not mapped: {source_unit}")


def _context_specs() -> dict[str, ContextSpec]:
    specs: dict[str, ContextSpec] = {}

    def add(
        record_id: str,
        classes: tuple[DependencyClass, ...],
        rationale: str,
        build: Callable[[EvidenceRecord, SourceContext], list[EvidenceNeighbor]],
    ) -> None:
        specs[record_id] = ContextSpec(classes, rationale, build)

    add(
        "gold-ev-35469c2bf6874d9057e05073805a8eb8",
        ("anaphora",),
        "“这一测试”需要最近的“图灵测试”先行词；仅保留首次命名该测试的一句。",
        _paragraph(
            20,
            "阿兰·图灵(Alan Mathison Turing)于1950年提出了图灵测试[2]，这是一种用于评估机器是否具有智能的测试方法。",
        ),
    )
    add(
        "gold-ev-a9911582c78e1c224211cd7b66f9ccbd",
        ("list_scope",),
        "单个能力条目需要精确的列表引导句，不需要其他兄弟条目。",
        _paragraph(20, "当然，为了实现计算机的类人行为智能，这类智能体需要具备以下能力："),
    )
    add(
        "gold-ev-8b932ed56da4f885f0b36b5b10ac5d0b",
        ("procedure_sequence",),
        "“首先”所属的实践任务由紧邻前句界定。",
        _paragraph(
            183, "该案例基于Python语言的Scikit-Learn机器学习库实现感知机模型，并完成二分类任务。"
        ),
    )
    add(
        "gold-ev-6ecdac141332a7cb4593a131ccfc4619",
        ("code_scope",),
        "“第一/二/三行代码”必须绑定被解释的三行代码，而不是章节标题。",
        _paragraphs(
            [
                (188, "model=Perceptron(fit_intercept=False,n_iter=20,shuffle=True)", "previous"),
                (189, "model.fit(x_train,y_train)", "previous"),
                (190, "acc=model.score(x_test,y_test)", "previous"),
            ]
        ),
    )
    add(
        "gold-ev-f333ceb9cf96f580ec2a714438e94d27",
        ("formula_symbols",),
        "冒号后的距离公式是该引导句不可分割的最小后邻接。",
        _formula_distance,
    )
    add(
        "gold-ev-06bba04373bb42548483d596702e7e36",
        ("anaphora", "procedure_sequence"),
        "“以上算法”直接指向紧邻的完整感知机算法。",
        _algorithm_context,
    )
    add(
        "gold-ev-899a0d03f4d2d6ef4b06b0d5d7c56532",
        ("table_scope",),
        "矩形网格已有表头，只需表题标明它是深度学习应用场景。",
        _table_caption(441, "表3.1 深度学习应用场景"),
    )
    for record_id in (
        "gold-ev-df677f4ff85b6c769ff0d151ecb5e4a9",
        "gold-ev-1b100e38a40ce3da44a2a2b530bef411",
        "gold-ev-537b4cddfec78d5d8f9a9da14e9d9654",
    ):
        add(
            record_id,
            ("table_scope",),
            "单独表行需要表题确定主题，并需要表头解释两列含义。",
            _table_row_context,
        )
    add(
        "gold-ev-520d6a9a5468ceca52daee6f7de7e998",
        ("anaphora",),
        "“其中”需要紧邻句先声明Transformer由编码器和解码器组成。",
        _paragraph(
            1056,
            "Transformer的整体结构如图3.18所示，主要包括编码器（Encoder）和解码器（Decoder）两部分。",
        ),
    )
    add(
        "gold-ev-035088153234134744f10d335a333cc0",
        ("table_scope",),
        "模型名称、年份、参数量网格需要表题界定为常见语言大模型参数量。",
        _table_caption(1535, "表5.1 常见语言大模型的参数量"),
    )
    add(
        "gold-ev-365b0dc23cff9bb089df5923e2212264",
        ("procedure_sequence",),
        "“最后”只需列表总领句即可确定它是涌现能力形成因素之一。",
        _paragraph(1538, "涌现能力的形成基于几个关键因素。"),
    )
    add(
        "gold-ev-c4302e898ef633d4ffb54fd4fc102e01",
        ("list_scope",),
        "第（5）项需要列表总领句；其他兄弟项不是理解该项所必需。",
        _paragraph(1539, "涌现能力的具体表现形式多样，主要包括以下几个方面："),
    )
    add(
        "gold-ev-5c3010622081949981b20043ffc9cdd4",
        ("causal_or_transition",),
        "“如GPT-4”是对紧邻训练耗时判断的例证。",
        _paragraph(
            1548, "即使在高性能计算资源的支持下，大模型的训练过程仍可能需要数周到数月的时间。"
        ),
    )
    add(
        "gold-ev-75d22d8bcfbf8c36e0066355d522c1c2",
        ("anaphora",),
        "“这些数据集”由同段紧邻前句的大规模训练数据集唯一指代。",
        _paragraph(
            1565, "大模型训练涉及使用大规模数据集来确保模型能够学习到尽可能多的语言模式和知识。"
        ),
    )
    add(
        "gold-ev-c381ce4313a28997c8fcb240b28fd058",
        ("table_scope", "anaphora"),
        "“如表5.2所示”必须携带紧随其后的LLaMA数据表。",
        _llama_table,
    )
    add(
        "gold-ev-bd7d355279cc06aac5841601572cfe04",
        ("anaphora",),
        "“这”直接承接同段上一句关于降维、特征选择和计算资源的判断。",
        _paragraph(
            1581,
            "同时，通过数据预处理，如通过降维和特征选择，可以有效减少数据的复杂度，从而降低模型训练和预测时所需的计算资源。",
        ),
    )
    add(
        "gold-ev-830849986bdfd10a7bd6767ec1a741fe",
        ("anaphora",),
        "“这一过程”由同段第一句对大模型对齐的定义解除指代。",
        _paragraph(
            1651,
            "大模型的对齐（Alignment of Large Models）是指确保人工智能模型的行为、输出和目标与人类的价值观、意图和期望相一致。",
        ),
    )
    add(
        "gold-ev-adecfc060732ab29684a055efdaacecc",
        ("anaphora", "list_scope"),
        "“这些详细的对齐方法”需要六个方法名称；保留方法标题，省略各方法说明。",
        _alignment_method_context,
    )
    add(
        "gold-ev-f622172906b893ec87c367db5ea298b3",
        ("table_scope",),
        "实验环境网格需要精确表题。",
        _table_caption(1696, "表5.3 实验环境"),
    )
    add(
        "gold-ev-03e5b59b045a42c014af9c4c1a02ce98",
        ("procedure_sequence",),
        "“首先配置”属于微调模型与初始模型效果比对步骤。",
        _paragraph(1725, "（4）在测试数据集上对微调后模型与初始模型做效果比对"),
    )
    add(
        "gold-ev-0e4d6445988271493ac03dfe6f5a8706",
        ("code_scope",),
        "单行代码需要最近的txt_base_list定义、索引循环和微调前分支标记。",
        _paragraphs(
            [
                (1748, "txt_base_list = []", "previous"),
                (1749, "for line in file:", "previous"),
                (
                    1750,
                    "txt_base_list.append(line.strip())  # strip()函数用于去除行末的换行符",
                    "previous",
                ),
                (1770, "for index in range(length):", "previous"),
                (1771, "# 微调前的 #", "previous"),
            ]
        ),
    )
    add(
        "gold-ev-33a7eeb84721f894d7224e97ba49e5f5",
        ("anaphora", "causal_or_transition"),
        "“这进一步证明”由同段紧邻的微调前后错误分类概率对比解除指代。",
        _paragraph(
            1817,
            "通过以上代码，我们可以绘制出柱状图，如图5.11所示，从图中可以清晰地看出，使用基础模型（即预训练模型）进行预测时，错误分类的概率在26%左右。然而经过微调后，使用微调模型进行预测时，错误分类的概率几乎为0。",
        ),
    )
    for record_id in (
        "gold-ev-6f70e7f62938bf78ba45c307de1adaef",
        "gold-ev-488132baeca3f0acc28854e2d6ab51b8",
        "gold-ev-e21068d14748c7ffc033f2ec37f88501",
        "gold-ev-b4635872d741e3a6c0e764547d376516",
    ):
        add(
            record_id,
            ("formula_symbols",),
            "公式仅补充紧邻的缺失符号定义或公式目的，不使用章节标题。",
            _formula_symbol_contexts,
        )
    add(
        "gold-ev-d103c316281ffd09d95fdc5c49d6ca90",
        ("table_scope",),
        "算法网格需要精确表题界定创新前沿分类。",
        _table_caption(2436, "表6.1 神经辐射场三维重建创新前沿研究"),
    )
    add(
        "gold-ev-ce0b58d4e9bed39607616060acc59261",
        ("procedure_sequence",),
        "“而后”依赖同段先完成边界框计算并进入瓦片缓存的步骤。",
        _paragraph(
            2457,
            "在渲染过程中，对于场景中的每一个高斯，3D高斯点染首先计算其边界框，这个边界框是根据二维投影协方差所确定的椭圆轴对齐边界，确保了高斯粒子在投影过程中的统计特性得以保留。当边界框与特定的瓦片区域相交时，相应的高斯粒子便被纳入该瓦片的缓存中，准备进行后续处理。",
        ),
    )
    add(
        "gold-ev-1c656f4df29556da71c7217626346570",
        ("causal_or_transition",),
        "“与此同时”承接同段前句对3D高斯点染既有里程碑地位的判断。",
        _paragraph(
            2462,
            "3D高斯点染三维重建技术以其创新的三维高斯表示和高效光栅化算法成为三维视觉领域的一个新里程碑。",
        ),
    )
    add(
        "gold-ev-8739e43b9385b01e5b9b8ca672c5bb19",
        ("table_scope",),
        "传感器优缺点网格需要精确表题。",
        _table_caption(2469, "表6.3多源融合SLAM常用传感器的优点和缺点"),
    )
    add(
        "gold-ev-cb0846fb32a5dd3d31485859a0290141",
        ("anaphora",),
        "“不同的传感器”由最近的SLAM常用LiDAR和相机陈述界定。",
        _paragraph(
            2466,
            "在SLAM系统中，最常用的传感器是激光雷达（Light Detection and Ranging，LiDAR）和相机。",
        ),
    )
    add(
        "gold-ev-9f350313b9d2181d3fb13d62b08713e1",
        ("anaphora",),
        "“它的”由紧邻且唯一的算法小标题R3-LIVE解除；该标题本身就是最小先行词。",
        _paragraph(2514, "（2）R3-LIVE[23]", relation="parent_heading"),
    )
    add(
        "gold-ev-d582a5e418494c5dd20d1ebd93eb6f76",
        ("anaphora",),
        "句首“所运用”缺失主语，需前句的智能可穿戴系统定义。",
        _paragraph(
            3393,
            "智能可穿戴系统基于实时检测系统的典型特征和便携穿戴式设计，以非侵入性或微创的方式持续监测个体生理特征，将高度自主性的检测和分析功能一体化集成，通过对日常穿戴物品的智能化设计实时监测人体状态。",
        ),
    )
    add(
        "gold-ev-9376f75284ba93d51d39a91efd40fe3c",
        ("table_scope",),
        "算法类型网格需要精确表题。",
        _table_caption(3400, "表9.2智能可穿戴系统中的人工智能算法类型"),
    )
    add(
        "gold-ev-37c2653936a801fb1ace1e57b6e94fd8",
        ("list_scope",),
        "第（1）项需要最近的三类应用总领句。",
        _paragraph(
            3401,
            "上述人工智能技术在智能可穿戴系统中的应用主要体现在数据挖掘、状态感知和个性化推荐上。",
        ),
    )
    add(
        "gold-ev-b5848cc84a30502e827293c903043ec2",
        ("table_scope",),
        "距离/目标像素网格需要精确表题。",
        _table_caption(3499, "表9.4 不同距离，不同种类下目标在像素平面所占的像素格大小"),
    )
    add(
        "gold-ev-8e7f5191ecc55a07bf9bd0c23da8a94c",
        ("anaphora",),
        "“前述技术应用”由更小的类别总领句准确概括，无需复制紧邻长段。",
        _paragraph(3495, "典型的侦察感知能力包括目标检测、行为识别以及毁伤检测等。"),
    )
    add(
        "gold-ev-98ed1e051c4b052625cc4314716bc71e",
        ("procedure_sequence",),
        "“最终”依赖同段紧邻的相对位姿和融合建图优化结果。",
        _paragraph(
            3537,
            "通过第3步的方位转换和局部约束，优化了无人车与无人机的相对位姿，最后再进行关键帧的全局捆集调整，经过相应优化后的地面机器人轨迹更加接近真实值，同时无人机与地面机器人融合建图也更加精准。",
        ),
    )
    add(
        "gold-ev-a63ff5485aa0ef4684620c1ae70da74f",
        ("table_scope",),
        "政策网格需要精确表题。",
        _table_caption(3674, "表9.5 近3年中国服务机器人行业政策汇总表[21]"),
    )
    add(
        "gold-ev-7aad9ec018e1c23c8450e4eb3eb8c8ef",
        ("causal_or_transition",),
        "“也”承接紧邻的劳动力短缺场景，构成服务机器人需求的另一项社会因素。",
        _paragraph(
            3677,
            "在此背景下，服务机器人能够提供互动类服务，进行配送、巡检、引导等工作，有效缓解劳动密集型产业里出现的劳动力缺失和人力成本上涨问题。",
        ),
    )
    add(
        "gold-ev-8678aa8ac3b7207a105135326cdd7109",
        ("anaphora",),
        "“软件组成”需要服务机器人系统架构这一主语和硬件/软件二分范围。",
        _paragraph(
            3680,
            "服务机器人的系统架构通常由硬件和软件两大部分组成，确保其能够高效、智能地执行各种服务任务[26][27]。",
        ),
    )
    add(
        "gold-ev-45b6cf06eb11c87041e80cb332e0fb08",
        ("list_scope",),
        "c）项由用户界面的三种形式总领句界定，不复制a）和b）条目。",
        _paragraph(
            3695,
            "用户界面作为服务机器人与外界交互的窗口，包括图形用户界面、语音交互界面和视觉交互界面。",
        ),
    )

    add(
        "gold-ev-203251897a3d2eef2cfa747cbbe59fc2",
        ("anaphora",),
        "“另一重要贡献”需要最近一项明斯基既有贡献作为先行内容。",
        _pdf(
            12,
            12,
            13,
            "1969年被授予图灵奖的是人工智能学者马文·明斯基。1975年，他首创框架理论(Frame Theory)。",
        ),
    )
    add(
        "gold-ev-b9857bdc1a027788ca9043a1001e9477",
        ("anaphora",),
        "“这段长达十余年的时间”需要最近的达特茅斯会议后第一次人工智能热潮作为时期锚点。",
        _pdf(12, 20, 20, "在1956年的达特茅斯会议之后，人工智能迎来了属于它的第一股热潮。"),
    )
    add(
        "gold-ev-4f40dc8586ca6829794799ceae4826cd",
        ("causal_or_transition",),
        "句首“因为”解释紧邻的“现有定义未界定机器智能”判断。",
        _pdf(14, 25, 27, "但如何来界定一台计算机(机器)是否具有智能，它们都没有提及。"),
    )
    add(
        "gold-ev-63b2a09988aa11ff3c6d483a0b162917",
        ("list_scope",),
        "分号后的2013年谷歌事件需要同句首项提供年份和时间线范围。",
        _pdf(
            18,
            10,
            12,
            "2013年，Facebook成立了人工智能实验室，探索深度学习领域，借此为Facebook用户提供更加智能化的产品体验；",
        ),
    )
    add(
        "gold-ev-d388ee889882727dcfcd0d06ad3e6df0",
        ("anaphora",),
        "“其中”只需紧邻的三座AI城市排名即可解除范围指代。",
        _pdf(19, 23, 24, "北京、杭州、上海成AI城市前三强。"),
    )
    add(
        "gold-ev-331fa0be83528a68d9fae5234494bbd4",
        ("table_scope",),
        "PDF表格需要其可见表号和表题。",
        _pdf(20, 17, 18, "表1-1\n部分职业的被淘汰概率"),
    )
    add(
        "gold-ev-e3e6a8faf68b6e239ae34c4cfdda2fb1",
        ("list_scope",),
        "“具有如下特征”必须携带紧随其后的三项特征。",
        _pdf(
            20,
            8,
            10,
            "①需要从业者具备较强的社交能力、协商能力及人际沟通能力。\n②需要从业者具备较强的同情心，并为他人提供真心实意的扶助。\n③创意性较强。",
            relation="next",
        ),
    )
    add(
        "gold-ev-ae3c8bbcde5fb29d3c00dbf07a25b002",
        ("list_scope",),
        "第（1）项需要最近的“通用人工智能与强人工智能的区别”总领句。",
        _pdf(21, 18, 18, "通用人工智能与强人工智能的区别如下："),
    )
    add(
        "gold-ev-038b4413d51c1b35b884caf999ec92e0",
        ("anaphora",),
        "“一些问题”由紧随其后的声音人格权侵权案例具体化。",
        _pdf(
            24,
            10,
            12,
            "2024年出现了全国首例“AI生成声音人格权侵权案”，这就提醒我们，在享受技术带来的便利时，也要重视声音权益保护和技术的规范应用。",
            relation="next",
            visual_order=True,
        ),
    )
    add(
        "gold-ev-e300c37f5e95ad19b9c0be5e4217db4a",
        ("anaphora", "procedure_sequence"),
        "摄像头与专注度模型步骤需要紧邻的课堂教学智能反馈系统作为主语。",
        _pdf(27, 11, 11, "课堂教学智能反馈系统可以分析学生的课堂专注度和学习状态。"),
    )
    add(
        "gold-ev-be751fca2dda9ea3000329029002714c",
        ("anaphora",),
        "生物识别趋势需要最近的支付身份识别场景。",
        _pdf(31, 17, 17, "用户身份识别是支付起点。"),
    )
    return specs


def _with_neighbors(record: EvidenceRecord, neighbors: list[EvidenceNeighbor]) -> EvidenceRecord:
    values = record.model_dump(mode="json")
    values.update(
        {
            "requires_parent": bool(neighbors),
            "necessary_neighbor_text": [neighbor.text for neighbor in neighbors],
            "necessary_neighbors": [neighbor.model_dump(mode="json") for neighbor in neighbors],
        }
    )
    return EvidenceRecord.model_validate(values)


def _review_pack(
    repository_root: Path,
    records: list[EvidenceRecord],
    audit_entries: list[AuditEntry],
    candidate_sha256: str,
) -> tuple[Path, str]:
    review_dir = repository_root / "storage_eval/ds2_p06_review" / candidate_sha256
    review_dir.mkdir(parents=True, exist_ok=True)
    audit_by_id = {str(entry["evidence_id"]): entry for entry in audit_entries}
    cards: list[str] = []
    for record in records:
        entry = audit_by_id[record.record_id]
        neighbors = (
            "<br>".join(
                html.escape(f"{neighbor.relation}: {neighbor.text}")
                for neighbor in record.necessary_neighbors
            )
            or "无（已判定为独立完整语义单元）"
        )
        classes = "、".join(str(value) for value in entry["dependency_classes"]) or "无"
        cards.append(
            f"""<article><h2>{record.record_id}</h2>
<p><b>Section：</b>{html.escape(record.source_span.section_path[-1])}；<b>依赖类：</b>{html.escape(classes)}；
<b>结论：</b>{html.escape(str(entry["outcome"]))}</p>
<h3>Gold 原文（未改）</h3><pre>{html.escape(record.gold_text)}</pre>
<h3>必要邻接</h3><p>{neighbors}</p>
<p><b>审计理由：</b>{html.escape(str(entry["rationale"]))}</p></article>"""
        )
    index = f"""<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><title>DS2 r7 必要邻接全量审计</title>
<style>body{{font-family:system-ui,sans-serif;max-width:1100px;margin:32px auto;line-height:1.6;background:#f4f6f8}}article{{background:white;border:1px solid #cbd5e1;border-radius:10px;padding:18px;margin:16px 0}}pre{{white-space:pre-wrap;background:#f8fafc;padding:12px}}code{{word-break:break-all}}</style></head><body>
<h1>DS2 r7 必要邻接全量审计</h1><p><b>Candidate SHA-256：</b><code>{candidate_sha256}</code></p>
<p>本页记录 120/120 的邻接判断。课程所有者已授权跳过本次上下文字段的逐条复核；本页仅作可追溯审计。正式提升仍需对完整 Candidate SHA-256 明确批准。</p>
{"".join(cards)}</body></html>"""
    atomic_write_text(review_dir / "index.html", index)
    return review_dir, sha256_file(review_dir / "index.html")


def build_r7(repository_root: Path) -> dict[str, object]:
    repository_root = repository_root.resolve()
    r6_path = repository_root / R6_PATH
    if sha256_file(r6_path) != R6_SHA256:
        raise ValueError("DS2 r6 Hash changed; refusing to apply the context-only revision")
    r6 = DS2EvidenceDataset.model_validate_json(r6_path.read_text(encoding="utf-8"))
    old_manifest = DS2CandidateManifest.model_validate_json(
        (repository_root / MANIFEST_PATH).read_text(encoding="utf-8")
    )
    valid_lineage = (
        old_manifest.candidate_revision == 6 and old_manifest.candidate_file_sha256 == R6_SHA256
    ) or (
        old_manifest.candidate_revision == 7
        and old_manifest.predecessor_candidate_file_sha256 == R6_SHA256
    )
    if not valid_lineage:
        raise ValueError("current DS2 manifest is not bound to immutable r6/r7 lineage")

    source = SourceContext(repository_root)
    specs = _context_specs()
    ids = {record.record_id for record in r6.evidence}
    if not set(specs).issubset(ids):
        raise ValueError(f"context spec references unknown records: {sorted(set(specs) - ids)}")

    records: list[EvidenceRecord] = []
    audit_entries: list[AuditEntry] = []
    for record in r6.evidence:
        spec = specs.get(record.record_id)
        outcome: Literal["standalone", "neighbors_required"]
        if spec is None:
            neighbors: list[EvidenceNeighbor] = []
            classes: tuple[DependencyClass, ...] = ()
            rationale = (
                "逐字检查后，该记录自身已给出主语、对象和完整命题；不存在外部代词、列表/表格范围、"
                "公式符号、代码、步骤或转折依赖，因此不附加帮助性背景。"
            )
            outcome = "standalone"
        else:
            neighbors = spec.build(record, source)
            if not neighbors:
                raise AssertionError(f"dependent record has no neighbor: {record.record_id}")
            classes = spec.dependency_classes
            rationale = spec.rationale
            outcome = "neighbors_required"
        revised = _with_neighbors(record, neighbors)
        records.append(revised)
        audit_entries.append(
            {
                "evidence_id": record.record_id,
                "dependency_classes": list(classes),
                "outcome": outcome,
                "rationale": rationale,
                "neighbor_text_sha256": [neighbor.text_sha256 for neighbor in neighbors],
            }
        )

    dataset = DS2EvidenceDataset(
        dataset_id=r6.dataset_id,
        dataset_version=r6.dataset_version,
        evidence=records,
    )
    r7_path = repository_root / R7_PATH
    atomic_write_json(r7_path, dataset.model_dump(mode="json"))
    candidate_sha256 = sha256_file(r7_path)

    decisions = DS2ReviewDecisions(
        candidate_file_sha256=R6_SHA256,
        reviewed_record_ids=[record.record_id for record in r6.evidence],
        completed_group_ids=["group-01", "group-02", "group-03", "group-04"],
        notes=(
            "The course owner reported the r6 review complete, then issued a global context-only "
            "directive: audit necessary neighbors for all 120 records, use nearest/minimum exact "
            "source context, and waive another item-level review. Exact r7 batch Hash approval "
            "remains mandatory."
        ),
    )
    decisions_path = repository_root / R6_DECISIONS_PATH
    atomic_write_json(decisions_path, decisions.model_dump(mode="json"))
    delegated_decisions = DS2ReviewDecisions(
        candidate_file_sha256=candidate_sha256,
        reviewed_record_ids=[record.record_id for record in records],
        completed_group_ids=["group-01", "group-02", "group-03", "group-04"],
        notes=(
            "All 120 r7 necessary-neighbor outcomes were completed under the course-owner's "
            "delegated nearest/minimum context directive. The owner explicitly waived another "
            "item-level review for this context-only revision. This decision artifact is not a "
            "batch approval; literal r7 Candidate SHA-256 approval remains required."
        ),
    )
    delegated_decisions_path = repository_root / R7_DECISIONS_PATH
    atomic_write_json(
        delegated_decisions_path,
        delegated_decisions.model_dump(mode="json"),
    )
    policy_path = repository_root / POLICY_PATH
    atomic_write_json(policy_path, POLICY)

    audit: JsonValue = {
        "schema_version": "courserag.ds2-necessary-neighbor-audit.v1",
        "dataset_id": "courserag-eval",
        "dataset_version": "v1",
        "predecessor_candidate_file_sha256": R6_SHA256,
        "candidate_file_sha256": candidate_sha256,
        "policy_sha256": sha256_file(policy_path),
        "record_count": 120,
        "standalone_record_count": sum(not record.necessary_neighbors for record in records),
        "neighbors_required_record_count": sum(
            bool(record.necessary_neighbors) for record in records
        ),
        "stable_ids_unchanged": [record.record_id for record in records]
        == [record.record_id for record in r6.evidence],
        "gold_content_hashes_unchanged": {
            record.record_id: record.content_sha256 for record in records
        }
        == {record.record_id: record.content_sha256 for record in r6.evidence},
        "course_owner_item_rereview_waived": True,
        "exact_batch_hash_approval_required": True,
        "entries": cast(list[JsonValue], audit_entries),
    }
    audit_path = repository_root / AUDIT_PATH
    atomic_write_json(audit_path, audit)

    review_dir, review_index_sha = _review_pack(
        repository_root,
        records,
        audit_entries,
        candidate_sha256,
    )
    record_hashes = {record.record_id: record_digest(record) for record in records}
    context_counts = Counter(
        dependency for entry in audit_entries for dependency in entry["dependency_classes"]
    )
    source_artifacts = [
        artifact
        for artifact in old_manifest.source_artifacts
        if artifact.path
        not in {
            "datasets/courserag_eval/v1/reviews/ds2_p06_review_decisions_r5.json",
            R6_DECISIONS_PATH.as_posix(),
            R7_DECISIONS_PATH.as_posix(),
            POLICY_PATH.as_posix(),
            AUDIT_PATH.as_posix(),
        }
    ]
    source_artifacts.extend(
        [
            _artifact(repository_root, R6_DECISIONS_PATH, "application/json"),
            _artifact(repository_root, R7_DECISIONS_PATH, "application/json"),
            _artifact(repository_root, POLICY_PATH, "application/json"),
            _artifact(repository_root, AUDIT_PATH, "application/json"),
        ]
    )
    manifest = DS2CandidateManifest(
        dataset_id=old_manifest.dataset_id,
        dataset_version=old_manifest.dataset_version,
        candidate_revision=7,
        candidate_relative_path=R7_PATH.as_posix(),
        candidate_file_sha256=candidate_sha256,
        candidate_record_sha256=record_hashes,
        source_artifacts=source_artifacts,
        upstream_approved_file_sha256=old_manifest.upstream_approved_file_sha256,
        semantic_source_count=2,
        record_counts={
            **old_manifest.record_counts,
            "necessary_neighbor_records": sum(
                bool(record.necessary_neighbors) for record in records
            ),
            "standalone_records": sum(not record.necessary_neighbors for record in records),
        },
        section_record_counts=old_manifest.section_record_counts,
        semantic_type_counts=old_manifest.semantic_type_counts,
        review_groups=old_manifest.review_groups,
        review_pack_relative_path=(review_dir / "index.html")
        .relative_to(repository_root)
        .as_posix(),
        review_pack_index_sha256=review_index_sha,
        review_asset_sha256=old_manifest.review_asset_sha256,
        renderer_profile=old_manifest.renderer_profile,
        canonical_rendered_pdf_sha256=old_manifest.canonical_rendered_pdf_sha256,
        stable_id_profile_sha256=old_manifest.stable_id_profile_sha256,
        normalization_profile_sha256=old_manifest.normalization_profile_sha256,
        predecessor_candidate_file_sha256=R6_SHA256,
        review_decisions_artifact=_artifact(repository_root, R7_DECISIONS_PATH, "application/json"),
        retained_reviewed_record_ids=[record.record_id for record in records],
        required_rereview_record_ids=[],
        duplicate_review_record_ids=old_manifest.duplicate_review_record_ids,
        constraints=[
            "Candidate only; exact course_owner approval of the full r7 file SHA-256 is required.",
            "All 120 records received a necessary-neighbor audit under the frozen minimal-context policy.",
            "The course owner waived another item-level review for this delegated context-only revision.",
            "Gold text, content Hashes, source units, BBoxes, semantic labels and stable Evidence IDs are unchanged from r6.",
            "Neighbors are exact source text; generic Section headings were removed unless the heading itself is the unique antecedent.",
            "Rendered formula neighbor text uses the approved formula-linearization profile and fixed LibreOffice render.",
            "Only two independent semantic sources exist; derived OCR fixtures remain provenance only.",
            "No P06 Evidence Builder, Chunker, Retriever, LLM Judge or system prediction defined Gold.",
            "Dev/Test remain empty and Test remains unlocked.",
        ],
    )
    manifest_path = repository_root / MANIFEST_PATH
    atomic_write_json(manifest_path, manifest.model_dump(mode="json"))

    history_path = repository_root / HISTORY_PATH
    history = CandidateRevisionHistory.model_validate_json(history_path.read_text(encoding="utf-8"))
    revisions = [
        revision.model_copy(
            update={
                "status": "superseded",
                "reason": (
                    "Course-owner review completed, then a global necessary-neighbor audit was "
                    "requested for every record; r7 supersedes r6 without changing Gold content."
                ),
            }
        )
        if revision.revision == 6
        else revision
        for revision in history.revisions
        if revision.revision != 7
    ]
    revisions.append(
        CandidateRevisionArtifact(
            revision=7,
            candidate_relative_path=R7_PATH.as_posix(),
            candidate_file_sha256=candidate_sha256,
            status="pending_course_owner_review",
            reason=(
                "All-record nearest/minimum necessary-neighbor audit complete; item-level re-review "
                "waived by course owner, exact batch Hash approval still pending."
            ),
        )
    )
    atomic_write_json(
        history_path,
        history.model_copy(update={"revisions": revisions}).model_dump(mode="json"),
    )

    governance_path = repository_root / DATASET_ROOT / "manifest.json"
    governance = json.loads(governance_path.read_text(encoding="utf-8"))
    governance["phase_input_status"]["p06"] = "candidate_r7_pending_course_owner_exact_hash"
    for component_key in ("gold_components", "gold_component_status"):
        components = governance.setdefault(component_key, {})
        if component_key == "gold_components" or "ds2" in components:
            components["ds2"] = "candidate_r7_pending_course_owner_exact_hash"
    atomic_write_json(governance_path, governance)

    return {
        "candidate_path": R7_PATH.as_posix(),
        "candidate_sha256": candidate_sha256,
        "manifest_sha256": sha256_file(manifest_path),
        "audit_path": AUDIT_PATH.as_posix(),
        "audit_sha256": sha256_file(audit_path),
        "review_pack": (review_dir / "index.html").relative_to(repository_root).as_posix(),
        "records": len(records),
        "neighbors_required": sum(bool(record.necessary_neighbors) for record in records),
        "standalone": sum(not record.necessary_neighbors for record in records),
        "dependency_class_counts": dict(sorted(context_counts.items())),
        "record_hashes_changed": sum(
            record_digest(before) != record_digest(after)
            for before, after in zip(r6.evidence, records, strict=True)
        ),
        "gold_content_hashes_changed": sum(
            before.content_sha256 != after.content_sha256
            for before, after in zip(r6.evidence, records, strict=True)
        ),
        "stable_ids_changed": sum(
            before.record_id != after.record_id
            for before, after in zip(r6.evidence, records, strict=True)
        ),
        "formula_profile_sha256": FORMULA_PROFILE_SHA256,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the all-record DS2 r7 context revision")
    parser.add_argument("--repository-root", type=Path, default=Path.cwd())
    args = parser.parse_args()
    print(json.dumps(build_r7(args.repository_root), ensure_ascii=False, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
