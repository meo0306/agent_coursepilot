"""Build the source-grounded P04 input work package without authoring DS1 Gold."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
from collections.abc import Mapping
from pathlib import Path
from uuid import NAMESPACE_URL, uuid5

import fitz
from docx import Document
from dotenv import load_dotenv
from sqlalchemy import select
from sqlalchemy.orm import Session

from courserag.evals.schemas import (
    CorpusFixtureManifest,
    DocxPaginationAnchorCandidate,
    DS0CorpusDataset,
    NativePDFPageCandidate,
    OCRRouteOnlyCandidate,
    P03EvalIdentityBinding,
    P03EvalIdentitySnapshot,
    P04InputWorkPackage,
    SectionCandidateSelection,
    TableCandidateSelection,
)
from courserag.jobs.artifacts import FileArtifactStore
from courserag.jobs.coordinator import DocumentInventoryStage
from courserag.jobs.stages import BuildStageRunner, StageContext, sha256_json
from courserag.persistence.base import utc_now
from courserag.persistence.models import (
    ArtifactRecord,
    BuildJobDocumentVersionRecord,
    BuildJobRecord,
    BuildStageRunRecord,
    DocumentVersionRecord,
    KnowledgeBaseRecord,
    SourceDocumentRecord,
)
from courserag.persistence.repositories import CourseRAGRepository
from evaluation.corpus_fixtures import (
    CLEAN_SOURCE_PAGES,
    COMPRESSED_SOURCE_PAGES,
    DOCX_PRIMARY,
    DOCX_STRESS,
    MIXED_RASTER_SOURCE_PAGES,
    PDF_MIXED,
    PDF_PRIMARY,
    PDF_SCAN_CLEAN,
    PDF_SCAN_COMPRESSED,
    sha256_file,
    sha256_text,
)
from evaluation.io import atomic_write_json

DATASET_ID = "courserag-eval"
DATASET_VERSION = "v1"
IDENTITY_SCHEME = "uuid5-url-courserag-eval-v1"
EVAL_DB_PORT = "55432"
EVAL_DB_NAME = "coursepilot_eval"

NATIVE_PDF_PAGES = [5, 10, 11, 12, 15, 18, 20, 21, 23, 24, 26, 27, 30, 31, 32]
PILOT_PDF_PAGES = {5, 10, 12, 20, 32}
PDF_PAGE_TAGS = {
    5: ["two_column_toc", "heading_dense"],
    10: ["mind_map", "figure", "qr_code", "header_footer"],
    12: ["figure", "body_text", "section_boundary", "header_footer"],
    20: ["body_text", "numbered_list", "table", "header_footer"],
    32: ["body_text", "qr_code", "section_boundary", "header_footer"],
}

DOCX_ANCHORS = [
    (19, "1.1.1", "1.1.1 类人行为智能：图灵测试法"),
    (144, "2.2", "2.2 感知机与多层感知机"),
    (1034, "3.5.3", "3.5.3 注意力机制和Transformer"),
    (1614, "5.3", "5.3 大模型的微调与对齐"),
    (3782, "9.7", "9.7 全面融入人类生活的大模型技术"),
]

PDF_SECTIONS = [
    ("1.1.1", "1.1.1 人工智能起源"),
    ("1.1.2", "1.1.2 人工智能的定义"),
    ("1.2.1", "1.2.1 人工智能发展历程"),
    ("1.2.2", "1.2.2 人工智能的发展趋势"),
    ("1.2.3", "1.2.3 通用人工智能的发展"),
    ("1.3.1", "1.3.1 AI+传媒"),
    ("1.3.4", "1.3.4 AI+教育"),
    ("1.4.1", "1.4.1 基本概念"),
]

DOCX_SECTIONS = [
    ("1.1.1", "1.1.1 类人行为智能：图灵测试法"),
    ("2.2", "2.2 感知机与多层感知机"),
    ("2.2.2", "2.2.2 模型与原理"),
    ("3.1.3", "3.1.3 深度学习应用场景"),
    ("3.5.3", "3.5.3 注意力机制和Transformer"),
    ("5.1.1", "5.1.1 大模型的定义与特征"),
    ("5.2.1", "5.2.1 数据准备与预处理"),
    ("5.3", "5.3 大模型的微调与对齐"),
    ("6.2.2", "6.2.2 神经辐射场三维重建"),
    ("6.2.3", "6.2.3 3D高斯点染三维重建"),
    ("9.1.2", "9.1.2 技术原理：可穿戴技术与人工智能"),
    ("9.7", "9.7 全面融入人类生活的大模型技术"),
]

DOCX_TABLES = [
    (0, "2.2.2 模型与原理"),
    (1, "3.1.3 深度学习应用场景"),
    (6, "5.1.1 大模型的定义与特征"),
    (8, "5.3.4 实践：大模型微调实现酒店住宿评价情感分类"),
    (10, "6.2.2 神经辐射场三维重建"),
    (12, "6.2.3 3D高斯点染三维重建"),
    (18, "9.1.2 技术原理：可穿戴技术与人工智能"),
    (20, "9.2 智能侦察"),
    (21, "9.5.1 智慧服务系统概述"),
]

PENDING_P04_FIELDS = [
    "reading_order",
    "regions_and_noise",
    "section_boundaries",
    "table_cells",
    "docx_physical_page_index",
    "docx_display_page_label",
    "docx_section_page_index",
    "renderer_manifest",
    "rendered_pdf_sha256",
    "alignment_confidence",
]


def deterministic_id(entity_type: str, *parts: str) -> str:
    identity = "/".join(("courserag-eval", DATASET_VERSION, entity_type, *parts))
    return str(uuid5(NAMESPACE_URL, identity))


def _canonical_sha256(value: object) -> str:
    content = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode(
        "utf-8"
    )
    return hashlib.sha256(content).hexdigest()


def _relative_file(repository_root: Path, relative_path: str) -> Path:
    root = repository_root.resolve()
    path = (root / relative_path).resolve()
    if not path.is_relative_to(root):
        raise ValueError("evaluation corpus path escapes the repository root")
    if not path.is_file():
        raise FileNotFoundError(path)
    return path


def load_frozen_inputs(
    repository_root: Path,
    *,
    approved_ds0_path: Path,
    fixture_manifest_path: Path,
) -> tuple[DS0CorpusDataset, CorpusFixtureManifest, dict[str, Path]]:
    approved = DS0CorpusDataset.model_validate_json(approved_ds0_path.read_text(encoding="utf-8"))
    fixture_manifest = CorpusFixtureManifest.model_validate_json(
        fixture_manifest_path.read_text(encoding="utf-8")
    )
    if len(approved.documents) != 6 or any(
        document.review_status != "approved" for document in approved.documents
    ):
        raise ValueError("Pre-P04 preparation requires all six Approved DS0 documents")
    paths: dict[str, Path] = {}
    for document in approved.documents:
        relative_path = document.repository_relative_path
        if relative_path is None:
            raise ValueError(f"DS0 document has no repository path: {document.document_id}")
        path = _relative_file(repository_root, relative_path)
        if sha256_file(path) != document.sha256:
            raise ValueError(f"DS0 file hash changed: {document.document_id}")
        paths[document.document_id] = path
    if {document.document_id for document in approved.documents} != {
        artifact.document_id for artifact in fixture_manifest.artifacts
    }:
        raise ValueError("Approved DS0 and fixture manifest document sets differ")
    return approved, fixture_manifest, paths


def _assert_fields(record: object, expected: Mapping[str, object], *, identity: str) -> None:
    for field_name, expected_value in expected.items():
        if getattr(record, field_name) != expected_value:
            raise ValueError(f"existing {identity} conflicts on {field_name}")


def _ensure_knowledge_base(
    session: Session, repository: CourseRAGRepository, course_id: str
) -> KnowledgeBaseRecord:
    record_id = deterministic_id("knowledge-base", course_id)
    expected = {"course_id": course_id, "name": course_id, "status": "empty"}
    record = session.get(KnowledgeBaseRecord, record_id)
    if record is None:
        record = KnowledgeBaseRecord(id=record_id, **expected)
        repository.add(record)
        repository.flush()
    else:
        _assert_fields(record, expected, identity=f"knowledge base {course_id}")
    return record


def _ensure_source_document(
    session: Session,
    repository: CourseRAGRepository,
    *,
    knowledge_base: KnowledgeBaseRecord,
    document_id: str,
    filename: str,
    document_type: str,
) -> SourceDocumentRecord:
    record_id = deterministic_id("source-document", document_id)
    expected = {
        "knowledge_base_id": knowledge_base.id,
        "legacy_document_id": document_id,
        "filename": filename,
        "document_type": document_type,
        "status": "registered",
    }
    record = session.get(SourceDocumentRecord, record_id)
    if record is None:
        record = SourceDocumentRecord(id=record_id, **expected)
        repository.add(record)
        repository.flush()
    else:
        _assert_fields(record, expected, identity=f"source document {document_id}")
    return record


def _ensure_document_version(
    session: Session,
    repository: CourseRAGRepository,
    *,
    source_document: SourceDocumentRecord,
    document: object,
    source_path: Path,
    retrieval_eligible: bool,
) -> DocumentVersionRecord:
    document_id = str(getattr(document, "document_id"))
    digest = str(getattr(document, "sha256"))
    record_id = deterministic_id("document-version", document_id, digest)
    source_metadata = {
        "dataset_id": DATASET_ID,
        "dataset_version": DATASET_VERSION,
        "dataset_document_id": document_id,
        "dataset_document_version": str(getattr(document, "document_version")),
        "document_role": str(getattr(document, "document_role")),
        "parent_document_id": getattr(document, "parent_document_id"),
        "mutually_exclusive_variant_group": str(
            getattr(document, "mutually_exclusive_variant_group")
        ),
        "retrieval_eligible": retrieval_eligible,
        "repository_relative_path": str(getattr(document, "repository_relative_path")),
    }
    expected = {
        "source_document_id": source_document.id,
        "version_number": 1,
        "content_sha256": digest,
        "object_uri": f"eval-corpus://{document_id}/{digest}",
        "mime_type": str(getattr(document, "mime_type")),
        "size_bytes": source_path.stat().st_size,
        "status": "registered",
        "source_metadata_json": source_metadata,
    }
    record = session.get(DocumentVersionRecord, record_id)
    if record is None:
        record = DocumentVersionRecord(id=record_id, **expected)
        repository.add(record)
        repository.flush()
    else:
        _assert_fields(record, expected, identity=f"document version {document_id}")
    if source_document.current_version_id not in {None, record.id}:
        raise ValueError(f"source document has a conflicting current version: {document_id}")
    source_document.current_version_id = record.id
    return record


def _ensure_build(
    session: Session,
    repository: CourseRAGRepository,
    artifact_store: FileArtifactStore,
    *,
    knowledge_base: KnowledgeBaseRecord,
    document_version: DocumentVersionRecord,
) -> tuple[BuildJobRecord, BuildStageRunRecord, ArtifactRecord]:
    request_payload = {
        "knowledge_base_id": knowledge_base.id,
        "document_version_ids": [document_version.id],
        "force": False,
        "pipeline": "pre_p04_eval_inventory",
        "pipeline_version": "1.0",
    }
    request_hash = sha256_json(request_payload)
    record_id = deterministic_id("build-job", document_version.id, request_hash)
    expected = {
        "knowledge_base_id": knowledge_base.id,
        "request_hash": request_hash,
        "build_mode": "full",
        "force_rebuild": False,
    }
    job = session.get(BuildJobRecord, record_id)
    if job is None:
        job = BuildJobRecord(id=record_id, status="queued", **expected)
        repository.add(job)
        repository.flush()
        repository.add(
            BuildJobDocumentVersionRecord(
                build_job_id=job.id, document_version_id=document_version.id
            )
        )
        repository.flush()
    else:
        _assert_fields(job, expected, identity=f"build job {record_id}")
        linked = repository.list_build_document_versions(job.id)
        if [item.id for item in linked] != [document_version.id]:
            raise ValueError("Pre-P04 Build must contain exactly one document version")

    if job.status != "succeeded":
        job.status = "running"
        job.started_at = job.started_at or utc_now()
        job.attempt_count += 1
        stage = DocumentInventoryStage((document_version,))
        context = StageContext(
            build_job_id=job.id,
            knowledge_base_id=knowledge_base.id,
            input_hashes=(document_version.content_sha256,),
            input_identities=(
                "|".join(
                    (document_version.id, document_version.object_uri, document_version.mime_type)
                ),
            ),
            config={"pipeline": "pre_p04_eval_inventory", "pipeline_version": "1.0"},
        )
        BuildStageRunner(repository, artifact_store).run(stage, context)
        job.status = "succeeded"
        job.completed_at = utc_now()
        repository.audit(
            "build.succeeded",
            "build_job",
            job.id,
            actor_id="pre_p04_eval_seeder",
            details={"scope": "p03_inventory"},
        )
        repository.flush()
    stage_run = session.scalar(
        select(BuildStageRunRecord)
        .where(
            BuildStageRunRecord.build_job_id == job.id,
            BuildStageRunRecord.stage_name == "legacy_document_inventory",
            BuildStageRunRecord.status.in_(("succeeded", "cached")),
        )
        .order_by(BuildStageRunRecord.attempt_number.desc())
        .limit(1)
    )
    if stage_run is None or stage_run.artifact_id is None:
        raise ValueError("succeeded Pre-P04 Build has no inventory Stage artifact")
    artifact = session.get(ArtifactRecord, stage_run.artifact_id)
    if artifact is None or not artifact_store.verify(artifact.uri, artifact.sha256):
        raise ValueError("Pre-P04 inventory Artifact is missing or corrupt")
    return job, stage_run, artifact


def seed_p03_eval_identities(
    session: Session,
    *,
    approved: DS0CorpusDataset,
    fixture_manifest: CorpusFixtureManifest,
    paths: dict[str, Path],
    artifact_root: Path,
) -> P03EvalIdentitySnapshot:
    repository = CourseRAGRepository(session)
    artifact_store = FileArtifactStore(artifact_root)
    fixtures = {artifact.document_id: artifact for artifact in fixture_manifest.artifacts}
    bindings: list[P03EvalIdentityBinding] = []
    for document in sorted(approved.documents, key=lambda item: item.document_id):
        fixture = fixtures[document.document_id]
        knowledge_base = _ensure_knowledge_base(session, repository, document.course_id or "")
        source_document = _ensure_source_document(
            session,
            repository,
            knowledge_base=knowledge_base,
            document_id=document.document_id,
            filename=document.filename,
            document_type="pdf" if document.mime_type == "application/pdf" else "docx",
        )
        version = _ensure_document_version(
            session,
            repository,
            source_document=source_document,
            document=document,
            source_path=paths[document.document_id],
            retrieval_eligible=fixture.retrieval_eligible,
        )
        job, stage_run, artifact = _ensure_build(
            session,
            repository,
            artifact_store,
            knowledge_base=knowledge_base,
            document_version=version,
        )
        bindings.append(
            P03EvalIdentityBinding(
                course_id=document.course_id or "",
                document_id=document.document_id,
                document_role=document.document_role,
                parent_document_id=document.parent_document_id,
                mutually_exclusive_variant_group=document.mutually_exclusive_variant_group or "",
                retrieval_eligible=fixture.retrieval_eligible,
                ds0_document_version=document.document_version,
                document_sha256=document.sha256,
                knowledge_base_id=knowledge_base.id,
                source_document_id=source_document.id,
                document_version_id=version.id,
                build_job_id=job.id,
                build_request_sha256=job.request_hash,
                stage_status=stage_run.status,
                stage_fingerprint_sha256=stage_run.fingerprint,
                artifact_uri=artifact.uri,
                artifact_sha256=artifact.sha256,
            )
        )
    snapshot = P03EvalIdentitySnapshot(
        dataset_id=DATASET_ID,
        dataset_version=DATASET_VERSION,
        bindings=bindings,
    )
    session.commit()
    return snapshot


def _normalize_lookup(text: str) -> str:
    return re.sub(r"\s+", "", text).replace("＋", "+").casefold()


def _exact_normalized_line(source_lines: list[str], expected: str, *, identity: str) -> str:
    expected_normalized = _normalize_lookup(expected)
    matches = [
        line.strip()
        for line in source_lines
        if _normalize_lookup(line.strip()) == expected_normalized
    ]
    if len(matches) != 1:
        raise ValueError(
            f"source must contain exactly one normalized line for {identity}: {expected}"
        )
    return matches[0]


def _identity_map(snapshot: P03EvalIdentitySnapshot) -> dict[str, P03EvalIdentityBinding]:
    return {binding.document_id: binding for binding in snapshot.bindings}


def _native_pdf_candidates(
    path: Path, binding: P03EvalIdentityBinding
) -> list[NativePDFPageCandidate]:
    candidates: list[NativePDFPageCandidate] = []
    paired_pages = set(CLEAN_SOURCE_PAGES + COMPRESSED_SOURCE_PAGES)
    with fitz.open(path) as document:
        if document.page_count != 32:
            raise ValueError("native PDF page tree no longer contains 32 pages")
        for page_number in NATIVE_PDF_PAGES:
            page_text = document[page_number - 1].get_text("text")
            if not page_text.strip():
                raise ValueError(f"native PDF candidate page has no text: {page_number}")
            tags = list(PDF_PAGE_TAGS.get(page_number, ["native_text_reference"]))
            if page_number in paired_pages:
                tags.append("paired_scan_parent")
            candidates.append(
                NativePDFPageCandidate(
                    selection_id=f"p04-native-pdf-page-{page_number:02d}",
                    document_id=PDF_PRIMARY,
                    document_version_id=binding.document_version_id,
                    page_number=page_number,
                    page_text_sha256=sha256_text(page_text),
                    pilot=page_number in PILOT_PDF_PAGES,
                    coverage_tags=tags,
                )
            )
    return candidates


def _docx_anchor_candidates(
    primary_path: Path,
    fixture_manifest: CorpusFixtureManifest,
    bindings: dict[str, P03EvalIdentityBinding],
) -> list[DocxPaginationAnchorCandidate]:
    document = Document(str(primary_path))
    fixture = next(
        artifact for artifact in fixture_manifest.artifacts if artifact.document_id == DOCX_STRESS
    )
    source_unit_map = {item.source_index: item for item in fixture.source_unit_map}
    candidates: list[DocxPaginationAnchorCandidate] = []
    for paragraph_index, section_anchor, expected_text in DOCX_ANCHORS:
        actual_text = document.paragraphs[paragraph_index].text.strip()
        if actual_text != expected_text:
            raise ValueError(f"DOCX anchor changed at paragraph:{paragraph_index}")
        digest = sha256_text(actual_text)
        source_index = f"paragraph:{paragraph_index}"
        mapping = source_unit_map.get(source_index)
        if mapping is None or mapping.text_sha256 != digest:
            raise ValueError(f"DOCX stress mapping changed for {source_index}")
        for document_id, mapped_source_unit in (
            (DOCX_PRIMARY, source_index),
            (DOCX_STRESS, mapping.target_index),
        ):
            candidates.append(
                DocxPaginationAnchorCandidate(
                    selection_id=(
                        f"p04-docx-anchor-{document_id.replace('_', '-')}-{section_anchor}"
                    ),
                    document_id=document_id,
                    document_version_id=bindings[document_id].document_version_id,
                    source_document_id=DOCX_PRIMARY,
                    source_paragraph_index=paragraph_index,
                    section_anchor=section_anchor,
                    source_text=actual_text,
                    source_text_sha256=digest,
                    mapped_source_unit=mapped_source_unit,
                )
            )
    return candidates


def _section_candidates(
    pdf_path: Path,
    docx_path: Path,
    bindings: dict[str, P03EvalIdentityBinding],
) -> list[SectionCandidateSelection]:
    with fitz.open(pdf_path) as document:
        pdf_lines = [line for page in document for line in page.get_text("text").splitlines()]
    candidates: list[SectionCandidateSelection] = []
    for anchor, title in PDF_SECTIONS:
        source_title = _exact_normalized_line(pdf_lines, title, identity=f"PDF section {anchor}")
        candidates.append(
            SectionCandidateSelection(
                selection_id=f"p04-section-pdf-{anchor}",
                document_id=PDF_PRIMARY,
                document_version_id=bindings[PDF_PRIMARY].document_version_id,
                section_anchor=anchor,
                source_title=source_title,
                source_title_sha256=sha256_text(source_title),
            )
        )
    document = Document(str(docx_path))
    paragraph_by_text: dict[str, list[int]] = {}
    for index, paragraph in enumerate(document.paragraphs):
        paragraph_by_text.setdefault(paragraph.text.strip(), []).append(index)
    for anchor, title in DOCX_SECTIONS:
        indexes = paragraph_by_text.get(title, [])
        if len(indexes) != 1:
            raise ValueError(f"DOCX section title is not unique: {title}")
        candidates.append(
            SectionCandidateSelection(
                selection_id=f"p04-section-docx-{anchor}",
                document_id=DOCX_PRIMARY,
                document_version_id=bindings[DOCX_PRIMARY].document_version_id,
                section_anchor=anchor,
                source_title=title,
                source_title_sha256=sha256_text(title),
                source_paragraph_index=indexes[0],
            )
        )
    return candidates


def _table_matrix(table: object) -> list[list[str]]:
    rows = getattr(table, "rows")
    return [[cell.text for cell in row.cells] for row in rows]


def _table_candidates(
    pdf_path: Path,
    docx_path: Path,
    bindings: dict[str, P03EvalIdentityBinding],
) -> list[TableCandidateSelection]:
    document = Document(str(docx_path))
    paragraph_lines = [paragraph.text.strip() for paragraph in document.paragraphs]
    candidates: list[TableCandidateSelection] = []
    for table_index, heading in DOCX_TABLES:
        source_heading = _exact_normalized_line(
            paragraph_lines, heading, identity=f"DOCX table {table_index} heading"
        )
        table = document.tables[table_index]
        matrix = _table_matrix(table)
        candidates.append(
            TableCandidateSelection(
                selection_id=f"p04-table-docx-{table_index:02d}",
                document_id=DOCX_PRIMARY,
                document_version_id=bindings[DOCX_PRIMARY].document_version_id,
                source_table_index=table_index,
                source_anchor_text=source_heading,
                source_content_sha256=_canonical_sha256(matrix),
                observed_row_count=len(matrix),
                observed_column_count=max(len(row) for row in matrix),
            )
        )
    with fitz.open(pdf_path) as pdf:
        page_text = pdf[19].get_text("text")
    pdf_anchor = "表1-1 部分职业的被淘汰概率"
    source_anchor = _exact_normalized_line(
        page_text.splitlines(), pdf_anchor, identity="PDF table on page 20"
    )
    candidates.append(
        TableCandidateSelection(
            selection_id="p04-table-pdf-page-20-table-1-1",
            document_id=PDF_PRIMARY,
            document_version_id=bindings[PDF_PRIMARY].document_version_id,
            source_page_number=20,
            source_anchor_text=source_anchor,
            source_content_sha256=sha256_text(page_text),
        )
    )
    return candidates


def _ocr_route_candidates(
    fixture_manifest: CorpusFixtureManifest,
    bindings: dict[str, P03EvalIdentityBinding],
) -> list[OCRRouteOnlyCandidate]:
    expected_sources = {
        PDF_SCAN_CLEAN: CLEAN_SOURCE_PAGES,
        PDF_SCAN_COMPRESSED: COMPRESSED_SOURCE_PAGES,
        PDF_MIXED: MIXED_RASTER_SOURCE_PAGES,
    }
    artifacts = {artifact.document_id: artifact for artifact in fixture_manifest.artifacts}
    candidates: list[OCRRouteOnlyCandidate] = []
    for document_id, source_pages in expected_sources.items():
        artifact = artifacts[document_id]
        page_by_source = {
            item.source_page_number: item
            for item in artifact.page_map
            if item.representation != "native"
        }
        if set(page_by_source) != set(source_pages):
            raise ValueError(f"OCR routing page map changed: {document_id}")
        for source_page in source_pages:
            page = page_by_source[source_page]
            candidates.append(
                OCRRouteOnlyCandidate(
                    selection_id=(
                        f"p04-ocr-route-{document_id.replace('_', '-')}-"
                        f"{page.output_page_number:02d}"
                    ),
                    document_id=document_id,
                    document_version_id=bindings[document_id].document_version_id,
                    document_page_number=page.output_page_number,
                    parent_document_id=PDF_PRIMARY,
                    source_page_number=source_page,
                    representation=page.representation,
                )
            )
    return candidates


def build_p04_input_work_package(
    *,
    snapshot: P03EvalIdentitySnapshot,
    snapshot_sha256: str,
    fixture_manifest: CorpusFixtureManifest,
    paths: dict[str, Path],
) -> P04InputWorkPackage:
    bindings = _identity_map(snapshot)
    return P04InputWorkPackage(
        record_id="p04-input-work-package-v1",
        review_status="candidate",
        candidate_source="deterministic_source_grounded_selection",
        dataset_id=DATASET_ID,
        dataset_version=DATASET_VERSION,
        p03_identity_snapshot_sha256=snapshot_sha256,
        identities=snapshot.bindings,
        native_pdf_pages=_native_pdf_candidates(paths[PDF_PRIMARY], bindings[PDF_PRIMARY]),
        docx_pagination_anchors=_docx_anchor_candidates(
            paths[DOCX_PRIMARY], fixture_manifest, bindings
        ),
        sections=_section_candidates(paths[PDF_PRIMARY], paths[DOCX_PRIMARY], bindings),
        tables=_table_candidates(paths[PDF_PRIMARY], paths[DOCX_PRIMARY], bindings),
        ocr_route_only_pages=_ocr_route_candidates(fixture_manifest, bindings),
        pending_p04_fields=PENDING_P04_FIELDS,
        constraints=[
            "This approval scope selects P04 annotation inputs; it does not approve DS1 Gold.",
            "DOCX pagination fields remain null until P04 freezes the Renderer Profile.",
            "OCR candidates are route-only; transcription and coordinates remain deferred to P05.",
            "Derived fixtures cannot be indexed with their primary parent document.",
            "The six evaluation documents represent only two independent semantic sources.",
            "No system Parser, Retriever, OCR, or QA prediction is stored in this work package.",
        ],
    )


def _update_candidate_manifest(manifest_path: Path) -> None:
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if manifest.get("gold_status") != "ds0_pilot_approved":
        raise ValueError("Pre-P04 candidate generation cannot change the current Gold boundary")
    status = manifest.setdefault("phase_input_status", {})
    current = status.get("p04")
    if current not in {None, "candidate_pending_course_owner"}:
        raise ValueError(f"unexpected P04 input status: {current!r}")
    status["p04"] = "candidate_pending_course_owner"
    atomic_write_json(manifest_path, manifest)


def prepare_candidate(
    session: Session,
    *,
    repository_root: Path,
    approved_ds0_path: Path,
    fixture_manifest_path: Path,
    identity_snapshot_path: Path,
    candidate_path: Path,
    dataset_manifest_path: Path,
    artifact_root: Path,
) -> dict[str, object]:
    approved, fixture_manifest, paths = load_frozen_inputs(
        repository_root,
        approved_ds0_path=approved_ds0_path,
        fixture_manifest_path=fixture_manifest_path,
    )
    snapshot = seed_p03_eval_identities(
        session,
        approved=approved,
        fixture_manifest=fixture_manifest,
        paths=paths,
        artifact_root=artifact_root,
    )
    atomic_write_json(identity_snapshot_path, snapshot.model_dump(mode="json"))
    snapshot_sha256 = sha256_file(identity_snapshot_path)
    candidate = build_p04_input_work_package(
        snapshot=snapshot,
        snapshot_sha256=snapshot_sha256,
        fixture_manifest=fixture_manifest,
        paths=paths,
    )
    atomic_write_json(candidate_path, candidate.model_dump(mode="json"))
    _update_candidate_manifest(dataset_manifest_path)
    return {
        "candidate_file_sha256": sha256_file(candidate_path),
        "identity_snapshot_sha256": snapshot_sha256,
        "knowledge_base_count": len({item.knowledge_base_id for item in snapshot.bindings}),
        "document_version_count": len(snapshot.bindings),
        "build_count": len(snapshot.bindings),
        "artifact_count": len({item.artifact_sha256 for item in snapshot.bindings}),
    }


def _configure_eval_environment(env_file: Path, port: str) -> None:
    load_dotenv(env_file, override=False)
    os.environ["POSTGRES_HOST"] = "127.0.0.1"
    os.environ["POSTGRES_PORT"] = port
    if os.environ.get("POSTGRES_DB") != EVAL_DB_NAME:
        raise ValueError(f"Pre-P04 preparation requires POSTGRES_DB={EVAL_DB_NAME}")


def _run_migrations(repository_root: Path) -> None:
    from alembic.config import Config

    from alembic import command

    command.upgrade(Config(str(repository_root / "alembic.ini")), "head")


def _database_session() -> Session:
    from coursepilot.db.session import CoursePilotSessionLocal, get_coursepilot_engine

    return CoursePilotSessionLocal(bind=get_coursepilot_engine())


def main() -> None:
    parser = argparse.ArgumentParser(description="Prepare the non-Gold P04 input work package.")
    parser.add_argument("command", choices=("migrate", "prepare"))
    parser.add_argument("--repository-root", type=Path, default=Path("."))
    parser.add_argument("--env-file", type=Path, default=Path(".env.eval"))
    parser.add_argument("--database-port", default=EVAL_DB_PORT)
    parser.add_argument(
        "--artifact-root", type=Path, default=Path("storage_eval/pre_p04/artifacts")
    )
    args = parser.parse_args()
    repository_root = args.repository_root.resolve()
    _configure_eval_environment(args.env_file, args.database_port)
    if args.command == "migrate":
        _run_migrations(repository_root)
        print(json.dumps({"alembic_revision": "0009_courserag_content"}, sort_keys=True))
        return
    dataset_root = repository_root / "datasets" / "courserag_eval" / "v1"
    session = _database_session()
    try:
        result = prepare_candidate(
            session,
            repository_root=repository_root,
            approved_ds0_path=dataset_root / "approved" / "ds0" / "pilot.json",
            fixture_manifest_path=dataset_root / "provenance" / "corpus_fixture_manifest.json",
            identity_snapshot_path=dataset_root / "provenance" / "p03_eval_identity_snapshot.json",
            candidate_path=dataset_root / "candidates" / "work_packages" / "p04_input.json",
            dataset_manifest_path=dataset_root / "manifest.json",
            artifact_root=args.artifact_root,
        )
    finally:
        session.close()
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))


if __name__ == "__main__":
    main()
