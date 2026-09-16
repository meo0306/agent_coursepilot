"""Create DS2 r6 from the course-owner's complete r5 review decisions.

The revision is deliberately a delta over immutable r5.  Only the 22 returned
records may change; the other 98 records and their review decisions are retained
byte-for-byte at record level.  Rendered formulas are transcribed from the fixed
LibreOffice PDF, never from P06 output.
"""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import shutil
from collections import Counter
from pathlib import Path
from typing import Literal

import fitz
from docx import Document
from pydantic import JsonValue

from courserag.evals.schemas import (
    DS2CandidateManifest,
    DS2EvidenceDataset,
    DS2ReviewDecisions,
    DS2ReviewGroup,
    EvidenceBBox,
    EvidenceNeighbor,
    EvidenceRecord,
    EvidenceSourceUnit,
)
from evaluation.contracts import (
    CandidateRevisionArtifact,
    CandidateRevisionHistory,
)
from evaluation.corpus_fixtures import sha256_file
from evaluation.datasets import record_digest
from evaluation.ds2_p06_data import (
    DATASET_ROOT,
    DOCX_RENDER_RUNS,
    PRIMARY_DOCX,
    PRIMARY_PDF,
    _artifact,
    _review_image,
    _stable_id,
)
from evaluation.io import atomic_write_json, atomic_write_text

R5_PATH = DATASET_ROOT / "candidates/ds2/p06_evidence_r5.json"
R6_PATH = DATASET_ROOT / "candidates/ds2/p06_evidence_r6.json"
MANIFEST_PATH = DATASET_ROOT / "provenance/ds2_p06_candidate_manifest.json"
HISTORY_PATH = DATASET_ROOT / "provenance/ds2_p06_candidate_revision_history.json"
DECISIONS_PATH = DATASET_ROOT / "reviews/ds2_p06_review_decisions_r5.json"
FORMULA_PROFILE_PATH = DATASET_ROOT / "provenance/ds2_formula_linearization_v1.json"
FEEDBACK_DIR = Path("storage_eval/ds2_p06_review/feedback")
PHASE_REPORT_PATH = Path("docs/refactor/phase_reports/ED_PRE_P06_DS2_candidate_review.md")

R5_SHA256 = "f4295f9783b131a8305713765907e169a83dc741d6c9af390b0b67b2f15b5bfc"

RETURNED_IDS = (
    "gold-ev-35469c2bf6874d9057e05073805a8eb8",
    "gold-ev-a9911582c78e1c224211cd7b66f9ccbd",
    "gold-ev-9adc0b9bee2ca3458e8bbfcf96ce0001",
    "gold-ev-0af37ab1ec6fa141fbd014388cd62950",
    "gold-ev-06bba04373bb42548483d596702e7e36",
    "gold-ev-df677f4ff85b6c769ff0d151ecb5e4a9",
    "gold-ev-1b100e38a40ce3da44a2a2b530bef411",
    "gold-ev-537b4cddfec78d5d8f9a9da14e9d9654",
    "gold-ev-f62c84f00da07a72487b548f43a65089",
    "gold-ev-520d6a9a5468ceca52daee6f7de7e998",
    "gold-ev-31d3879f19d532f859137666a7ad43ea",
    "gold-ev-c4302e898ef633d4ffb54fd4fc102e01",
    "gold-ev-75d22d8bcfbf8c36e0066355d522c1c2",
    "gold-ev-c381ce4313a28997c8fcb240b28fd058",
    "gold-ev-bd7d355279cc06aac5841601572cfe04",
    "gold-ev-830849986bdfd10a7bd6767ec1a741fe",
    "gold-ev-9439d8c1692cec2a57de2900a4effc62",
    "gold-ev-994f07594a8b689a4f0e38c1f73f3642",
    "gold-ev-6849d144b6969ff7d21bc7554362755f",
    "gold-ev-ad22bcc78de7f1db4688cc4e6aecfce5",
    "gold-ev-ce0b58d4e9bed39607616060acc59261",
    "gold-ev-9f350313b9d2181d3fb13d62b08713e1",
)

FORMULA_PROFILE: JsonValue = {
    "profile": "courserag.ds2-rendered-formula-linearization.v1",
    "authority": "fixed LibreOffice 7.4.7.2 rendered PDF visual objects",
    "rules": [
        "Preserve visible mathematical identifiers, operators, subscripts and superscripts.",
        "Use underscore-plus-braces for multi-character subscripts and caret for powers.",
        "Use square brackets with semicolon-delimited rows for rendered matrices/vectors.",
        "Do not repair or simplify a formula even when the source appears mathematically unusual.",
        "Bind every display formula to its rendered-PDF BBox and formula number.",
    ],
}


def _sha256_text(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _normalized_hash(value: str) -> str:
    import unicodedata

    normalized = " ".join(unicodedata.normalize("NFKC", value).split())
    return _sha256_text(normalized)


def _validated_record(record: EvidenceRecord, **updates: object) -> EvidenceRecord:
    values = record.model_dump(mode="json")
    values.update(updates)
    semantic = values["semantic_unit_type"]
    if semantic != "other":
        values["other_type_review_reason"] = None
    return EvidenceRecord.model_validate(values)


def _paragraph_neighbor(
    record: EvidenceRecord,
    paragraphs: list[str],
    *,
    paragraph_index: int,
    text: str,
    page_number: int,
    relation: Literal["parent_heading", "previous", "next"] = "previous",
) -> EvidenceNeighbor:
    paragraph = paragraphs[paragraph_index]
    if paragraph.count(text) != 1:
        raise ValueError(f"neighbor is not unique in paragraph {paragraph_index}: {text}")
    start = paragraph.index(text)
    span = record.source_span.model_copy(
        update={
            "page_start": page_number,
            "page_end": page_number,
            "block_start": f"paragraph:{paragraph_index}",
            "block_end": f"paragraph:{paragraph_index}",
            "char_start": start,
            "char_end": start + len(text),
        }
    )
    return EvidenceNeighbor(
        relation=relation,
        source_span=span,
        text=text,
        text_sha256=_sha256_text(text),
    )


def _structural_neighbor(
    record: EvidenceRecord,
    *,
    source_unit_id: str,
    text: str,
    page_number: int,
    relation: Literal["parent_heading", "previous", "next"],
    page_end: int | None = None,
) -> EvidenceNeighbor:
    return EvidenceNeighbor(
        relation=relation,
        source_span=record.source_span.model_copy(
            update={
                "page_start": page_number,
                "page_end": page_end or page_number,
                "block_start": source_unit_id,
                "block_end": source_unit_id,
                "char_start": None,
                "char_end": None,
            }
        ),
        text=text,
        text_sha256=_sha256_text(text),
    )


def _with_neighbors(
    record: EvidenceRecord,
    neighbors: list[EvidenceNeighbor],
    *,
    semantic: str | None = None,
    requires_parent: bool = True,
) -> EvidenceRecord:
    return _validated_record(
        record,
        semantic_unit_type=semantic or record.semantic_unit_type,
        requires_parent=requires_parent,
        necessary_neighbor_text=[neighbor.text for neighbor in neighbors],
        necessary_neighbors=[neighbor.model_dump(mode="json") for neighbor in neighbors],
    )


def _formula_bbox(
    record: EvidenceRecord,
    *,
    page_number: int,
    bbox: tuple[float, float, float, float],
    formula_id: str,
) -> EvidenceBBox:
    reference = record.bboxes[0]
    return EvidenceBBox(
        page_number=page_number,
        bbox=bbox,
        page_width=reference.page_width,
        page_height=reference.page_height,
        coordinate_space="rendered_pdf_points_top_left",
        rendered_pdf_sha256=reference.rendered_pdf_sha256,
        source_region_id=formula_id,
    )


def _replace_content(
    record: EvidenceRecord,
    *,
    pieces: list[tuple[str, str, int | None, int | None]],
    gold_text: str,
    assembly: str,
    semantic: str,
    bboxes: list[EvidenceBBox] | None = None,
    requires_parent: bool | None = None,
    neighbors: list[EvidenceNeighbor] | None = None,
) -> EvidenceRecord:
    units = [
        EvidenceSourceUnit(
            source_unit_id=source_unit_id,
            char_start=char_start,
            char_end=char_end,
            text_sha256=_sha256_text(text),
        )
        for source_unit_id, text, char_start, char_end in pieces
    ]
    content_hash = _sha256_text(gold_text)
    evidence_id = _stable_id(
        course_id=record.course_id or "",
        document_version=record.source_span.document_version,
        source_units=units,
        text_assembly=assembly,
        content_sha256=content_hash,
    )
    one_bounded_unit = len(units) == 1 and units[0].char_start is not None
    source_span = record.source_span.model_copy(
        update={
            "block_start": units[0].source_unit_id,
            "block_end": units[-1].source_unit_id,
            "char_start": units[0].char_start if one_bounded_unit else None,
            "char_end": units[0].char_end if one_bounded_unit else None,
        }
    )
    resolved_neighbors = list(record.necessary_neighbors) if neighbors is None else neighbors
    resolved_parent = record.requires_parent if requires_parent is None else requires_parent
    return _validated_record(
        record,
        record_id=evidence_id,
        evidence_id=evidence_id,
        candidate_source="source_grounded_fixed_renderer_transcription",
        source_span=source_span.model_dump(mode="json"),
        gold_text=gold_text,
        content_sha256=content_hash,
        normalized_content_sha256=_normalized_hash(gold_text),
        source_units=[unit.model_dump(mode="json") for unit in units],
        text_assembly=assembly,
        semantic_unit_type=semantic,
        requires_parent=resolved_parent,
        necessary_neighbor_text=[neighbor.text for neighbor in resolved_neighbors],
        necessary_neighbors=[neighbor.model_dump(mode="json") for neighbor in resolved_neighbors],
        bboxes=[box.model_dump(mode="json") for box in (bboxes or list(record.bboxes))],
    )


def _merge_feedback(
    repository_root: Path,
    candidate: DS2EvidenceDataset,
    review_groups: list[DS2ReviewGroup],
) -> tuple[DS2ReviewDecisions, dict[str, str]]:
    record_ids = {record.record_id for record in candidate.evidence}
    reviewed: list[str] = []
    returned: list[str] = []
    notes: dict[str, str] = {}
    feedback_hashes: dict[str, str] = {}
    expected_groups: dict[str, set[str]] = {
        group.group_id: set(group.record_ids) for group in review_groups
    }
    paths = sorted((repository_root / FEEDBACK_DIR).glob("ds2_group-*_review.json"))
    if len(paths) != 4:
        raise ValueError("exactly four DS2 r5 review-group files are required")
    for path in paths:
        values = json.loads(path.read_text(encoding="utf-8"))
        if values.get("candidate_file_sha256") != R5_SHA256:
            raise ValueError(f"feedback is not bound to DS2 r5: {path}")
        group_id = str(values.get("group_id"))
        expected_record_ids = expected_groups.get(group_id)
        if expected_record_ids is None:
            raise ValueError(f"unexpected review group: {group_id}")
        group_reviewed = list(values.get("reviewed_record_ids", []))
        group_returned = list(values.get("returned_record_ids", []))
        if set(group_reviewed) & set(group_returned):
            raise ValueError(f"conflicting review decision in {group_id}")
        if set(group_reviewed) | set(group_returned) != expected_record_ids:
            raise ValueError(f"review group is incomplete: {group_id}")
        reviewed.extend(group_reviewed)
        returned.extend(group_returned)
        notes.update(
            {str(key): str(value) for key, value in values.get("record_notes", {}).items()}
        )
        feedback_hashes[path.name] = sha256_file(path)
    if set(reviewed) | set(returned) != record_ids or len(reviewed) != 98:
        raise ValueError("merged r5 review must contain 98 accepted and 22 returned records")
    if (
        tuple(record.record_id for record in candidate.evidence if record.record_id in returned)
        != RETURNED_IDS
    ):
        raise ValueError("course-owner returned set differs from the audited r6 correction set")
    decisions = DS2ReviewDecisions(
        candidate_file_sha256=R5_SHA256,
        reviewed_record_ids=[
            record.record_id for record in candidate.evidence if record.record_id in reviewed
        ],
        returned_record_ids=list(RETURNED_IDS),
        completed_group_ids=["group-01", "group-02", "group-03", "group-04"],
        record_notes=notes,
        notes=(
            "Merged verbatim from four complete course_owner r5 review exports. "
            "The 98 accepted decisions are retained for r6; the 22 returned records require re-review."
        ),
    )
    return decisions, feedback_hashes


def _correct_records(
    candidate: DS2EvidenceDataset,
    paragraphs: list[str],
    tables: list[list[list[str]]],
) -> tuple[list[EvidenceRecord], dict[str, str]]:
    by_id = {record.record_id: record for record in candidate.evidence}
    corrected: dict[str, EvidenceRecord] = {}

    record = by_id[RETURNED_IDS[0]]
    context = (
        "阿兰·图灵(Alan Mathison Turing)于1950年提出了图灵测试[2]，这是一种用于评估机器是否具有智能的测试方法。"
        "该测试的基本思想是，如果一个机器能够在人类对话中表现得与人类一样，以至于无法区分机器和人类的回答，那么这台机器就可以被认为具有智能。"
    )
    corrected[record.record_id] = _with_neighbors(
        record,
        [_paragraph_neighbor(record, paragraphs, paragraph_index=20, text=context, page_number=6)],
    )

    record = by_id[RETURNED_IDS[1]]
    context = "当然，为了实现计算机的类人行为智能，这类智能体需要具备以下能力："
    corrected[record.record_id] = _with_neighbors(
        record,
        [_paragraph_neighbor(record, paragraphs, paragraph_index=20, text=context, page_number=6)],
        semantic="list",
    )

    record = by_id[RETURNED_IDS[2]]
    start = record.source_span.char_start
    if start is None:
        raise ValueError("returned paragraph 26 unit lost its source offset")
    text = paragraphs[26][start:]
    if (
        text
        != "要通过完整的图灵测试，机器人需要具备：（1）计算机视觉和语音识别来感知世界；（2）机器人技术来操控物体和移动。"
    ):
        raise ValueError("paragraph 26 full list changed")
    corrected[record.record_id] = _replace_content(
        record,
        pieces=[("paragraph:26", text, start, len(paragraphs[26]))],
        gold_text=text,
        assembly="single_source_unit",
        semantic="list",
    )

    algorithm = """输入：样本集合{(X_i,Y_i)}_{i=1}^s
输出：分类超平面的法向量ω和b
初始化：ω_1←0∈R^n，b_1←0
for t=1,2,…,T：
    随机挑选样本序号i=1,…,s
    if 分类错误Y_{i(t)}(ω_t^T X_{i(t)})≤0：
        更新ω_{t+1}←ω_t+ηY_{i(t)}X_{i(t)}
            b_{t+1}←b_t+ηY_{i(t)}
    else：
        ω_{t+1}←ω_t
返回ω_{T+1}"""
    record = by_id[RETURNED_IDS[3]]
    corrected_table = _replace_content(
        record,
        pieces=[("table:0:rendered-visible-algorithm", algorithm, None, None)],
        gold_text=algorithm,
        assembly="rendered_formula_linearization",
        semantic="procedure",
        requires_parent=False,
        neighbors=[],
    )
    corrected[record.record_id] = corrected_table

    record = by_id[RETURNED_IDS[4]]
    algorithm_neighbor = _structural_neighbor(
        record,
        source_unit_id="table:0:rendered-visible-algorithm",
        text=algorithm,
        page_number=50,
        relation="previous",
    )
    table_ref = next(
        ref for ref in corrected_table.upstream_record_refs if ref.record_type == "table"
    )
    corrected[record.record_id] = _validated_record(
        _with_neighbors(record, [algorithm_neighbor]),
        upstream_record_refs=[
            *[ref.model_dump(mode="json") for ref in record.upstream_record_refs],
            table_ref.model_dump(mode="json"),
        ],
    )

    table_header = "\t".join(tables[1][0])
    table_contexts = [
        _paragraph_neighbor(
            by_id[RETURNED_IDS[5]],
            paragraphs,
            paragraph_index=440,
            text="深度学习方法在众多领域中均有所应用，且显著优于其他方法。表3-1简单列写了部分应用场景。",
            page_number=92,
        ),
        _paragraph_neighbor(
            by_id[RETURNED_IDS[5]],
            paragraphs,
            paragraph_index=441,
            text="表3.1 深度学习应用场景",
            page_number=92,
        ),
        _structural_neighbor(
            by_id[RETURNED_IDS[5]],
            source_unit_id="table:1:header",
            text=table_header,
            page_number=92,
            relation="previous",
        ),
    ]
    for record_id in RETURNED_IDS[5:8]:
        record = by_id[record_id]
        neighbors = [
            neighbor.model_copy(
                update={
                    "source_span": neighbor.source_span.model_copy(
                        update={"section_path": record.source_span.section_path}
                    )
                }
            )
            for neighbor in table_contexts
        ]
        corrected[record.record_id] = _with_neighbors(record, neighbors)

    record = by_id[RETURNED_IDS[8]]
    text = (
        "注意力机制中，Q、K、V通常来自两个不同地方，其中Q来自解码器（Decoder）中的当前位置的隐藏状态或输出，"
        "K和V通常来自编码器（Encoder）中的所有位置的隐藏状态或输出，包含了输入序列的信息；"
    )
    corrected[record.record_id] = _replace_content(
        record,
        pieces=[("paragraph:1042:inline-equations", text, None, None)],
        gold_text=text,
        assembly="rendered_formula_linearization",
        semantic="principle",
    )

    record = by_id[RETURNED_IDS[9]]
    context = (
        "Transformer的整体结构如图3.18所示，主要包括编码器（Encoder）和解码器（Decoder）两部分。"
    )
    corrected[record.record_id] = _with_neighbors(
        record,
        [
            _paragraph_neighbor(
                record, paragraphs, paragraph_index=1056, text=context, page_number=148
            )
        ],
    )

    record = by_id[RETURNED_IDS[10]]
    text = "编码器-解码器注意力机制层的Q矩阵由解码器上一个子网络的输出计算得到，K矩阵和V矩阵由编码器输出编码矩阵计算得到。"
    corrected[record.record_id] = _replace_content(
        record,
        pieces=[("paragraph:1078:inline-equations", text, None, None)],
        gold_text=text,
        assembly="rendered_formula_linearization",
        semantic="principle",
    )

    record = by_id[RETURNED_IDS[11]]
    lead_in = "涌现能力的具体表现形式多样，主要包括以下几个方面："
    previous_item = paragraphs[1543]
    corrected[record.record_id] = _with_neighbors(
        record,
        [
            _paragraph_neighbor(
                record, paragraphs, paragraph_index=1539, text=lead_in, page_number=232
            ),
            _paragraph_neighbor(
                record, paragraphs, paragraph_index=1543, text=previous_item, page_number=233
            ),
        ],
    )

    record = by_id[RETURNED_IDS[12]]
    context = "大模型训练涉及使用大规模数据集来确保模型能够学习到尽可能多的语言模式和知识。"
    corrected[record.record_id] = _with_neighbors(
        record,
        [
            _paragraph_neighbor(
                record, paragraphs, paragraph_index=1565, text=context, page_number=239
            )
        ],
    )

    record = by_id[RETURNED_IDS[13]]
    table_7_text = "\n".join("\t".join(row) for row in tables[7])
    corrected[record.record_id] = _with_neighbors(
        record,
        [
            _structural_neighbor(
                record,
                source_unit_id="table:7",
                text=table_7_text,
                page_number=240,
                relation="next",
            )
        ],
    )

    record = by_id[RETURNED_IDS[14]]
    context = "同时，通过数据预处理，如通过降维和特征选择，可以有效减少数据的复杂度，从而降低模型训练和预测时所需的计算资源。"
    corrected[record.record_id] = _with_neighbors(
        record,
        [
            _paragraph_neighbor(
                record, paragraphs, paragraph_index=1581, text=context, page_number=242
            )
        ],
    )

    record = by_id[RETURNED_IDS[15]]
    context = paragraphs[1650]
    corrected[record.record_id] = _with_neighbors(
        record,
        [
            _paragraph_neighbor(
                record,
                paragraphs,
                paragraph_index=1650,
                text=context,
                page_number=260,
                relation="parent_heading",
            )
        ],
    )

    record = by_id[RETURNED_IDS[16]]
    prose = (
        "特别的，T(t)为累积透射率，表示射线从采样近平面t_n到采样远平面t_f期间射线未被阻挡的概率："
    )
    formula = "T(t)=exp(-∫_{t_n}^{t}σ(r(s))ds)"
    corrected[record.record_id] = _replace_content(
        record,
        pieces=[
            ("paragraph:2420:inline-equations", prose, None, None),
            ("paragraph:2421:equation:14", formula, None, None),
        ],
        gold_text=f"{prose}\n{formula}",
        assembly="rendered_formula_linearization",
        semantic="formula",
        bboxes=[
            *record.bboxes,
            _formula_bbox(
                record,
                page_number=341,
                bbox=(233.5, 441.5508, 362.6, 467.7508),
                formula_id="rendered-equation:14",
            ),
        ],
    )

    record = by_id[RETURNED_IDS[17]]
    formula = "L=∑_{r∈R}||Ĉ(r)-C_{gt}(r)||_2^2"
    prose = "式中，C_{gt}(r)代表与r相关的训练图像像素的真实颜色，R代表与待合成图像相关的射线集合。"
    corrected[record.record_id] = _replace_content(
        record,
        pieces=[
            ("paragraph:2427:equation:16", formula, None, None),
            ("paragraph:2428:inline-equations", prose, None, None),
        ],
        gold_text=f"{formula}\n{prose}",
        assembly="rendered_formula_linearization",
        semantic="formula",
        bboxes=[
            _formula_bbox(
                record,
                page_number=342,
                bbox=(242.9, 416.2508, 353.2, 446.1508),
                formula_id="rendered-equation:16",
            ),
            *record.bboxes,
        ],
    )

    record = by_id[RETURNED_IDS[18]]
    prose = (
        "在3D高斯点染的高斯投影过程中，三维空间中的高斯均值的μ通过标准的透视投影方法被映射到像素空间，"
        "这需要首先将μ转换为摄像机坐标系中的齐次坐标t∈R^4、归一化设备坐标系中的归一化坐标t̃∈R^4，"
        "并最终转换到像素坐标系中的像素坐标μ′∈R^2。"
    )
    formula = "t=T_cw[μ;1]，t̃=Pt，μ′=[(1/2)(w·t̃_x/t̃_w+1)+c_x；(1/2)(h·t̃_y/t̃_w+1)+c_y]"
    corrected[record.record_id] = _replace_content(
        record,
        pieces=[
            ("paragraph:2444:inline-equations", prose, None, None),
            ("paragraph:2445:equation:18", formula, None, None),
        ],
        gold_text=f"{prose}\n{formula}",
        assembly="rendered_formula_linearization",
        semantic="formula",
        bboxes=[
            *record.bboxes,
            _formula_bbox(
                record,
                page_number=346,
                bbox=(188.2, 526.7508, 407.9, 602.5508),
                formula_id="rendered-equation:18",
            ),
        ],
        requires_parent=True,
    )

    record = by_id[RETURNED_IDS[19]]
    formula_19 = "J=[[f_x/t_z,0,-f_x·t_x/t_z^2];[0,f_y/t_z,-f_y·t_y/t_z^2]]"
    prose = (
        "式中，f_x和f_y代表摄像机焦距。根据相机坐标系到世界坐标系的视角转换矩阵W，"
        "转换后的二维协方差矩阵Σ′∈R^{2×2}也随之确定："
    )
    formula_20 = "Σ′=JWΣW^T J^T"
    corrected[record.record_id] = _replace_content(
        record,
        pieces=[
            ("paragraph:2448:equation:19", formula_19, None, None),
            ("paragraph:2449:inline-equations", prose, None, None),
            ("paragraph:2450:equation:20", formula_20, None, None),
        ],
        gold_text=f"{formula_19}\n{prose}\n{formula_20}",
        assembly="rendered_formula_linearization",
        semantic="formula",
        bboxes=[
            _formula_bbox(
                record,
                page_number=347,
                bbox=(241.0, 212.9508, 355.1, 288.7508),
                formula_id="rendered-equation:19",
            ),
            *record.bboxes,
            _formula_bbox(
                record,
                page_number=347,
                bbox=(260.2, 422.8508, 335.9, 437.8508),
                formula_id="rendered-equation:20",
            ),
        ],
    )

    record = by_id[RETURNED_IDS[20]]
    corrected[record.record_id] = _validated_record(record, semantic_unit_type="procedure")

    record = by_id[RETURNED_IDS[21]]
    context = paragraphs[2514]
    corrected[record.record_id] = _with_neighbors(
        record,
        [
            _paragraph_neighbor(
                record,
                paragraphs,
                paragraph_index=2514,
                text=context,
                page_number=360,
                relation="parent_heading",
            )
        ],
    )

    if set(corrected) != set(RETURNED_IDS):
        raise AssertionError("r6 correction map does not cover the exact returned set")
    old_to_new = {old_id: corrected[old_id].record_id for old_id in RETURNED_IDS}
    records = [corrected.get(record.record_id, record) for record in candidate.evidence]
    return records, old_to_new


def _record_card(record: EvidenceRecord, predecessor_id: str, note: str) -> str:
    units = "<br>".join(
        html.escape(
            f"{unit.source_unit_id} chars={unit.char_start}:{unit.char_end} sha={unit.text_sha256}"
        )
        for unit in record.source_units
    )
    neighbors = (
        "<br>".join(
            html.escape(f"{neighbor.relation}: {neighbor.text}")
            for neighbor in record.necessary_neighbors
        )
        or "无"
    )
    payload = html.escape(
        json.dumps(record.model_dump(mode="json"), ensure_ascii=False, indent=2, sort_keys=True)
    )
    return f"""
<article id="{record.record_id}">
  <h2>{record.record_id}</h2>
  <p><b>r5 记录：</b>{predecessor_id}<br><b>退回意见：</b>{html.escape(note)}</p>
  <p><b>课程 / Section：</b>{html.escape(record.course_id or "")} / {html.escape(".".join(record.source_span.section_path))}<br>
  <b>来源 / 语义：</b>{record.source_type} / {record.semantic_unit_type}<br>
  <b>Span：</b>{html.escape(str(record.source_span.block_start))} → {html.escape(str(record.source_span.block_end))}, pages {record.source_span.page_start}-{record.source_span.page_end}<br>
  <b>组装：</b>{record.text_assembly}；<b>需要父上下文：</b>{str(record.requires_parent).lower()}</p>
  <img src="assets/{record.record_id}.png" alt="{record.record_id} BBox overlay">
  <h3>修订后的 Gold 原文</h3><pre>{html.escape(record.gold_text)}</pre>
  <p><b>Source Units</b><br>{units}</p>
  <p><b>必要邻接</b><br>{neighbors}</p>
  <details><summary>完整 Candidate JSON</summary><pre>{payload}</pre></details>
  <label><input type="radio" name="{record.record_id}" value="approve">通过</label>
  <label><input type="radio" name="{record.record_id}" value="return">退回</label>
  <input class="note" data-id="{record.record_id}" placeholder="复核备注">
</article>"""


def _delta_group_html(
    *,
    candidate_sha256: str,
    group: DS2ReviewGroup,
    records: list[EvidenceRecord],
    predecessor_by_new: dict[str, str],
    notes: dict[str, str],
) -> str:
    cards = "\n".join(
        _record_card(
            record,
            predecessor_by_new[record.record_id],
            notes[predecessor_by_new[record.record_id]],
        )
        for record in records
    )
    record_ids = [record.record_id for record in records]
    return f"""<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>DS2 r6 {group.group_id}</title><style>
body{{font-family:system-ui,sans-serif;margin:0;background:#f4f6f8;color:#17202a}}header,main{{max-width:1120px;margin:auto;padding:20px}}
header{{background:#17202a;color:white;max-width:none}}header>div{{max-width:1120px;margin:auto}}article{{background:white;border:1px solid #cbd5e1;border-radius:10px;padding:18px;margin:18px 0}}
img{{max-width:100%;border:1px solid #94a3b8}}pre{{white-space:pre-wrap;word-break:break-word;background:#f8fafc;padding:12px;overflow:auto}}.note{{width:60%;margin-left:12px}}a{{color:#93c5fd}}
</style></head><body><header><div><h1>{html.escape(group.title)} — r6 退回项复核</h1>
<p>Candidate SHA-256: <code>{candidate_sha256}</code>；本页只列出本组需复核的 {len(records)} 条，其余记录沿用 r5 已通过决定。</p><a href="index.html">返回总览</a></div></header>
<main>{cards}<button id="export">导出本组 r6 复核 JSON</button></main>
<script>document.getElementById('export').onclick=()=>{{const reviewed=[],returned=[],notes={{}};for(const id of {json.dumps(record_ids, ensure_ascii=False)}){{const value=document.querySelector(`input[name="${{id}}"]:checked`)?.value;if(value==='approve')reviewed.push(id);if(value==='return')returned.push(id);const note=document.querySelector(`.note[data-id="${{id}}"]`)?.value||'';if(note)notes[id]=note;}}const data={{schema_version:'courserag.ds2-review-group.v1',candidate_file_sha256:'{candidate_sha256}',group_id:'{group.group_id}',reviewed_record_ids:reviewed,returned_record_ids:returned,record_notes:notes}};const a=document.createElement('a');a.href=URL.createObjectURL(new Blob([JSON.stringify(data,null,2)],{{type:'application/json'}}));a.download='ds2_r6_{group.group_id}_review.json';a.click();}};</script></body></html>"""


def _review_pack(
    repository_root: Path,
    records: list[EvidenceRecord],
    groups: list[DS2ReviewGroup],
    candidate_sha256: str,
    old_to_new: dict[str, str],
    decisions: DS2ReviewDecisions,
) -> tuple[Path, str, dict[str, str]]:
    review_dir = repository_root / "storage_eval/ds2_p06_review" / candidate_sha256
    assets = review_dir / "assets"
    assets.mkdir(parents=True, exist_ok=True)
    old_assets = repository_root / "storage_eval/ds2_p06_review" / R5_SHA256 / "assets"
    predecessor_by_new = {new: old for old, new in old_to_new.items()}
    required = set(predecessor_by_new)
    with (
        fitz.open(repository_root / PRIMARY_PDF) as pdf_document,
        fitz.open(repository_root / DOCX_RENDER_RUNS[0]) as docx_rendered,
    ):
        for record in records:
            output = assets / f"{record.record_id}.png"
            if record.record_id not in required:
                source = old_assets / f"{record.record_id}.png"
                if not source.is_file():
                    raise FileNotFoundError(f"retained r5 review asset is missing: {source}")
                shutil.copyfile(source, output)
            else:
                _review_image(
                    record,
                    pdf_document=pdf_document,
                    docx_rendered=docx_rendered,
                    output=output,
                )
    by_id = {record.record_id: record for record in records}
    for group in groups:
        delta = [by_id[record_id] for record_id in group.record_ids if record_id in required]
        atomic_write_text(
            review_dir / f"{group.group_id}.html",
            _delta_group_html(
                candidate_sha256=candidate_sha256,
                group=group,
                records=delta,
                predecessor_by_new=predecessor_by_new,
                notes=decisions.record_notes,
            ),
        )
    counts = {
        group.group_id: sum(record_id in required for record_id in group.record_ids)
        for group in groups
    }
    links = "\n".join(
        f'<li><a href="{group.group_id}.html">{html.escape(group.title)}</a>（需复核 {counts[group.group_id]} 条）</li>'
        for group in groups
    )
    index = f"""<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><title>DS2 r6 差量复核</title>
<style>body{{font-family:system-ui,sans-serif;max-width:900px;margin:40px auto;line-height:1.65}}code{{word-break:break-all}}li{{margin:14px 0}}</style></head><body>
<h1>P06 前正式 DS2 Candidate r6 差量复核</h1><p><b>Candidate SHA-256：</b><code>{candidate_sha256}</code></p>
<p>r5 的 98 条“通过”决定已保留；仅复核 22 条修订记录。第 3、4 组无退回项，不需要重复审核。</p>
<ol>{links}</ol>
<p>重点检查：公式/算法线性化是否与红框可见内容一致，新增必要邻接是否刚好消除代词或列表依赖，语义类型是否正确。复核完成后仍需对上面的完整 SHA-256 明确批准，系统才会生成 Approved DS2。</p>
</body></html>"""
    atomic_write_text(review_dir / "index.html", index)
    hashes = {
        record.record_id: sha256_file(assets / f"{record.record_id}.png") for record in records
    }
    return review_dir, sha256_file(review_dir / "index.html"), hashes


def _update_governance(dataset_root: Path) -> None:
    path = dataset_root / "manifest.json"
    values = json.loads(path.read_text(encoding="utf-8"))
    values.setdefault("gold_components", {})["ds2"] = "candidate_r6_pending_course_owner"
    values.setdefault("phase_input_status", {})["p06"] = "candidate_r6_pending_course_owner"
    atomic_write_json(path, values)


def build_r6(repository_root: Path) -> dict[str, object]:
    repository_root = repository_root.resolve()
    r5_path = repository_root / R5_PATH
    if sha256_file(r5_path) != R5_SHA256:
        raise ValueError("immutable DS2 r5 Hash differs from the reviewed Candidate")
    candidate = DS2EvidenceDataset.model_validate_json(r5_path.read_text(encoding="utf-8"))
    current_manifest = DS2CandidateManifest.model_validate_json(
        (repository_root / MANIFEST_PATH).read_text(encoding="utf-8")
    )
    document = Document(str(repository_root / PRIMARY_DOCX))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    tables = [
        [[cell.text for cell in row.cells] for row in table.rows] for table in document.tables
    ]
    records, old_to_new = _correct_records(candidate, paragraphs, tables)
    inverse_ids = {new: old for old, new in old_to_new.items()}
    if current_manifest.candidate_revision == 5:
        if current_manifest.candidate_file_sha256 != R5_SHA256:
            raise ValueError("current DS2 r5 manifest is not bound to the reviewed Candidate")
        r5_groups = current_manifest.review_groups
    elif current_manifest.candidate_revision == 6:
        if current_manifest.predecessor_candidate_file_sha256 != R5_SHA256:
            raise ValueError("current DS2 r6 manifest has the wrong predecessor")
        r5_groups = [
            DS2ReviewGroup(
                group_id=group.group_id,
                title=group.title,
                record_ids=[
                    inverse_ids.get(record_id, record_id) for record_id in group.record_ids
                ],
            )
            for group in current_manifest.review_groups
        ]
    else:
        raise ValueError("current DS2 manifest is neither reviewed r5 nor pending r6")
    decisions, feedback_hashes = _merge_feedback(repository_root, candidate, r5_groups)
    decisions_path = repository_root / DECISIONS_PATH
    atomic_write_json(decisions_path, decisions.model_dump(mode="json"))
    formula_profile_path = repository_root / FORMULA_PROFILE_PATH
    atomic_write_json(formula_profile_path, FORMULA_PROFILE)

    r5_hashes = {record.record_id: record_digest(record) for record in candidate.evidence}
    r6_hashes = {record.record_id: record_digest(record) for record in records}
    retained_ids = [record_id for record_id in decisions.reviewed_record_ids]
    if any(r5_hashes[record_id] != r6_hashes.get(record_id) for record_id in retained_ids):
        raise ValueError("r6 changed a record already accepted in r5")
    changed_old_ids = [old for old in RETURNED_IDS if old_to_new[old] != old]
    required_ids = [old_to_new[old] for old in RETURNED_IDS]
    if len(set(required_ids)) != 22 or set(r6_hashes) != set(retained_ids) | set(required_ids):
        raise ValueError("r6 identity partition does not preserve 98 and revise 22")

    dataset = DS2EvidenceDataset(
        dataset_id=candidate.dataset_id,
        dataset_version=candidate.dataset_version,
        evidence=records,
    )
    r6_path = repository_root / R6_PATH
    candidate_text = (
        json.dumps(dataset.model_dump(mode="json"), ensure_ascii=False, indent=2, sort_keys=True)
        + "\n"
    )
    if r6_path.exists() and r6_path.read_text(encoding="utf-8") != candidate_text:
        raise ValueError("DS2 r6 is immutable; create a new revision instead of overwriting it")
    atomic_write_text(r6_path, candidate_text)
    candidate_sha256 = sha256_file(r6_path)

    groups = [
        DS2ReviewGroup(
            group_id=group.group_id,
            title=group.title,
            record_ids=[old_to_new.get(record_id, record_id) for record_id in group.record_ids],
        )
        for group in r5_groups
    ]
    review_dir, review_index_hash, asset_hashes = _review_pack(
        repository_root,
        records,
        groups,
        candidate_sha256,
        old_to_new,
        decisions,
    )
    section_counts = Counter(
        f"{record.source_span.document_id}:{record.source_span.section_path[-1]}"
        for record in records
    )
    semantic_counts = Counter(record.semantic_unit_type for record in records)
    source_counts = Counter(record.source_type for record in records)
    generated_artifact_paths = {DECISIONS_PATH.as_posix(), FORMULA_PROFILE_PATH.as_posix()}
    source_artifacts = [
        *[
            artifact
            for artifact in current_manifest.source_artifacts
            if artifact.path not in generated_artifact_paths
        ],
        _artifact(repository_root, decisions_path.relative_to(repository_root), "application/json"),
        _artifact(
            repository_root, formula_profile_path.relative_to(repository_root), "application/json"
        ),
    ]
    manifest = DS2CandidateManifest(
        dataset_id=current_manifest.dataset_id,
        dataset_version=current_manifest.dataset_version,
        candidate_revision=6,
        candidate_relative_path=R6_PATH.as_posix(),
        candidate_file_sha256=candidate_sha256,
        candidate_record_sha256=r6_hashes,
        source_artifacts=source_artifacts,
        upstream_approved_file_sha256=current_manifest.upstream_approved_file_sha256,
        record_counts={
            "total": 120,
            "docx": 80,
            "pdf": 40,
            "native_text": source_counts["native_text"],
            "ocr_derived": source_counts["ocr_derived"],
            "table_gold_derived": 10,
            "semantic_table": semantic_counts["table"],
            "procedure_from_table_gold": 1,
        },
        section_record_counts=dict(sorted(section_counts.items())),
        semantic_type_counts=dict(sorted(semantic_counts.items())),
        review_groups=groups,
        review_pack_relative_path=(review_dir / "index.html")
        .relative_to(repository_root)
        .as_posix(),
        review_pack_index_sha256=review_index_hash,
        review_asset_sha256=asset_hashes,
        renderer_profile=current_manifest.renderer_profile,
        canonical_rendered_pdf_sha256=current_manifest.canonical_rendered_pdf_sha256,
        stable_id_profile_sha256=current_manifest.stable_id_profile_sha256,
        normalization_profile_sha256=current_manifest.normalization_profile_sha256,
        predecessor_candidate_file_sha256=R5_SHA256,
        review_decisions_artifact=_artifact(
            repository_root, decisions_path.relative_to(repository_root), "application/json"
        ),
        retained_reviewed_record_ids=retained_ids,
        required_rereview_record_ids=required_ids,
        duplicate_review_record_ids=current_manifest.duplicate_review_record_ids,
        constraints=[
            "Candidate only; exact course_owner approval of the full r6 file SHA-256 is required.",
            "The 98 accepted r5 record Hashes and decisions are retained; only 22 returned records changed.",
            "Rendered equations use courserag.ds2-rendered-formula-linearization.v1 and bind fixed-render BBoxes.",
            "DOCX TableGold 0 is classified as a procedure after owner review; ten TableGold-derived records remain, with nine semantic tables and one procedure.",
            "Only two independent semantic sources exist; derived OCR fixtures remain provenance only.",
            "No P06 Evidence Builder, Chunker, Retriever, LLM Judge or system prediction defined Gold.",
            "Dev/Test remain empty and Test remains unlocked.",
        ],
    )
    manifest_path = repository_root / MANIFEST_PATH
    atomic_write_json(manifest_path, manifest.model_dump(mode="json"))

    history_path = repository_root / HISTORY_PATH
    history = CandidateRevisionHistory.model_validate_json(history_path.read_text(encoding="utf-8"))
    revisions = list(history.revisions)
    if len(revisions) == 5 and revisions[-1].candidate_file_sha256 == R5_SHA256:
        revisions[-1] = CandidateRevisionArtifact(
            revision=5,
            candidate_relative_path=R5_PATH.as_posix(),
            candidate_file_sha256=R5_SHA256,
            status="rejected",
            reason=(
                "course_owner completed all four review groups: 98 records accepted and 22 returned "
                "for source-grounded formula, algorithm, semantic-type or necessary-context corrections."
            ),
        )
        revisions.append(
            CandidateRevisionArtifact(
                revision=6,
                candidate_relative_path=R6_PATH.as_posix(),
                candidate_file_sha256=candidate_sha256,
                status="pending_course_owner_review",
                reason=(
                    "r6 retains all 98 accepted r5 record Hashes and revises only the 22 returned "
                    "records using original OOXML and the frozen renderer snapshot."
                ),
            )
        )
    elif not (
        len(revisions) == 6
        and revisions[-2].candidate_file_sha256 == R5_SHA256
        and revisions[-2].status == "rejected"
        and revisions[-1].candidate_file_sha256 == candidate_sha256
        and revisions[-1].status == "pending_course_owner_review"
    ):
        raise ValueError("DS2 revision history is not at immutable r5 or idempotent r6")
    atomic_write_json(
        history_path,
        CandidateRevisionHistory(
            dataset_id=history.dataset_id,
            dataset_version=history.dataset_version,
            revisions=revisions,
        ).model_dump(mode="json"),
    )
    _update_governance(repository_root / DATASET_ROOT)

    report = f"""# ED-PRE06 正式 DS2 Candidate r6 差量复核报告

- r5 审核结果：98 条通过，22 条退回；四组均已完整提交。
- r6 Candidate：`{R6_PATH.as_posix()}`
- r6 Candidate SHA-256：`{candidate_sha256}`
- 保留记录：98 条，记录 Hash 与 r5 完全一致。
- 需复核记录：22 条；其中 {len(changed_old_ids)} 条稳定 Evidence ID 因原文/源单元变化而更新。
- 审核入口：`{(review_dir / "index.html").relative_to(repository_root).as_posix()}`
- r5 合并决定：`{DECISIONS_PATH.as_posix()}`
- 公式线性化规则：`{FORMULA_PROFILE_PATH.as_posix()}`

## 本轮边界

r6 仍全部为 Candidate，未创建或修改 `approved/ds2`，未运行 P06 B1/B2，Dev/Test 为空且 Test 未锁定。仅需复核 22 条修订记录；98 条 r5 已通过决定被保留。表 0 经 course_owner 认定为算法流程，因此十条 TableGold 来源记录现在按语义计为 9 条 `table` 和 1 条 `procedure`。

## 审批要求

复核无误后必须回复：`批准正式 DS2 Candidate {candidate_sha256}`。泛化的“批准”不会触发 Gold 提升。

## 差量复核范围

- 第 1 组：12 条，含图灵测试上下文、完整列表、感知机算法、表 3.1 行上下文、Q/K/V 符号及涌现能力列表。
- 第 2 组：10 条，含大模型数据集/表 5.2/对齐上下文、公式 14/16/18/19/20、瓦片排序语义类型和 R3-LIVE 上下文。
- 第 3、4 组：0 条；各 30 条 r5 通过决定全部保留，无需重复审核。

复核时先比较页面红框与“修订后的 Gold 原文”，再检查语义类型、`requires_parent` 和必要邻接。公式采用线性化表达，不要求排版外观一致，但变量、上下标、运算符、矩阵行列和公式范围必须与红框一致。若任一条仍有问题，应在该条选择“退回”并写明记录 ID 与原因。

## 验证记录

- r6 生成重复执行：Candidate SHA-256 均为 `{candidate_sha256}`；Manifest 内容可确定性复现。
- 差量完整性：98/98 已通过记录 Hash 与 r5 相同；22 条退回记录全部进入 r6 必复核集合；8 条稳定 ID 更新；四组仍覆盖 120 条唯一记录。
- 视觉检查：固定渲染页 50、146、152、341、342、346、347 已逐页检查；算法及公式红框覆盖可见对象，无页面越界。
- 专项测试：`28 passed`。
- 全量测试：`351 passed, 5 skipped`；skip 均为既有 gated 测试。
- Schema：成功导出 38 份；旧 P02 Pilot 继续验证通过。
- Ruff：315 个文件 format check 通过，lint 通过。
- Mypy：222 个源文件 `Success: no issues found`。

首次全量测试在生产 Compose 容器中受 `COURSERAG_OCR_MAX_WORKERS='1'` 字符串覆盖和未自动加载 `pytest-asyncio` 影响，属于测试环境问题；移除该容器覆盖并显式加载仓库已有异步插件后，全量套件通过。没有为通过测试修改产品配置。
"""
    atomic_write_text(repository_root / PHASE_REPORT_PATH, report)
    return {
        "candidate_file_sha256": candidate_sha256,
        "candidate_record_count": 120,
        "retained_reviewed_count": len(retained_ids),
        "required_rereview_count": len(required_ids),
        "changed_stable_id_count": len(changed_old_ids),
        "review_pack": (review_dir / "index.html").relative_to(repository_root).as_posix(),
        "feedback_file_sha256": feedback_hashes,
        "manifest_sha256": sha256_file(manifest_path),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the course-owner-reviewed DS2 r6 delta")
    parser.add_argument("--repository-root", type=Path, default=Path.cwd())
    args = parser.parse_args()
    print(json.dumps(build_r6(args.repository_root), ensure_ascii=False, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
