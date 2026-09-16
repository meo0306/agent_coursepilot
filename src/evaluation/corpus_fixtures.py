"""Build the deterministic, source-derived Pre-P03 CourseRAG evaluation corpus.

The four derived files are parser/OCR/pagination/migration fixtures.  They are
not independent semantic sources and must never be indexed together with their
primary parent document for retrieval or QA evaluation.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import re
import zipfile
from collections.abc import Iterator
from datetime import UTC, datetime
from pathlib import Path
from xml.etree import ElementTree

import fitz
from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw

from courserag.evals.schemas import (
    CorpusDocument,
    CorpusFixtureArtifact,
    CorpusFixtureManifest,
    DS0CorpusDataset,
    DS1SampleSelection,
    DS1SamplingPlan,
    FixturePageMap,
    FixtureSourceUnitMap,
)
from evaluation.io import atomic_write_json

GENERATOR = "courserag-pre-p03-fixture-builder"
GENERATOR_VERSION = "1.0.1"
SOURCE_DOCX_SHA256 = "c93df4cb4bb533e381647e7591104ff06ecc6a486582d77b3e6d63511428066b"
SOURCE_PDF_SHA256 = "15344771664d40cf190be7790be3d266eb94a2798a4c2608fcf4d1ffa8fd231d"
SOURCE_PDF_PAGE_COUNT = 32
SOURCE_DOCX_FILE_PROPERTY_PAGE_COUNT = 183

COURSE_DOCX = "course_ai_algorithms_systems"
COURSE_PDF = "course_ai_general_education"
DOCX_PRIMARY = "doc_ai_algorithms_systems"
PDF_PRIMARY = "doc_ai_general_education_excerpt"
DOCX_STRESS = "doc_ai_algorithms_systems_structure_stress"
PDF_SCAN_CLEAN = "doc_ai_general_education_scan_clean"
PDF_SCAN_COMPRESSED = "doc_ai_general_education_scan_compressed"
PDF_MIXED = "doc_ai_general_education_mixed"

CLEAN_SOURCE_PAGES = [21, 23, 24, 26, 30]
COMPRESSED_SOURCE_PAGES = [11, 15, 18, 27, 31]
MIXED_RASTER_SOURCE_PAGES = [9, 22, 25, 29, 32]
DOCX_ANCHORS = ["1.1.1", "2.2", "3.5.3", "5.3", "9.7"]

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
ZIP_TIMESTAMP = (2000, 1, 1, 0, 0, 0)
FIXED_CORE_TIMESTAMP = datetime(2000, 1, 1, tzinfo=UTC)
A4_CONTENT_WIDTH_DXA = 10_034
TWO_COLUMN_GAP_DXA = 504
TWO_COLUMN_WIDTH_DXA = (A4_CONTENT_WIDTH_DXA - TWO_COLUMN_GAP_DXA) // 2


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def eval_version(sha256: str) -> str:
    return f"eval-v1-{sha256[:16]}"


def _pdf_metadata() -> dict[str, str]:
    return {
        "author": "",
        "creationDate": "",
        "creator": GENERATOR,
        "keywords": "",
        "modDate": "",
        "producer": f"{GENERATOR}/{GENERATOR_VERSION}",
        "subject": "",
        "title": "",
        "trapped": "",
    }


def _render_page_image(
    page: fitz.Page,
    *,
    dpi: int,
    grayscale: bool,
    jpeg_quality: int | None,
) -> tuple[bytes, str]:
    pixmap = page.get_pixmap(
        dpi=dpi,
        colorspace=fitz.csGRAY if grayscale else fitz.csRGB,
        alpha=False,
    )
    if jpeg_quality is None:
        return pixmap.tobytes("png"), "raster_lossless"
    mode = "L" if grayscale else "RGB"
    image = Image.frombytes(mode, (pixmap.width, pixmap.height), pixmap.samples)
    buffer = io.BytesIO()
    image.save(
        buffer,
        format="JPEG",
        quality=jpeg_quality,
        optimize=False,
        progressive=False,
    )
    return buffer.getvalue(), "raster_jpeg"


def _insert_raster_page(
    output: fitz.Document,
    source_page: fitz.Page,
    *,
    dpi: int,
    grayscale: bool,
    jpeg_quality: int | None,
) -> str:
    image_bytes, representation = _render_page_image(
        source_page,
        dpi=dpi,
        grayscale=grayscale,
        jpeg_quality=jpeg_quality,
    )
    target_page = output.new_page(width=source_page.rect.width, height=source_page.rect.height)
    target_page.insert_image(target_page.rect, stream=image_bytes, keep_proportion=False)
    return representation


def _save_pdf(document: fitz.Document, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.set_metadata(_pdf_metadata())
    document.save(
        output_path,
        garbage=4,
        clean=True,
        deflate=True,
        deflate_images=True,
        no_new_id=True,
        preserve_metadata=False,
    )


def build_raster_subset_pdf(
    source_path: Path,
    output_path: Path,
    *,
    source_pages: list[int],
    dpi: int,
    grayscale: bool,
    jpeg_quality: int | None,
) -> list[FixturePageMap]:
    page_map: list[FixturePageMap] = []
    with fitz.open(source_path) as source, fitz.open() as output:
        for output_number, source_number in enumerate(source_pages, start=1):
            representation = _insert_raster_page(
                output,
                source[source_number - 1],
                dpi=dpi,
                grayscale=grayscale,
                jpeg_quality=jpeg_quality,
            )
            page_map.append(
                FixturePageMap(
                    output_page_number=output_number,
                    source_page_number=source_number,
                    representation=representation,
                )
            )
        _save_pdf(output, output_path)
    return page_map


def build_mixed_pdf(
    source_path: Path,
    output_path: Path,
    *,
    raster_source_pages: list[int],
    dpi: int,
    grayscale: bool,
    jpeg_quality: int,
) -> list[FixturePageMap]:
    raster_page_set = set(raster_source_pages)
    page_map: list[FixturePageMap] = []
    with fitz.open(source_path) as source, fitz.open() as output:
        for source_number in range(1, source.page_count + 1):
            if source_number in raster_page_set:
                representation = _insert_raster_page(
                    output,
                    source[source_number - 1],
                    dpi=dpi,
                    grayscale=grayscale,
                    jpeg_quality=jpeg_quality,
                )
            else:
                output.insert_pdf(source, from_page=source_number - 1, to_page=source_number - 1)
                representation = "native"
            page_map.append(
                FixturePageMap(
                    output_page_number=source_number,
                    source_page_number=source_number,
                    representation=representation,
                )
            )
        _save_pdf(output, output_path)
    return page_map


def _iter_body_blocks(document: DocumentObject) -> Iterator[Paragraph | Table]:
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _heading_level(paragraph: Paragraph) -> int | None:
    style_name = paragraph.style.name if paragraph.style is not None else ""
    match = re.search(r"(?:Heading|标题)\s*([1-9])", style_name, re.IGNORECASE)
    if match:
        return min(int(match.group(1)), 3)
    return None


def _is_list(paragraph: Paragraph) -> bool:
    paragraph_properties = paragraph._p.pPr
    if paragraph_properties is not None and paragraph_properties.numPr is not None:
        return True
    style_name = paragraph.style.name if paragraph.style is not None else ""
    return "List" in style_name or "列表" in style_name


def _add_bookmark(paragraph: Paragraph, *, bookmark_id: int, name: str) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _set_cell_width_dxa(cell, width_dxa: int) -> None:  # noqa: ANN001
    properties = cell._tc.get_or_add_tcPr()
    width = properties.first_child_found_in("w:tcW")
    if width is None:
        width = OxmlElement("w:tcW")
        properties.append(width)
    width.set(qn("w:w"), str(width_dxa))
    width.set(qn("w:type"), "dxa")


def _set_table_fixed_layout(table: Table) -> None:
    properties = table._tbl.tblPr
    layout = properties.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")


def _set_table_width_dxa(table: Table, width_dxa: int) -> None:
    properties = table._tbl.tblPr
    width = properties.first_child_found_in("w:tblW")
    if width is None:
        width = OxmlElement("w:tblW")
        properties.append(width)
    width.set(qn("w:w"), str(width_dxa))
    width.set(qn("w:type"), "dxa")


def _add_page_field(paragraph: Paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def _set_columns(section, count: int) -> None:  # noqa: ANN001
    section_properties = section._sectPr
    columns = section_properties.xpath("./w:cols")
    element = columns[0] if columns else OxmlElement("w:cols")
    if not columns:
        section_properties.append(element)
    element.set(qn("w:num"), str(count))
    element.set(qn("w:space"), "504")


def _restart_page_numbering(section) -> None:  # noqa: ANN001
    section_properties = section._sectPr
    page_number_type = section_properties.xpath("./w:pgNumType")
    element = page_number_type[0] if page_number_type else OxmlElement("w:pgNumType")
    if not page_number_type:
        section_properties.append(element)
    element.set(qn("w:start"), "1")


def _configure_section(
    section,  # noqa: ANN001
    *,
    header_text: str,
    column_count: int,
    restart_page_number: bool,
) -> None:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)
    _set_columns(section, column_count)
    if restart_page_number:
        _restart_page_numbering(section)

    section.header.is_linked_to_previous = False
    header = section.header.paragraphs[0]
    header.text = header_text
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "Arial"
    header.runs[0].font.size = Pt(8)

    section.footer.is_linked_to_previous = False
    footer = section.footer.paragraphs[0]
    footer.clear()
    _add_page_field(footer)
    footer.runs[0].font.name = "Arial"
    footer.runs[0].font.size = Pt(8)


def _configure_docx_styles(document: DocumentObject) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.08
    for level, size in ((1, 15), (2, 12), (3, 10.5)):
        style = document.styles[f"Heading {level}"]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)


def _canonicalize_docx(source_path: Path, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with (
        zipfile.ZipFile(source_path, "r") as source,
        zipfile.ZipFile(
            output_path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
        ) as target,
    ):
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, ZIP_TIMESTAMP)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = 0
            target.writestr(info, source.read(name), compress_type=zipfile.ZIP_DEFLATED)


def build_docx_structure_stress(
    source_path: Path,
    output_path: Path,
) -> list[FixtureSourceUnitMap]:
    source = Document(str(source_path))
    output = Document()
    if output.paragraphs:
        paragraph = output.paragraphs[0]._element
        paragraph.getparent().remove(paragraph)
    _configure_docx_styles(output)

    header_text = next(
        (paragraph.text.strip() for paragraph in source.paragraphs if paragraph.text.strip()),
        source_path.stem,
    )
    blocks = list(_iter_body_blocks(source))
    boundaries = {len(blocks) // 4, len(blocks) // 2, (len(blocks) * 3) // 4}
    source_unit_map: list[FixtureSourceUnitMap] = []
    paragraph_index = 0
    table_index = 0
    section_number = 0

    _configure_section(
        output.sections[0],
        header_text=header_text,
        column_count=1,
        restart_page_number=False,
    )

    for block_index, block in enumerate(blocks):
        if block_index in boundaries and block_index > 0:
            section_number += 1
            section = output.add_section(WD_SECTION.NEW_PAGE)
            _configure_section(
                section,
                header_text=header_text,
                column_count=2 if section_number % 2 else 1,
                restart_page_number=section_number in {1, 3},
            )

        if isinstance(block, Paragraph):
            level = _heading_level(block)
            if level is not None:
                target = output.add_paragraph(style=f"Heading {level}")
            elif _is_list(block):
                target = output.add_paragraph(style="List Number")
            else:
                target = output.add_paragraph(style="Normal")
            target.add_run(block.text)
            bookmark_name = f"src_p_{paragraph_index:06d}"
            _add_bookmark(target, bookmark_id=paragraph_index + 1, name=bookmark_name)
            source_unit_map.append(
                FixtureSourceUnitMap(
                    source_kind="paragraph",
                    source_index=f"paragraph:{paragraph_index}",
                    target_index=f"bookmark:{bookmark_name}",
                    text_sha256=sha256_text(block.text),
                )
            )
            paragraph_index += 1
            continue

        row_count = len(block.rows)
        column_count = max((len(row.cells) for row in block.rows), default=1)
        target_table = output.add_table(rows=row_count, cols=column_count)
        target_table.style = "Table Grid"
        target_table.autofit = False
        _set_table_fixed_layout(target_table)
        available_width_dxa = (
            A4_CONTENT_WIDTH_DXA if section_number % 2 == 0 else TWO_COLUMN_WIDTH_DXA
        )
        _set_table_width_dxa(target_table, available_width_dxa)
        cell_width_dxa = max(720, available_width_dxa // column_count)
        for row_number, source_row in enumerate(block.rows):
            for column_number, source_cell in enumerate(source_row.cells):
                target_cell = target_table.cell(row_number, column_number)
                target_cell.text = source_cell.text
                target_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                _set_cell_width_dxa(target_cell, cell_width_dxa)
                source_unit_map.append(
                    FixtureSourceUnitMap(
                        source_kind="table_cell",
                        source_index=(f"table:{table_index}:row:{row_number}:cell:{column_number}"),
                        target_index=(f"table:{table_index}:row:{row_number}:cell:{column_number}"),
                        text_sha256=sha256_text(source_cell.text),
                    )
                )
        table_index += 1

    properties = output.core_properties
    properties.author = ""
    properties.category = ""
    properties.comments = ""
    properties.created = FIXED_CORE_TIMESTAMP
    properties.identifier = ""
    properties.keywords = ""
    properties.last_modified_by = ""
    properties.modified = FIXED_CORE_TIMESTAMP
    properties.subject = ""
    properties.title = header_text

    temporary_path = output_path.with_suffix(".uncanonicalized.docx")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output.save(str(temporary_path))
    try:
        _canonicalize_docx(temporary_path, output_path)
    finally:
        temporary_path.unlink(missing_ok=True)
    return source_unit_map


def _read_docx_file_property_page_count(path: Path) -> int | None:
    with zipfile.ZipFile(path) as package:
        root = ElementTree.fromstring(package.read("docProps/app.xml"))
    for element in root.iter():
        if element.tag.endswith("}Pages") and element.text and element.text.isdigit():
            return int(element.text)
    return None


def _artifact(
    *,
    course_id: str,
    document_id: str,
    document_role: str,
    parent_document_id: str | None,
    repository_relative_path: str,
    file_path: Path,
    mime_type: str,
    page_count: int | None,
    page_count_basis: str,
    transform_type: str,
    transform_parameters: dict[str, str | int | float | bool | list[int] | list[str]],
    page_map: list[FixturePageMap],
    source_unit_map: list[FixtureSourceUnitMap],
    mutually_exclusive_variant_group: str,
    retrieval_eligible: bool,
) -> CorpusFixtureArtifact:
    return CorpusFixtureArtifact.model_validate(
        {
            "course_id": course_id,
            "document_id": document_id,
            "document_role": document_role,
            "parent_document_id": parent_document_id,
            "repository_relative_path": repository_relative_path,
            "sha256": sha256_file(file_path),
            "size_bytes": file_path.stat().st_size,
            "mime_type": mime_type,
            "page_count": page_count,
            "page_count_basis": page_count_basis,
            "transform_type": transform_type,
            "transform_parameters": transform_parameters,
            "page_map": page_map,
            "source_unit_map": source_unit_map,
            "mutually_exclusive_variant_group": mutually_exclusive_variant_group,
            "retrieval_eligible": retrieval_eligible,
        }
    )


def _candidate_document(
    artifact: CorpusFixtureArtifact,
    *,
    manifest_sha256: str,
    file_property_page_count: int | None = None,
) -> CorpusDocument:
    derived = artifact.document_role == "derived_fixture"
    if artifact.document_id in {DOCX_PRIMARY, PDF_PRIMARY}:
        included_in = [
            "ds0",
            "future_ds1",
            "future_ds2",
            "future_ds3",
            "future_ds4",
            "future_ds5",
        ]
    elif artifact.document_id in {PDF_SCAN_CLEAN, PDF_SCAN_COMPRESSED}:
        included_in = ["ds0", "ds1_parsing", "ds1_ocr", "parser_ocr_migration_fixture"]
    elif artifact.document_id == PDF_MIXED:
        included_in = ["ds0", "ds1_parsing", "parser_ocr_migration_fixture"]
    else:
        included_in = ["ds0", "ds1_parsing", "docx_pagination_migration_fixture"]
    return CorpusDocument(
        record_id=f"ds0-{artifact.document_id.replace('_', '-')}",
        review_status="candidate",
        candidate_source="deterministic_owner_source_transform",
        course_id=artifact.course_id,
        document_id=artifact.document_id,
        filename=Path(artifact.repository_relative_path).name,
        repository_relative_path=artifact.repository_relative_path,
        document_role=artifact.document_role,
        parent_document_id=artifact.parent_document_id,
        mime_type=artifact.mime_type,
        sha256=artifact.sha256,
        document_version=eval_version(artifact.sha256),
        page_count=artifact.page_count,
        page_count_basis=artifact.page_count_basis,
        file_property_page_count=file_property_page_count,
        derivation_manifest_sha256=manifest_sha256 if derived else None,
        mutually_exclusive_variant_group=artifact.mutually_exclusive_variant_group,
        included_in=included_in,
        redistribution_status="owner_authorized_repository_distribution",
    )


def _sampling_plan() -> DS1SamplingPlan:
    return DS1SamplingPlan(
        dataset_id="courserag-eval",
        dataset_version="v1",
        selections=[
            DS1SampleSelection(
                selection_id="ds1-pilot-native-pdf-pages",
                document_id=PDF_PRIMARY,
                sample_kind="native_pdf_page",
                source_page_numbers=[5, 10, 12, 20, 32],
                document_page_numbers=[5, 10, 12, 20, 32],
                notes=["Candidate selection only; no Page Gold has been authored or approved."],
            ),
            DS1SampleSelection(
                selection_id="ds1-pilot-docx-structure-anchors",
                document_id=DOCX_PRIMARY,
                sample_kind="docx_structure_anchor",
                section_anchors=DOCX_ANCHORS,
                notes=[
                    "Bind to P03 DocumentVersion/Build identities before drafting Gold.",
                    "Do not assign DOCX page numbers until the P04 renderer profile is fixed.",
                ],
            ),
            DS1SampleSelection(
                selection_id="ds1-pilot-ocr-clean-pages",
                document_id=PDF_SCAN_CLEAN,
                sample_kind="ocr_page",
                source_page_numbers=CLEAN_SOURCE_PAGES,
                document_page_numbers=[1, 2, 3, 4, 5],
                notes=[
                    "Candidate OCR sample only; transcription and regions require human review."
                ],
            ),
        ],
        blockers=[
            "P03 must provide the real DocumentVersion and Build identities.",
            "P04 must freeze a DOCX renderer version, font manifest, and render profile before pagination Gold.",
            "P05 human review is required for OCR transcription, reading order, and region coordinates.",
        ],
    )


def build_pre_p03_corpus(
    *,
    repository_root: Path,
    output_dir: Path,
    manifest_path: Path,
    ds0_candidate_path: Path,
    sampling_plan_path: Path,
) -> CorpusFixtureManifest:
    repository_root = repository_root.resolve()
    source_dir = repository_root / "data" / "sample_files"
    source_docx = source_dir / "教材-人工智能：从算法到系统.docx"
    source_pdf = source_dir / "人工智能通识教程.pdf"
    if sha256_file(source_docx) != SOURCE_DOCX_SHA256:
        raise ValueError("owner-supplied DOCX hash does not match the frozen Pre-P03 source")
    if sha256_file(source_pdf) != SOURCE_PDF_SHA256:
        raise ValueError("owner-supplied PDF hash does not match the frozen Pre-P03 source")
    with fitz.open(source_pdf) as pdf:
        if pdf.page_count != SOURCE_PDF_PAGE_COUNT:
            raise ValueError("owner-supplied PDF no longer has the verified 32-page page tree")
    property_page_count = _read_docx_file_property_page_count(source_docx)
    if property_page_count != SOURCE_DOCX_FILE_PROPERTY_PAGE_COUNT:
        raise ValueError("owner-supplied DOCX file-property page count changed")

    output_dir.mkdir(parents=True, exist_ok=True)
    clean_path = output_dir / f"{PDF_SCAN_CLEAN}.pdf"
    compressed_path = output_dir / f"{PDF_SCAN_COMPRESSED}.pdf"
    mixed_path = output_dir / f"{PDF_MIXED}.pdf"
    stress_path = output_dir / f"{DOCX_STRESS}.docx"

    clean_map = build_raster_subset_pdf(
        source_pdf,
        clean_path,
        source_pages=CLEAN_SOURCE_PAGES,
        dpi=300,
        grayscale=False,
        jpeg_quality=None,
    )
    compressed_map = build_raster_subset_pdf(
        source_pdf,
        compressed_path,
        source_pages=COMPRESSED_SOURCE_PAGES,
        dpi=180,
        grayscale=True,
        jpeg_quality=70,
    )
    mixed_map = build_mixed_pdf(
        source_pdf,
        mixed_path,
        raster_source_pages=MIXED_RASTER_SOURCE_PAGES,
        dpi=200,
        grayscale=True,
        jpeg_quality=75,
    )
    stress_map = build_docx_structure_stress(source_docx, stress_path)

    pdf_group = "course_ai_general_education:document_variants"
    docx_group = "course_ai_algorithms_systems:document_variants"
    artifacts = [
        _artifact(
            course_id=COURSE_DOCX,
            document_id=DOCX_PRIMARY,
            document_role="primary",
            parent_document_id=None,
            repository_relative_path="data/sample_files/教材-人工智能：从算法到系统.docx",
            file_path=source_docx,
            mime_type=DOCX_MIME,
            page_count=None,
            page_count_basis="pending_fixed_renderer",
            transform_type="identity",
            transform_parameters={"file_property_page_count": property_page_count},
            page_map=[],
            source_unit_map=[],
            mutually_exclusive_variant_group=docx_group,
            retrieval_eligible=True,
        ),
        _artifact(
            course_id=COURSE_PDF,
            document_id=PDF_PRIMARY,
            document_role="primary",
            parent_document_id=None,
            repository_relative_path="data/sample_files/人工智能通识教程.pdf",
            file_path=source_pdf,
            mime_type=PDF_MIME,
            page_count=SOURCE_PDF_PAGE_COUNT,
            page_count_basis="verified_pdf_page_tree",
            transform_type="identity",
            transform_parameters={},
            page_map=[
                FixturePageMap(
                    output_page_number=page,
                    source_page_number=page,
                    representation="native",
                )
                for page in range(1, SOURCE_PDF_PAGE_COUNT + 1)
            ],
            source_unit_map=[],
            mutually_exclusive_variant_group=pdf_group,
            retrieval_eligible=True,
        ),
        _artifact(
            course_id=COURSE_PDF,
            document_id=PDF_SCAN_CLEAN,
            document_role="derived_fixture",
            parent_document_id=PDF_PRIMARY,
            repository_relative_path=f"storage_eval/courserag_corpus/v1/{clean_path.name}",
            file_path=clean_path,
            mime_type=PDF_MIME,
            page_count=5,
            page_count_basis="derived_manifest",
            transform_type="raster_subset_lossless",
            transform_parameters={
                "dpi": 300,
                "color_space": "rgb",
                "image_codec": "png",
                "source_pages": CLEAN_SOURCE_PAGES,
            },
            page_map=clean_map,
            source_unit_map=[],
            mutually_exclusive_variant_group=pdf_group,
            retrieval_eligible=False,
        ),
        _artifact(
            course_id=COURSE_PDF,
            document_id=PDF_SCAN_COMPRESSED,
            document_role="derived_fixture",
            parent_document_id=PDF_PRIMARY,
            repository_relative_path=(f"storage_eval/courserag_corpus/v1/{compressed_path.name}"),
            file_path=compressed_path,
            mime_type=PDF_MIME,
            page_count=5,
            page_count_basis="derived_manifest",
            transform_type="raster_subset_compressed",
            transform_parameters={
                "dpi": 180,
                "color_space": "grayscale",
                "image_codec": "jpeg",
                "jpeg_quality": 70,
                "jpeg_optimize": False,
                "jpeg_progressive": False,
                "source_pages": COMPRESSED_SOURCE_PAGES,
            },
            page_map=compressed_map,
            source_unit_map=[],
            mutually_exclusive_variant_group=pdf_group,
            retrieval_eligible=False,
        ),
        _artifact(
            course_id=COURSE_PDF,
            document_id=PDF_MIXED,
            document_role="derived_fixture",
            parent_document_id=PDF_PRIMARY,
            repository_relative_path=f"storage_eval/courserag_corpus/v1/{mixed_path.name}",
            file_path=mixed_path,
            mime_type=PDF_MIME,
            page_count=SOURCE_PDF_PAGE_COUNT,
            page_count_basis="derived_manifest",
            transform_type="mixed_native_raster",
            transform_parameters={
                "dpi": 200,
                "color_space": "grayscale",
                "image_codec": "jpeg",
                "jpeg_quality": 75,
                "jpeg_optimize": False,
                "jpeg_progressive": False,
                "raster_source_pages": MIXED_RASTER_SOURCE_PAGES,
            },
            page_map=mixed_map,
            source_unit_map=[],
            mutually_exclusive_variant_group=pdf_group,
            retrieval_eligible=False,
        ),
        _artifact(
            course_id=COURSE_DOCX,
            document_id=DOCX_STRESS,
            document_role="derived_fixture",
            parent_document_id=DOCX_PRIMARY,
            repository_relative_path=f"storage_eval/courserag_corpus/v1/{stress_path.name}",
            file_path=stress_path,
            mime_type=DOCX_MIME,
            page_count=None,
            page_count_basis="pending_fixed_renderer",
            transform_type="docx_structure_stress",
            transform_parameters={
                "style_preset": "compact_reference_guide",
                "section_count": 4,
                "column_pattern": ["one", "two", "one", "two"],
                "page_number_restart_sections": [2, 4],
                "header_source": "first_non_empty_source_paragraph",
                "footer": "page_field_only",
                "images_copied": False,
                "table_layout": "fixed_dxa_no_fixed_row_height",
            },
            page_map=[],
            source_unit_map=stress_map,
            mutually_exclusive_variant_group=docx_group,
            retrieval_eligible=False,
        ),
    ]
    manifest = CorpusFixtureManifest(
        generator=GENERATOR,
        generator_version=GENERATOR_VERSION,
        artifacts=artifacts,
    )
    atomic_write_json(manifest_path, manifest.model_dump(mode="json"))
    manifest_sha256 = sha256_file(manifest_path)

    documents = [
        _candidate_document(
            artifact,
            manifest_sha256=manifest_sha256,
            file_property_page_count=(
                property_page_count if artifact.document_id == DOCX_PRIMARY else None
            ),
        )
        for artifact in artifacts
    ]
    candidate_dataset = DS0CorpusDataset(
        dataset_id="courserag-eval",
        dataset_version="v1",
        documents=documents,
    )
    atomic_write_json(ds0_candidate_path, candidate_dataset.model_dump(mode="json"))
    atomic_write_json(sampling_plan_path, _sampling_plan().model_dump(mode="json"))
    return manifest


def _document_path(repository_root: Path, artifact: CorpusFixtureArtifact) -> Path:
    return repository_root / Path(artifact.repository_relative_path)


def verify_docx_source_map(
    *,
    repository_root: Path,
    artifact: CorpusFixtureArtifact,
) -> None:
    source_path = repository_root / "data" / "sample_files" / "教材-人工智能：从算法到系统.docx"
    output_path = _document_path(repository_root, artifact)
    source = Document(str(source_path))
    output = Document(str(output_path))
    bookmark_text: dict[str, str] = {}
    for paragraph in output.paragraphs:
        bookmarks = paragraph._p.xpath("./w:bookmarkStart")
        for bookmark in bookmarks:
            name = bookmark.get(qn("w:name"))
            if name:
                bookmark_text[name] = paragraph.text

    for mapping in artifact.source_unit_map:
        if mapping.source_kind == "paragraph":
            source_index = int(mapping.source_index.split(":")[1])
            source_text = source.paragraphs[source_index].text
            bookmark_name = mapping.target_index.split(":", 1)[1]
            target_text = bookmark_text[bookmark_name]
        else:
            parts = mapping.source_index.split(":")
            table_number, row_number, cell_number = int(parts[1]), int(parts[3]), int(parts[5])
            source_text = source.tables[table_number].cell(row_number, cell_number).text
            target_text = output.tables[table_number].cell(row_number, cell_number).text
        if sha256_text(source_text) != mapping.text_sha256:
            raise ValueError(f"source-unit hash changed: {mapping.source_index}")
        if target_text != source_text:
            raise ValueError(f"source-unit text mismatch: {mapping.source_index}")

    with zipfile.ZipFile(output_path) as package:
        names = package.namelist()
        if any(name.startswith("word/media/") for name in names):
            raise ValueError("DOCX stress fixture unexpectedly contains copied media")
        document_xml = package.read("word/document.xml")
    root = ElementTree.fromstring(document_xml)
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    sections = root.findall(".//w:sectPr", namespace)
    two_column_sections = [
        section
        for section in sections
        if (columns := section.find("w:cols", namespace)) is not None
        and columns.get(qn("w:num")) == "2"
    ]
    page_restarts = [
        section
        for section in sections
        if (page_number := section.find("w:pgNumType", namespace)) is not None
        and page_number.get(qn("w:start")) == "1"
    ]
    if len(sections) < 4 or len(two_column_sections) < 2 or len(page_restarts) < 2:
        raise ValueError("DOCX stress fixture is missing required section/column/page restarts")
    if root.findall(".//w:trHeight", namespace):
        raise ValueError("DOCX stress fixture must not use fixed table row heights")
    allowed_table_widths = {A4_CONTENT_WIDTH_DXA, TWO_COLUMN_WIDTH_DXA}
    for table in root.findall(".//w:tbl", namespace):
        table_width = table.find("./w:tblPr/w:tblW", namespace)
        if table_width is None or table_width.get(qn("w:type")) != "dxa":
            raise ValueError("DOCX stress fixture table is missing a fixed DXA width")
        width_dxa = int(table_width.get(qn("w:w"), "0"))
        if width_dxa not in allowed_table_widths:
            raise ValueError(
                f"DOCX stress fixture table width exceeds its supported layout: {width_dxa}"
            )


def verify_pre_p03_corpus(
    *,
    repository_root: Path,
    manifest_path: Path,
    render_dir: Path | None = None,
) -> dict[str, int | str]:
    manifest = CorpusFixtureManifest.model_validate_json(manifest_path.read_text(encoding="utf-8"))
    artifact_by_id = {artifact.document_id: artifact for artifact in manifest.artifacts}
    for artifact in manifest.artifacts:
        path = _document_path(repository_root, artifact)
        if sha256_file(path) != artifact.sha256:
            raise ValueError(f"artifact hash mismatch: {artifact.document_id}")

    source_pdf_path = _document_path(repository_root, artifact_by_id[PDF_PRIMARY])
    rendered_pages = 0
    with fitz.open(source_pdf_path) as source:
        for document_id in (PDF_SCAN_CLEAN, PDF_SCAN_COMPRESSED, PDF_MIXED):
            artifact = artifact_by_id[document_id]
            with fitz.open(_document_path(repository_root, artifact)) as document:
                if document.page_count != artifact.page_count:
                    raise ValueError(f"page-count mismatch: {document_id}")
                for mapping in artifact.page_map:
                    page = document[mapping.output_page_number - 1]
                    source_page = source[mapping.source_page_number - 1]
                    if (
                        abs(page.rect.width - source_page.rect.width) > 0.01
                        or abs(page.rect.height - source_page.rect.height) > 0.01
                    ):
                        raise ValueError(
                            f"page geometry mismatch: {document_id} page {page.number + 1}"
                        )
                    has_text = bool(page.get_text("text").strip())
                    expected_text = mapping.representation == "native"
                    if has_text != expected_text:
                        raise ValueError(
                            f"text-layer mismatch: {document_id} page {page.number + 1}"
                        )
                    rendered_pages += 1
            if render_dir is not None:
                render_pdf_montage(
                    _document_path(repository_root, artifact),
                    render_dir / f"{document_id}_montage.png",
                )

    verify_docx_source_map(
        repository_root=repository_root,
        artifact=artifact_by_id[DOCX_STRESS],
    )
    return {
        "artifact_count": len(manifest.artifacts),
        "semantic_source_count": manifest.semantic_source_count,
        "verified_pdf_pages": rendered_pages,
        "verified_docx_source_units": len(artifact_by_id[DOCX_STRESS].source_unit_map),
        "manifest_sha256": sha256_file(manifest_path),
    }


def render_pdf_montage(source_path: Path, output_path: Path) -> None:
    thumbnails: list[Image.Image] = []
    labels: list[str] = []
    with fitz.open(source_path) as document:
        for page in document:
            pixmap = page.get_pixmap(dpi=72, colorspace=fitz.csRGB, alpha=False)
            image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
            image.thumbnail((320, 450))
            thumbnails.append(image.copy())
            labels.append(str(page.number + 1))
    columns = 4
    cell_width, cell_height = 340, 480
    rows = (len(thumbnails) + columns - 1) // columns
    montage = Image.new("RGB", (columns * cell_width, rows * cell_height), "white")
    draw = ImageDraw.Draw(montage)
    for index, image in enumerate(thumbnails):
        x = (index % columns) * cell_width + 10
        y = (index // columns) * cell_height + 20
        montage.paste(image, (x, y))
        draw.text((x, 2 + (index // columns) * cell_height), labels[index], fill="black")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    montage.save(output_path, format="PNG", optimize=False)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build or verify the Pre-P03 evaluation corpus.")
    parser.add_argument("command", choices=("build", "verify"))
    parser.add_argument("--repository-root", type=Path, default=Path.cwd())
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("storage_eval/courserag_corpus/v1"),
    )
    parser.add_argument(
        "--manifest-path",
        type=Path,
        default=Path("datasets/courserag_eval/v1/provenance/corpus_fixture_manifest.json"),
    )
    parser.add_argument(
        "--ds0-candidate-path",
        type=Path,
        default=Path("datasets/courserag_eval/v1/candidates/ds0/pilot.json"),
    )
    parser.add_argument(
        "--sampling-plan-path",
        type=Path,
        default=Path("datasets/courserag_eval/v1/plans/ds1_sampling_plan.json"),
    )
    parser.add_argument(
        "--render-dir",
        type=Path,
        default=Path("storage_eval/courserag_corpus/v1/qa_render"),
    )
    args = parser.parse_args()
    repository_root = args.repository_root.resolve()
    if args.command == "build":
        manifest = build_pre_p03_corpus(
            repository_root=repository_root,
            output_dir=repository_root / args.output_dir,
            manifest_path=repository_root / args.manifest_path,
            ds0_candidate_path=repository_root / args.ds0_candidate_path,
            sampling_plan_path=repository_root / args.sampling_plan_path,
        )
        print(f"built {len(manifest.artifacts)} evaluation documents")
        return
    result = verify_pre_p03_corpus(
        repository_root=repository_root,
        manifest_path=repository_root / args.manifest_path,
        render_dir=repository_root / args.render_dir,
    )
    print(result)


if __name__ == "__main__":
    main()
