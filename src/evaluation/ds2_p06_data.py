"""Build the source-grounded formal DS2 Candidate used by the P06 Pilot.

The annotation path is deliberately independent from the P06 Evidence Builder and Chunker. It
reads the two owner-provided primary sources, Poppler geometry, the fixed DOCX renderer snapshot,
and already-approved DS1 records. The command creates Candidate and review artifacts only.
"""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
import unicodedata
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

import fitz
from docx import Document
from docx.document import Document as DocxDocument
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw

from courserag.evals.schemas import (
    BODY_OCR_REGION_ROLES,
    DocxPaginationGold,
    DS1CandidateManifest,
    DS1ParsingDataset,
    DS2CandidateManifest,
    DS2EvidenceDataset,
    DS2ReviewGroup,
    EvidenceBBox,
    EvidenceNeighbor,
    EvidenceOCRProvenance,
    EvidenceRecord,
    EvidenceSourceUnit,
    OCRGold,
    OCRRegion,
    PageGold,
    PageRegion,
    SectionGold,
    TableGold,
    UpstreamApprovedRecordRef,
)
from courserag.parsers.pagination import RendererProfile, canonicalize_pdf
from evaluation.contracts import (
    CandidateRevisionArtifact,
    CandidateRevisionHistory,
    HashedArtifact,
    ReviewStatus,
    SourceSpan,
)
from evaluation.corpus_fixtures import sha256_file
from evaluation.datasets import canonical_json_bytes, record_digest
from evaluation.io import atomic_write_json, atomic_write_text

DATASET_ROOT = Path("datasets/courserag_eval/v1")
PRIMARY_PDF = Path("data/sample_files/人工智能通识教程.pdf")
PRIMARY_DOCX = Path("data/sample_files/教材-人工智能：从算法到系统.docx")
DOCX_RENDER_RUNS = (
    Path("storage_eval/p04_primary_render/run1/教材-人工智能：从算法到系统.pdf"),
    Path("storage_eval/p04_primary_render/run2/教材-人工智能：从算法到系统.pdf"),
)
POPPLER_BBOX_DIR = Path("storage_eval/ds1_p04_poppler_bbox")
RENDERER_PROFILE_PATH = Path("resources/renderers/libreoffice_headless_v1/profile.json")
RENDERER_FONT_LOCK_PATH = Path("resources/renderers/libreoffice_headless_v1/fonts.lock.json")
P04_APPROVED = DATASET_ROOT / "approved/ds1/p04_native_docx.json"
P05_APPROVED = DATASET_ROOT / "approved/ds1/p05_ocr.json"
P04_MANIFEST = DATASET_ROOT / "provenance/ds1_p04_candidate_manifest.json"
P05_MANIFEST = DATASET_ROOT / "provenance/p05_ocr_candidate_manifest.json"
P04_APPROVAL = DATASET_ROOT / "provenance/ds1_p04_approval.json"
P05_APPROVAL = DATASET_ROOT / "provenance/p05_ocr_approval.json"
R1_CANDIDATE_PATH = DATASET_ROOT / "candidates/ds2/p06_evidence_r1.json"
R2_CANDIDATE_PATH = DATASET_ROOT / "candidates/ds2/p06_evidence_r2.json"
R3_CANDIDATE_PATH = DATASET_ROOT / "candidates/ds2/p06_evidence_r3.json"
PREDECESSOR_CANDIDATE_PATH = DATASET_ROOT / "candidates/ds2/p06_evidence_r4.json"
CANDIDATE_PATH = DATASET_ROOT / "candidates/ds2/p06_evidence_r5.json"
CANDIDATE_MANIFEST_PATH = DATASET_ROOT / "provenance/ds2_p06_candidate_manifest.json"
REVISION_HISTORY_PATH = DATASET_ROOT / "provenance/ds2_p06_candidate_revision_history.json"
PHASE_REPORT_PATH = Path("docs/refactor/phase_reports/ED_PRE_P06_DS2_candidate_review.md")

DOCX_DOCUMENT_ID = "doc_ai_algorithms_systems"
DOCX_DOCUMENT_VERSION = "eval-v1-c93df4cb4bb533e3"
DOCX_DOCUMENT_SHA256 = "c93df4cb4bb533e381647e7591104ff06ecc6a486582d77b3e6d63511428066b"
DOCX_COURSE_ID = "course_ai_algorithms_systems"
PDF_DOCUMENT_ID = "doc_ai_general_education_excerpt"
PDF_DOCUMENT_VERSION = "eval-v1-15344771664d40cf"
PDF_DOCUMENT_SHA256 = "15344771664d40cf190be7790be3d266eb94a2798a4c2608fcf4d1ffa8fd231d"
PDF_COURSE_ID = "course_ai_general_education"

DOCX_SECTIONS = (
    "1.1.1",
    "2.2",
    "2.2.2",
    "3.1.3",
    "3.5.3",
    "5.1.1",
    "5.2.1",
    "5.3",
    "5.3.4",
    "6.2.2",
    "6.2.3",
    "6.3",
    "9.1.2",
    "9.2.2",
    "9.5.1",
    "9.7",
)
PDF_SECTIONS = (
    "1.1.1",
    "1.1.2",
    "1.2.1",
    "1.2.2",
    "1.2.3",
    "1.3.1",
    "1.3.4",
    "1.3.7",
)
OCR_PAGE_TO_SECTION = {
    11: "1.1.1",
    15: "1.2.1",
    18: "1.2.1",
    21: "1.2.3",
    22: "1.3.1",
    23: "1.3.1",
    24: "1.3.1",
    26: "1.3.4",
    27: "1.3.4",
    32: "1.3.7",
}
TABLE_TO_SECTION = {
    0: "2.2.2",
    1: "3.1.3",
    6: "5.1.1",
    8: "5.3.4",
    10: "6.2.2",
    12: "6.3",
    18: "9.1.2",
    20: "9.2.2",
    21: "9.5.1",
}
HEADING_NUMBER = re.compile(r"^(?P<number>\d+(?:\.\d+)+)(?:\s|\u3000)")
SENTENCE_END = re.compile(r"[^。！？；!?;]+[。！？；!?;]|[^。！？；!?;]+$")
BLOCK_ID = re.compile(r"^poppler:p(?P<page>\d+):b(?P<block>\d+)$")

# Exact source-grounded strings rejected during the pre-delivery visual audit.
# They are not corrected or rewritten: the malformed Poppler reading sequence is
# excluded and another complete unit from the same frozen Section is selected.
PDF_NATIVE_EXCLUDED_CONTENT_SHA256 = {
    "08cee3982adc09002a325ab715f551465c5e2e6881f6bb108bb3e23c13ea8e77": (
        "page 27 figure insertion leaves the sentence incomplete"
    ),
    "644032eb27fdd78b5af003ca0eaf919bcf5b6f90b011800c5e8508fcc9b97b2e": (
        "page 13 unit begins with an orphaned closing quotation mark"
    ),
    "8ee2fb4eb1951ad5821181877cfa830a1f3a2a01d460fc7c088175786a106a7a": (
        "page 24 multi-column object order does not form a readable sentence"
    ),
}

STABLE_ID_PROFILE = {
    "profile": "courserag.ds2-stable-evidence-id.v1",
    "prefix": "gold-ev-",
    "digest": "sha256-first-32-hex",
    "fields": [
        "course_id",
        "document_version",
        "ordered_source_units",
        "text_assembly",
        "content_sha256",
    ],
}
NORMALIZATION_PROFILE = {
    "profile": "courserag.ds2-text-normalization.v1",
    "unicode": "NFKC",
    "whitespace": "collapse-to-ascii-space",
    "case": "preserve",
}


@dataclass(frozen=True)
class SourcePiece:
    source_unit_id: str
    text: str
    order: float
    section_number: str
    semantic_unit_type: str
    char_start: int | None = None
    char_end: int | None = None
    assembly: str = "single_source_unit"
    table_gold: TableGold | None = None


@dataclass(frozen=True)
class PdfBlock:
    page_number: int
    block_number: int
    bbox: tuple[float, float, float, float]
    text: str
    page_width: float
    page_height: float

    @property
    def source_unit_id(self) -> str:
        return f"poppler:p{self.page_number:03d}:b{self.block_number:04d}"


@dataclass(frozen=True)
class RenderLine:
    page_number: int
    bbox: tuple[float, float, float, float]
    text: str
    normalized: str


@dataclass(frozen=True)
class RenderPageIndex:
    width: float
    height: float
    text: str
    lines: tuple[RenderLine, ...]
    ranges: tuple[tuple[int, int], ...]


@dataclass(frozen=True)
class LocatedPiece:
    piece: SourcePiece
    bboxes: tuple[EvidenceBBox, ...]


@dataclass(frozen=True)
class PdfNativeUnit:
    blocks: tuple[PdfBlock, ...]
    source_units: tuple[EvidenceSourceUnit, ...]
    text: str
    semantic_unit_type: str


def _sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _sha256_bytes(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def _canonical_hash(value: object) -> str:
    return hashlib.sha256(canonical_json_bytes(value)).hexdigest()


def _normalized_text(text: str) -> str:
    return " ".join(unicodedata.normalize("NFKC", text).split())


def _section_path(section_number: str) -> list[str]:
    parts = section_number.split(".")
    return [".".join(parts[:index]) for index in range(1, len(parts) + 1)]


def _match_text(text: str) -> str:
    return "".join(character for character in _normalized_text(text) if not character.isspace())


def _relative(repository_root: Path, path: Path) -> str:
    resolved = path.resolve()
    root = repository_root.resolve()
    if not resolved.is_relative_to(root):
        raise ValueError(f"path escapes repository root: {resolved}")
    return resolved.relative_to(root).as_posix()


def _artifact(repository_root: Path, path: Path, media_type: str) -> HashedArtifact:
    resolved = (repository_root / path).resolve()
    if not resolved.is_file() or not resolved.is_relative_to(repository_root.resolve()):
        raise FileNotFoundError(path)
    return HashedArtifact(
        path=_relative(repository_root, resolved),
        sha256=sha256_file(resolved),
        size_bytes=resolved.stat().st_size,
        media_type=media_type,
    )


def _upstream_ref(
    component: Literal["ds1_p04_native_docx", "ds1_p05_ocr"],
    record: PageGold | SectionGold | TableGold | OCRGold,
) -> UpstreamApprovedRecordRef:
    return UpstreamApprovedRecordRef(
        dataset_component=component,
        record_type=record.record_type,
        record_id=record.record_id,
        record_sha256=record_digest(record),
    )


def _semantic_type(text: str) -> str:
    stripped = text.strip()
    if "\t" in stripped:
        return "table"
    if re.match(r"^(?:[（(]?\d+[）).、]|[①②③④⑤⑥⑦⑧⑨⑩]|[一二三四五六七八九十]+、)", stripped):
        return "list"
    if any(token in stripped for token in ("步骤", "首先", "然后", "最后", "流程", "算法如下")):
        return "procedure"
    if any(token in stripped for token in ("区别", "相比", "不同于", "优点", "缺点", "比较")):
        return "comparison"
    if any(token in stripped for token in ("例如", "举例", "如图", "比如")):
        return "example"
    if any(token in stripped for token in ("定义", "是指", "是一种", "称为")):
        return "definition"
    if any(token in stripped for token in ("原理", "机制", "因此", "因为", "基于")):
        return "principle"
    if any(token in stripped for token in ("应用", "用于", "适用于", "场景", "能够")):
        return "application"
    return "other"


def _complete_segments(
    text: str, source_unit_id: str, order: float, section: str
) -> list[SourcePiece]:
    pieces: list[SourcePiece] = []
    for index, match in enumerate(SENTENCE_END.finditer(text)):
        raw = match.group(0)
        leading = len(raw) - len(raw.lstrip())
        trailing = len(raw.rstrip())
        start = match.start() + leading
        end = match.start() + trailing
        segment = text[start:end]
        if len(_match_text(segment)) < 24:
            continue
        pieces.append(
            SourcePiece(
                source_unit_id=source_unit_id,
                text=segment,
                order=order + index / 1000,
                section_number=section,
                semantic_unit_type=_semantic_type(segment),
                char_start=start,
                char_end=end,
            )
        )
    if not pieces and len(_match_text(text)) >= 24:
        stripped = text.strip()
        start = text.index(stripped)
        pieces.append(
            SourcePiece(
                source_unit_id=source_unit_id,
                text=stripped,
                order=order,
                section_number=section,
                semantic_unit_type=_semantic_type(stripped),
                char_start=start,
                char_end=start + len(stripped),
            )
        )
    return pieces


def _docx_structure(
    path: Path,
) -> tuple[
    DocxDocument,
    list[Paragraph],
    dict[str, tuple[int, int, str]],
    dict[int, int],
]:
    document = Document(str(path))
    paragraphs = list(document.paragraphs)
    headings: list[tuple[int, str, str]] = []
    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        match = HEADING_NUMBER.match(text)
        if match is not None:
            headings.append((index, match.group("number"), text))
    by_number: dict[str, tuple[int, int, str]] = {}
    for index, number, title in headings:
        level = number.count(".") + 1
        boundary = len(paragraphs)
        for next_index, next_number, _ in headings:
            if next_index > index and next_number.count(".") + 1 <= level:
                boundary = next_index
                break
        if number in DOCX_SECTIONS:
            if number in by_number:
                raise ValueError(f"DOCX target heading is not unique: {number}")
            by_number[number] = (index, boundary, title)
    if set(by_number) != set(DOCX_SECTIONS):
        raise ValueError(f"missing DOCX target headings: {set(DOCX_SECTIONS) - set(by_number)}")

    table_after_paragraph: dict[int, int] = {}
    paragraph_index = -1
    table_index = -1
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph_index += 1
        elif child.tag.endswith("}tbl"):
            table_index += 1
            table_after_paragraph[table_index] = paragraph_index
    return document, paragraphs, by_number, table_after_paragraph


def _docx_table_grid(table: DocxTable) -> list[list[str]]:
    return [
        [
            "\n".join(paragraph.text.rstrip() for paragraph in cell.paragraphs).strip()
            for cell in row.cells
        ]
        for row in table.rows
    ]


def _table_text(table: TableGold) -> str:
    return "\n".join("\t".join(cell for cell in row) for row in table.cells)


def _candidate_docx_pieces(
    document: DocxDocument,
    paragraphs: Sequence[Paragraph],
    sections: dict[str, tuple[int, int, str]],
    table_after_paragraph: dict[int, int],
    approved_tables: dict[int, TableGold],
) -> dict[str, list[SourcePiece]]:
    pieces_by_section: dict[str, list[SourcePiece]] = {}
    excluded_ranges = {
        "2.2": (sections["2.2.2"][0], sections["2.2.2"][1]),
        "5.3": (sections["5.3.4"][0], sections["5.3.4"][1]),
    }
    for section in DOCX_SECTIONS:
        start, end, _ = sections[section]
        pieces: list[SourcePiece] = []
        excluded = excluded_ranges.get(section)
        for paragraph_index in range(start + 1, end):
            if excluded is not None and excluded[0] <= paragraph_index < excluded[1]:
                continue
            text = paragraphs[paragraph_index].text.strip()
            if not text or HEADING_NUMBER.match(text):
                continue
            if re.match(r"^[图表]\s*\d", text):
                continue
            pieces.extend(
                _complete_segments(
                    text,
                    f"paragraph:{paragraph_index}",
                    float(paragraph_index),
                    section,
                )
            )
        table_indexes = [index for index, owner in TABLE_TO_SECTION.items() if owner == section]
        for table_index in table_indexes:
            table_gold = approved_tables[table_index]
            table_order = float(table_after_paragraph[table_index]) + 0.5
            pieces.append(
                SourcePiece(
                    source_unit_id=f"table:{table_index}",
                    text=_table_text(table_gold),
                    order=table_order,
                    section_number=section,
                    semantic_unit_type="table",
                    assembly="table_tsv",
                    table_gold=table_gold,
                )
            )
            if section == "3.1.3":
                for row_index, row in enumerate(table_gold.cells[1:], start=1):
                    row_text = "\t".join(row)
                    pieces.append(
                        SourcePiece(
                            source_unit_id=f"table:{table_index}:row:{row_index}",
                            text=row_text,
                            order=table_order + row_index / 1000,
                            section_number=section,
                            semantic_unit_type="application",
                            assembly="table_tsv",
                            table_gold=table_gold,
                        )
                    )
        pieces_by_section[section] = sorted(pieces, key=lambda item: item.order)
    del document
    return pieces_by_section


def _build_render_index(document: fitz.Document) -> dict[int, RenderPageIndex]:
    result: dict[int, RenderPageIndex] = {}
    for page in document:
        lines: list[RenderLine] = []
        stream = ""
        ranges: list[tuple[int, int]] = []
        grouped_words: dict[tuple[int, int], list[tuple[float, float, float, float, str]]] = (
            defaultdict(list)
        )
        for word in page.get_text("words", sort=True):
            x0, y0, x1, y1, text, block_number, line_number, _ = word
            grouped_words[(int(block_number), int(line_number))].append(
                (float(x0), float(y0), float(x1), float(y1), str(text))
            )
        ordered_groups = sorted(
            grouped_words.values(),
            key=lambda words: (min(word[1] for word in words), min(word[0] for word in words)),
        )
        for words in ordered_groups:
            words.sort(key=lambda word: word[0])
            text = " ".join(word[4] for word in words)
            normalized = _match_text(text)
            if not normalized:
                continue
            bbox = (
                round(min(word[0] for word in words), 4),
                round(min(word[1] for word in words), 4),
                round(max(word[2] for word in words), 4),
                round(max(word[3] for word in words), 4),
            )
            start = len(stream)
            stream += normalized
            ranges.append((start, len(stream)))
            lines.append(RenderLine(page.number + 1, bbox, text, normalized))
        result[page.number + 1] = RenderPageIndex(
            width=round(float(page.rect.width), 4),
            height=round(float(page.rect.height), 4),
            text=stream,
            lines=tuple(lines),
            ranges=tuple(ranges),
        )
    return result


def _locate_rendered_text(
    text: str,
    pages: dict[int, RenderPageIndex],
    *,
    rendered_pdf_sha256: str,
    preferred_page: int | None = None,
) -> tuple[EvidenceBBox, ...] | None:
    needle = _match_text(text)
    if len(needle) < 8:
        return None
    matches: list[tuple[int, int, int]] = []
    for page_number, page in pages.items():
        cursor = 0
        while True:
            offset = page.text.find(needle, cursor)
            if offset < 0:
                break
            matches.append((page_number, offset, offset + len(needle)))
            cursor = offset + 1
    if not matches:
        return None
    if preferred_page is not None:
        matches.sort(key=lambda item: (abs(item[0] - preferred_page), item[0], item[1]))
    elif len(matches) != 1:
        return None
    page_number, start, end = matches[0]
    page = pages[page_number]
    selected = [
        line
        for line, (line_start, line_end) in zip(page.lines, page.ranges, strict=True)
        if line_end > start and line_start < end
    ]
    if not selected:
        return None
    x0 = min(line.bbox[0] for line in selected)
    y0 = min(line.bbox[1] for line in selected)
    x1 = max(line.bbox[2] for line in selected)
    y1 = max(line.bbox[3] for line in selected)
    return (
        EvidenceBBox(
            page_number=page_number,
            bbox=(x0, y0, x1, y1),
            page_width=page.width,
            page_height=page.height,
            coordinate_space="rendered_pdf_points_top_left",
            rendered_pdf_sha256=rendered_pdf_sha256,
        ),
    )


def _locate_docx_tables(
    rendered: fitz.Document,
    pages: dict[int, RenderPageIndex],
    tables: Iterable[TableGold],
    canonical_sha256: str,
) -> dict[str, tuple[EvidenceBBox, ...]]:
    located: dict[str, tuple[EvidenceBBox, ...]] = {}
    for table in tables:
        if table.rendered_page_fragments:
            page_numbers = [
                fragment.physical_page_index for fragment in table.rendered_page_fragments
            ]
        else:
            tokens = [
                _match_text(cell)
                for row in table.cells
                for cell in row
                if len(_match_text(cell)) >= 4
            ]
            scored: list[tuple[int, int]] = []
            for page_number, page in pages.items():
                score = sum(1 for token in tokens if token in page.text)
                if score:
                    scored.append((score, page_number))
            if not scored:
                raise ValueError(f"DOCX table cannot be located in fixed render: {table.record_id}")
            scored.sort(reverse=True)
            best_score = scored[0][0]
            page_numbers = sorted(page for score, page in scored if score >= max(2, best_score - 1))
            if len(page_numbers) > 2:
                page_numbers = page_numbers[:2]
        boxes: list[EvidenceBBox] = []
        for page_number in page_numbers:
            page = rendered[page_number - 1]
            rects: list[fitz.Rect] = []
            for row in table.cells:
                for cell in row:
                    stripped = cell.strip()
                    if len(_match_text(stripped)) < 3:
                        continue
                    rects.extend(page.search_for(stripped))
            if rects:
                union = fitz.Rect(rects[0])
                for rect in rects[1:]:
                    union.include_rect(rect)
                # Text search locates only non-empty cells.  Expand to the
                # nearest full-width table rules so empty cells, the first
                # column, and cross-page fragments remain inside the Evidence
                # BBox.  These are vector rules from the frozen renderer PDF,
                # not Parser/P06 predictions.
                minimum_rule_width = max(100.0, union.width * 0.8)
                horizontal_rules = [
                    drawing["rect"]
                    for drawing in page.get_drawings()
                    if drawing["rect"].width >= minimum_rule_width and drawing["rect"].height <= 2.0
                ]
                rules_above = [rule for rule in horizontal_rules if rule.y0 <= union.y0 + 2.0]
                rules_below = [rule for rule in horizontal_rules if rule.y1 >= union.y1 - 2.0]
                if rules_above and rules_below:
                    upper = max(rules_above, key=lambda rule: rule.y0)
                    lower = min(rules_below, key=lambda rule: rule.y1)
                    if lower.y1 > upper.y0:
                        union = fitz.Rect(
                            min(union.x0, upper.x0, lower.x0),
                            upper.y0,
                            max(union.x1, upper.x1, lower.x1),
                            lower.y1,
                        )
                bbox = tuple(round(float(value), 4) for value in union)
            else:
                page_index = pages[page_number]
                matching_lines = [
                    line
                    for line in page_index.lines
                    if any(
                        len(token) >= 4 and token in line.normalized
                        for token in (_match_text(cell) for row in table.cells for cell in row)
                    )
                ]
                if not matching_lines:
                    raise ValueError(
                        f"DOCX table page has no independently matched cells: {table.record_id}"
                    )
                bbox = (
                    min(line.bbox[0] for line in matching_lines),
                    min(line.bbox[1] for line in matching_lines),
                    max(line.bbox[2] for line in matching_lines),
                    max(line.bbox[3] for line in matching_lines),
                )
            page_index = pages[page_number]
            boxes.append(
                EvidenceBBox(
                    page_number=page_number,
                    bbox=bbox,
                    page_width=page_index.width,
                    page_height=page_index.height,
                    coordinate_space="rendered_pdf_points_top_left",
                    rendered_pdf_sha256=canonical_sha256,
                    source_region_id=table.record_id,
                )
            )
        located[table.record_id] = tuple(boxes)
    return located


def _choose_evenly(pieces: Sequence[LocatedPiece], count: int) -> list[LocatedPiece]:
    if len(pieces) < count:
        raise ValueError(
            f"only {len(pieces)} independently located units are available for {count}"
        )
    if count == 1:
        return [pieces[len(pieces) // 2]]
    indexes = [round(position * (len(pieces) - 1) / (count - 1)) for position in range(count)]
    result: list[LocatedPiece] = []
    used: set[int] = set()
    for index in indexes:
        while index in used and index + 1 < len(pieces):
            index += 1
        if index in used:
            index = next(candidate for candidate in range(len(pieces)) if candidate not in used)
        used.add(index)
        result.append(pieces[index])
    return sorted(result, key=lambda item: item.piece.order)


def _even_indexes(length: int, count: int) -> list[int]:
    if length < count:
        raise ValueError(f"only {length} independently located units are available for {count}")
    if count == 1:
        return [length // 2]
    indexes = [round(position * (length - 1) / (count - 1)) for position in range(count)]
    used: set[int] = set()
    for position, index in enumerate(indexes):
        while index in used and index + 1 < length:
            index += 1
        if index in used:
            index = next(candidate for candidate in range(length) if candidate not in used)
        indexes[position] = index
        used.add(index)
    return sorted(indexes)


def _stable_id(
    *,
    course_id: str,
    document_version: str,
    source_units: Sequence[EvidenceSourceUnit],
    text_assembly: str,
    content_sha256: str,
) -> str:
    identity = {
        "course_id": course_id,
        "document_version": document_version,
        "ordered_source_units": [unit.model_dump(mode="json") for unit in source_units],
        "text_assembly": text_assembly,
        "content_sha256": content_sha256,
    }
    return f"gold-ev-{_canonical_hash(identity)[:32]}"


def _heading_neighbor(
    *,
    document_id: str,
    document_version: str,
    document_sha256: str,
    section_number: str,
    heading_text: str,
    heading_unit_id: str,
    page_number: int,
) -> EvidenceNeighbor:
    return EvidenceNeighbor(
        relation="parent_heading",
        source_span=SourceSpan(
            document_id=document_id,
            document_version=document_version,
            document_sha256=document_sha256,
            section_path=_section_path(section_number),
            page_start=page_number,
            page_end=page_number,
            block_start=heading_unit_id,
            block_end=heading_unit_id,
        ),
        text=heading_text,
        text_sha256=_sha256_text(heading_text),
    )


def _record_from_piece(
    located: LocatedPiece,
    *,
    course_id: str,
    document_id: str,
    document_version: str,
    document_sha256: str,
    section_number: str,
    heading_text: str,
    heading_unit_id: str,
    upstream: Sequence[UpstreamApprovedRecordRef],
) -> EvidenceRecord:
    piece = located.piece
    source_unit = EvidenceSourceUnit(
        source_unit_id=piece.source_unit_id,
        char_start=piece.char_start,
        char_end=piece.char_end,
        text_sha256=_sha256_text(piece.text),
    )
    content_hash = _sha256_text(piece.text)
    normalized_hash = _sha256_text(_normalized_text(piece.text))
    source_span = SourceSpan(
        document_id=document_id,
        document_version=document_version,
        document_sha256=document_sha256,
        section_path=_section_path(section_number),
        page_start=min(box.page_number for box in located.bboxes),
        page_end=max(box.page_number for box in located.bboxes),
        block_start=piece.source_unit_id,
        block_end=piece.source_unit_id,
        char_start=piece.char_start,
        char_end=piece.char_end,
    )
    requires_parent = piece.semantic_unit_type in {
        "list",
        "table",
        "formula",
        "other",
    } or piece.text.startswith(("其中", "因此", "上述", "这", "（", "("))
    neighbors: list[EvidenceNeighbor] = []
    if requires_parent:
        neighbors.append(
            _heading_neighbor(
                document_id=document_id,
                document_version=document_version,
                document_sha256=document_sha256,
                section_number=section_number,
                heading_text=heading_text,
                heading_unit_id=heading_unit_id,
                page_number=located.bboxes[0].page_number,
            )
        )
    evidence_id = _stable_id(
        course_id=course_id,
        document_version=document_version,
        source_units=[source_unit],
        text_assembly=piece.assembly,
        content_sha256=content_hash,
    )
    return EvidenceRecord(
        record_id=evidence_id,
        evidence_id=evidence_id,
        candidate_source="source_grounded_assisted",
        annotation_profile="formal_ds2_v1",
        course_id=course_id,
        source_span=source_span,
        source_type="native_text",
        gold_text=piece.text,
        content_sha256=content_hash,
        normalized_content_sha256=normalized_hash,
        source_units=[source_unit],
        text_assembly=piece.assembly,
        upstream_record_refs=list(upstream),
        semantic_unit_type=piece.semantic_unit_type,
        requires_parent=requires_parent,
        necessary_neighbor_text=[neighbor.text for neighbor in neighbors],
        necessary_neighbors=neighbors,
        bboxes=list(located.bboxes),
        ocr_confidence=None,
        other_type_review_reason=(
            "Source-grounded complete statement does not fit a narrower frozen semantic label; "
            "course_owner must confirm or relabel it."
            if piece.semantic_unit_type == "other"
            else None
        ),
    )


def _docx_records(
    repository_root: Path,
    approved_p04: DS1ParsingDataset,
    p04_manifest: DS1CandidateManifest,
) -> tuple[list[EvidenceRecord], str]:
    document, paragraphs, sections, table_after = _docx_structure(repository_root / PRIMARY_DOCX)
    approved_tables: dict[int, TableGold] = {}
    for record in approved_p04.records:
        if not isinstance(record, TableGold) or record.source_span.document_id != DOCX_DOCUMENT_ID:
            continue
        block_start = record.source_span.block_start
        if block_start is None:
            raise ValueError(f"Approved DOCX TableGold has no stable block: {record.record_id}")
        approved_tables[int(block_start.split(":")[1])] = record
    if set(approved_tables) != set(TABLE_TO_SECTION):
        raise ValueError("Approved P04 DOCX table set differs from the frozen DS2 table coverage")
    pieces = _candidate_docx_pieces(document, paragraphs, sections, table_after, approved_tables)

    render_contents = [(repository_root / path).read_bytes() for path in DOCX_RENDER_RUNS]
    canonical_contents = [canonicalize_pdf(content) for content in render_contents]
    canonical_hashes = [_sha256_bytes(content) for content in canonical_contents]
    expected_canonical = p04_manifest.canonical_rendered_pdf_sha256[DOCX_DOCUMENT_ID]
    if canonical_hashes != [expected_canonical, expected_canonical]:
        raise ValueError("fixed DOCX Renderer canonical Hash is no longer repeatable")
    with fitz.open(stream=render_contents[0], filetype="pdf") as rendered:
        if len(rendered) != 636:
            raise ValueError("formal DOCX renderer baseline must remain 636 physical pages")
        render_index = _build_render_index(rendered)
        table_boxes = _locate_docx_tables(
            rendered,
            render_index,
            approved_tables.values(),
            expected_canonical,
        )
        section_heading_pages: dict[str, int] = {}
        approved_anchor_pages = {
            record.section_path[-1]: record.expected_physical_page_index
            for record in approved_p04.records
            if isinstance(record, DocxPaginationGold)
            and record.document_id == DOCX_DOCUMENT_ID
            and record.section_path
        }
        previous_heading_page: int | None = None
        for section, (heading_index, _, heading_text) in sections.items():
            located_heading = _locate_rendered_text(
                heading_text,
                render_index,
                rendered_pdf_sha256=expected_canonical,
                preferred_page=approved_anchor_pages.get(section, previous_heading_page),
            )
            if located_heading is None:
                raise ValueError(f"DOCX section heading cannot be uniquely rendered: {section}")
            section_heading_pages[section] = located_heading[0].page_number
            if (
                previous_heading_page is not None
                and located_heading[0].page_number < previous_heading_page
            ):
                raise ValueError(f"DOCX target heading pages are not monotonic at {section}")
            previous_heading_page = located_heading[0].page_number
            if paragraphs[heading_index].text.strip() != heading_text:
                raise ValueError(f"DOCX heading source changed: {section}")

        section_gold = {
            record.title.split(maxsplit=1)[0].replace("\u3000", ""): record
            for record in approved_p04.records
            if isinstance(record, SectionGold) and record.document_id == DOCX_DOCUMENT_ID
        }
        table_refs = {
            index: _upstream_ref("ds1_p04_native_docx", table)
            for index, table in approved_tables.items()
        }
        records: list[EvidenceRecord] = []
        for section in DOCX_SECTIONS:
            heading_index, _, heading_text = sections[section]
            located_pieces: list[LocatedPiece] = []
            mandatory: LocatedPiece | None = None
            for piece in pieces[section]:
                if piece.table_gold is not None:
                    located_bboxes: tuple[EvidenceBBox, ...] | None = table_boxes[
                        piece.table_gold.record_id
                    ]
                else:
                    located_bboxes = _locate_rendered_text(
                        piece.text,
                        render_index,
                        rendered_pdf_sha256=expected_canonical,
                        preferred_page=section_heading_pages[section],
                    )
                if located_bboxes is None:
                    continue
                located_piece = LocatedPiece(piece, located_bboxes)
                located_pieces.append(located_piece)
                if piece.table_gold is not None and piece.source_unit_id.count(":") == 1:
                    mandatory = located_piece
            selectable = [piece for piece in located_pieces if piece is not mandatory]
            selected = _choose_evenly(selectable, 4 if mandatory is not None else 5)
            if mandatory is not None:
                selected.append(mandatory)
            selected.sort(key=lambda item: item.piece.order)
            if len(selected) != 5:
                raise AssertionError(f"DOCX section {section} did not produce five Evidence units")
            section_refs: list[UpstreamApprovedRecordRef] = []
            approved_section = section_gold.get(section)
            if approved_section is not None:
                section_refs.append(_upstream_ref("ds1_p04_native_docx", approved_section))
            for located_piece in selected:
                refs = list(section_refs)
                if located_piece.piece.table_gold is not None:
                    block_start = located_piece.piece.table_gold.source_span.block_start
                    if block_start is None:
                        raise ValueError("selected DOCX TableGold has no stable block")
                    table_index = int(block_start.split(":")[1])
                    refs.append(table_refs[table_index])
                records.append(
                    _record_from_piece(
                        located_piece,
                        course_id=DOCX_COURSE_ID,
                        document_id=DOCX_DOCUMENT_ID,
                        document_version=DOCX_DOCUMENT_VERSION,
                        document_sha256=DOCX_DOCUMENT_SHA256,
                        section_number=section,
                        heading_text=heading_text,
                        heading_unit_id=f"paragraph:{heading_index}",
                        upstream=refs,
                    )
                )
    if len(records) != 80:
        raise AssertionError(f"expected 80 DOCX Evidence records, received {len(records)}")
    return records, expected_canonical


def _poppler_page(repository_root: Path, page_number: int) -> tuple[float, float, list[PdfBlock]]:
    path = repository_root / POPPLER_BBOX_DIR / f"page-{page_number:03d}.html"
    if not path.is_file():
        raise FileNotFoundError(path)
    root = ET.fromstring(path.read_text(encoding="utf-8"))
    namespace = {"x": "http://www.w3.org/1999/xhtml"}
    page = root.find(".//x:page", namespace)
    if page is None:
        raise ValueError(f"Poppler page is missing: {path}")
    width = round(float(page.attrib["width"]), 4)
    height = round(float(page.attrib["height"]), 4)
    blocks: list[PdfBlock] = []
    for block_number, block in enumerate(page.findall(".//x:block", namespace)):
        lines: list[str] = []
        for line in block.findall("./x:line", namespace):
            words = [word.text or "" for word in line.findall("./x:word", namespace)]
            joined = ""
            for word in words:
                if (
                    joined
                    and joined[-1:].isascii()
                    and word[:1].isascii()
                    and joined[-1:].isalnum()
                    and word[:1].isalnum()
                ):
                    joined += " "
                joined += word
            if joined.strip():
                lines.append(joined.strip())
        text = "\n".join(lines).strip()
        if not text:
            continue
        bbox = (
            round(float(block.attrib["xMin"]), 4),
            round(float(block.attrib["yMin"]), 4),
            round(float(block.attrib["xMax"]), 4),
            round(float(block.attrib["yMax"]), 4),
        )
        blocks.append(PdfBlock(page_number, block_number, bbox, text, width, height))
    return width, height, blocks


def _block_range(section: SectionGold) -> tuple[int, int, int, int]:
    start_match = BLOCK_ID.match(section.source_span.block_start or "")
    end_match = BLOCK_ID.match(section.source_span.block_end or "")
    if start_match is None or end_match is None:
        raise ValueError(f"PDF Section has invalid Poppler boundaries: {section.record_id}")
    return (
        int(start_match.group("page")),
        int(start_match.group("block")),
        int(end_match.group("page")),
        int(end_match.group("block")),
    )


def _blocks_for_section(repository_root: Path, section: SectionGold) -> list[PdfBlock]:
    start_page, start_block, end_page, end_block = _block_range(section)
    result: list[PdfBlock] = []
    for page_number in range(start_page, end_page + 1):
        _, _, blocks = _poppler_page(repository_root, page_number)
        for block in blocks:
            if page_number == start_page and block.block_number < start_block:
                continue
            if page_number == end_page and block.block_number > end_block:
                continue
            if block.bbox[1] < block.page_height * 0.08 or block.bbox[3] > block.page_height * 0.93:
                continue
            if HEADING_NUMBER.match(block.text.strip()) or re.match(
                r"^图\s*\d", block.text.strip()
            ):
                continue
            result.append(block)
    return result


def _pdf_complete_units(blocks: Sequence[PdfBlock]) -> list[PdfNativeUnit]:
    filtered: list[PdfBlock] = []
    for block in blocks:
        stripped = block.text.strip()
        if len(_match_text(stripped)) < 20 and not stripped.endswith(
            ("。", "！", "？", ";", "；", "!", "?")
        ):
            continue
        filtered.append(block)
    stream = ""
    ranges: list[tuple[int, int, PdfBlock]] = []
    for block in filtered:
        start = len(stream)
        stream += block.text
        ranges.append((start, len(stream), block))
        stream += "\n"
    units: list[PdfNativeUnit] = []
    for match in SENTENCE_END.finditer(stream):
        raw = match.group(0)
        leading = len(raw) - len(raw.lstrip())
        trailing = len(raw.rstrip())
        start = match.start() + leading
        end = match.start() + trailing
        if end <= start or not stream[start:end].rstrip().endswith(
            ("。", "！", "？", ";", "；", "!", "?")
        ):
            continue
        parts: list[str] = []
        source_units: list[EvidenceSourceUnit] = []
        used_blocks: list[PdfBlock] = []
        for block_start, block_end, block in ranges:
            overlap_start = max(start, block_start)
            overlap_end = min(end, block_end)
            if overlap_end <= overlap_start:
                continue
            local_start = overlap_start - block_start
            local_end = overlap_end - block_start
            part = block.text[local_start:local_end].strip()
            if not part:
                continue
            part_offset = block.text.find(part, local_start, local_end + 1)
            parts.append(part)
            source_units.append(
                EvidenceSourceUnit(
                    source_unit_id=block.source_unit_id,
                    char_start=part_offset,
                    char_end=part_offset + len(part),
                    text_sha256=_sha256_text(part),
                )
            )
            used_blocks.append(block)
        text = "\n".join(parts)
        if len(_match_text(text)) < 24 or not source_units:
            continue
        if _sha256_text(text) in PDF_NATIVE_EXCLUDED_CONTENT_SHA256:
            continue
        readable = sum(
            character.isalnum()
            or "\u4e00" <= character <= "\u9fff"
            or character.isspace()
            or character in "，。！？；：、（）()《》“”‘’[]+-/%.,:;!?"
            for character in text
        ) / len(text)
        if readable < 0.9:
            continue
        units.append(
            PdfNativeUnit(
                blocks=tuple(used_blocks),
                source_units=tuple(source_units),
                text=text,
                semantic_unit_type=_semantic_type(text),
            )
        )
    return units


def _pdf_native_record(
    unit: PdfNativeUnit,
    *,
    section_number: str,
    section: SectionGold,
    upstream: Sequence[UpstreamApprovedRecordRef],
) -> EvidenceRecord:
    source_units = list(unit.source_units)
    assembly = "single_source_unit" if len(source_units) == 1 else "ordered_source_units_newline"
    content_hash = _sha256_text(unit.text)
    evidence_id = _stable_id(
        course_id=PDF_COURSE_ID,
        document_version=PDF_DOCUMENT_VERSION,
        source_units=source_units,
        text_assembly=assembly,
        content_sha256=content_hash,
    )
    bboxes = [
        EvidenceBBox(
            page_number=block.page_number,
            bbox=block.bbox,
            page_width=block.page_width,
            page_height=block.page_height,
            coordinate_space="pdf_points_top_left",
            source_region_id=block.source_unit_id,
        )
        for block in unit.blocks
    ]
    requires_parent = unit.semantic_unit_type in {"list", "other"} or unit.text.startswith(
        ("其中", "因此", "上述", "这", "（", "(")
    )
    neighbors = (
        [
            _heading_neighbor(
                document_id=PDF_DOCUMENT_ID,
                document_version=PDF_DOCUMENT_VERSION,
                document_sha256=PDF_DOCUMENT_SHA256,
                section_number=section_number,
                heading_text=section.title,
                heading_unit_id=section.source_span.block_start or "",
                page_number=unit.blocks[0].page_number,
            )
        ]
        if requires_parent
        else []
    )
    return EvidenceRecord(
        record_id=evidence_id,
        evidence_id=evidence_id,
        candidate_source="source_grounded_poppler_blocks",
        annotation_profile="formal_ds2_v1",
        course_id=PDF_COURSE_ID,
        source_span=SourceSpan(
            document_id=PDF_DOCUMENT_ID,
            document_version=PDF_DOCUMENT_VERSION,
            document_sha256=PDF_DOCUMENT_SHA256,
            section_path=_section_path(section_number),
            page_start=unit.blocks[0].page_number,
            page_end=unit.blocks[-1].page_number,
            block_start=unit.blocks[0].source_unit_id,
            block_end=unit.blocks[-1].source_unit_id,
            char_start=(source_units[0].char_start if len(source_units) == 1 else None),
            char_end=(source_units[0].char_end if len(source_units) == 1 else None),
        ),
        source_type="native_text",
        gold_text=unit.text,
        content_sha256=content_hash,
        normalized_content_sha256=_sha256_text(_normalized_text(unit.text)),
        source_units=source_units,
        text_assembly=assembly,
        upstream_record_refs=list(upstream),
        semantic_unit_type=unit.semantic_unit_type,
        requires_parent=requires_parent,
        necessary_neighbor_text=[neighbor.text for neighbor in neighbors],
        necessary_neighbors=neighbors,
        bboxes=bboxes,
        ocr_confidence=None,
        other_type_review_reason=(
            "Source-grounded complete statement does not fit a narrower frozen semantic label; "
            "course_owner must confirm or relabel it."
            if unit.semantic_unit_type == "other"
            else None
        ),
    )


def _source_region_for_ocr_region(
    region_id: str,
    page_gold: PageGold | None,
) -> PageRegion | None:
    if page_gold is None:
        return None
    for region in [*page_gold.regions, *page_gold.noise_regions]:
        if region.region_id == region_id:
            return region
    return None


def _ocr_record_for_page(approved_p05: DS1ParsingDataset, page_number: int) -> OCRGold:
    matches = [
        record
        for record in approved_p05.records
        if isinstance(record, OCRGold) and record.source_span.page_start == page_number
    ]
    if len(matches) != 1:
        raise ValueError(f"Approved P05 must contain exactly one OCR record for page {page_number}")
    return matches[0]


def _ocr_evidence(
    *,
    page_number: int,
    section_number: str,
    section: SectionGold,
    ocr_record: OCRGold,
    page_gold: PageGold | None,
) -> tuple[EvidenceRecord, list[tuple[str, int, int]]]:
    start_page, start_block, end_page, end_block = _block_range(section)
    eligible: list[tuple[OCRRegion, str]] = []
    for region_id in ocr_record.reading_order:
        region = next(region for region in ocr_record.regions if region.region_id == region_id)
        if region.role not in BODY_OCR_REGION_ROLES or not region.gold_text:
            continue
        source_region = _source_region_for_ocr_region(region.source_page_region_id or "", page_gold)
        block_id = source_region.block_id if source_region is not None else None
        if block_id is None:
            block_id = f"ocr-source:p{page_number:03d}:{region.region_id}"
        match = BLOCK_ID.match(block_id)
        if match is not None:
            block_number = int(match.group("block"))
            if page_number == start_page and block_number < start_block:
                continue
            if page_number == end_page and block_number > end_block:
                continue
        eligible.append((region, block_id))
    if page_number > start_page:
        first_complete_end = next(
            (
                index
                for index, (region, _) in enumerate(eligible)
                if str(region.gold_text).rstrip().endswith(("。", "！", "？", ";", "；"))
            ),
            None,
        )
        if first_complete_end is not None:
            eligible = eligible[first_complete_end + 1 :]
    usable: list[tuple[OCRRegion, str]] = []
    for region, block_id in eligible:
        region_text = region.gold_text or ""
        is_atomic_structure = (
            region.role in {"title", "list", "table", "formula"}
            and len(_match_text(region_text)) >= 6
        )
        if (
            len(_match_text(region_text)) < 20
            and not region_text.rstrip().endswith(("。", "！", "？", ";", "；"))
            and not is_atomic_structure
        ):
            continue
        usable.append((region, block_id))
    stream = ""
    region_ranges: list[tuple[int, int, OCRRegion, str]] = []
    # A heading/list immediately before a body sentence is a separate semantic
    # unit.  Do not silently concatenate it into that sentence merely because
    # structural text normally has no terminal punctuation.
    for region, block_id in usable:
        if region.role != "body":
            continue
        region_text = region.gold_text or ""
        start = len(stream)
        stream += region_text
        region_ranges.append((start, len(stream), region, block_id))
        stream += "\n"
    selected: list[tuple[OCRRegion, str, str, int, int]] = []
    for sentence in SENTENCE_END.finditer(stream):
        raw = sentence.group(0)
        leading = len(raw) - len(raw.lstrip())
        trailing = len(raw.rstrip())
        sentence_start = sentence.start() + leading
        sentence_end = sentence.start() + trailing
        candidate_parts: list[tuple[OCRRegion, str, str, int, int]] = []
        for region_start, region_end, region, block_id in region_ranges:
            overlap_start = max(sentence_start, region_start)
            overlap_end = min(sentence_end, region_end)
            if overlap_end <= overlap_start:
                continue
            region_text = region.gold_text or ""
            local_start = overlap_start - region_start
            local_end = overlap_end - region_start
            part = region_text[local_start:local_end].strip()
            if not part:
                continue
            part_offset = region_text.find(part, local_start, local_end + 1)
            candidate_parts.append((region, block_id, part, part_offset, part_offset + len(part)))
        candidate_text = "\n".join(part for _, _, part, _, _ in candidate_parts)
        if len(_match_text(candidate_text)) >= 24 and candidate_text.rstrip().endswith(
            ("。", "！", "？", ";", "；")
        ):
            selected = candidate_parts
            break
    if not selected:
        # Some approved OCR pages end exactly where a new subsection begins, so
        # their only complete unit is a visible list item/title rather than a
        # punctuated sentence (page 15 is the canonical example).  Prefer an
        # atomic structural unit over joining it to the following, unavailable
        # page or accepting a dangling body fragment.
        structural_roles = ("list", "table", "formula", "title")
        for role in structural_roles:
            for region, block_id in usable:
                region_text = (region.gold_text or "").strip()
                if region.role != role or len(_match_text(region_text)) < 6:
                    continue
                if re.fullmatch(r"[\d.\s]+", region_text):
                    continue
                selected = [(region, block_id, region_text, 0, len(region_text))]
                break
            if selected:
                break
    if not selected:
        raise ValueError(f"no complete body OCR unit is available on source page {page_number}")
    gold_text = "\n".join(part for _, _, part, _, _ in selected)
    units = [
        EvidenceSourceUnit(
            source_unit_id=str(region.region_id),
            char_start=char_start,
            char_end=char_end,
            text_sha256=_sha256_text(part),
        )
        for region, _, part, char_start, char_end in selected
    ]
    content_hash = _sha256_text(gold_text)
    evidence_id = _stable_id(
        course_id=PDF_COURSE_ID,
        document_version=PDF_DOCUMENT_VERSION,
        source_units=units,
        text_assembly="ocr_regions_newline",
        content_sha256=content_hash,
    )
    bboxes = [
        EvidenceBBox(
            page_number=page_number,
            bbox=tuple(float(value) for value in region.source_bbox_pdf_points),
            page_width=ocr_record.source_page_width_points,
            page_height=ocr_record.source_page_height_points,
            coordinate_space="pdf_points_top_left",
            source_region_id=str(region.region_id),
        )
        for region, _, _, _, _ in selected
    ]
    block_ids = [block_id for _, block_id, _, _, _ in selected]
    excluded_source_ranges: list[tuple[str, int, int]] = []
    for region, block_id, part, char_start, char_end in selected:
        source_region = _source_region_for_ocr_region(region.source_page_region_id or "", page_gold)
        source_text = source_region.text if source_region is not None else None
        if source_text is not None and source_text == region.gold_text:
            excluded_source_ranges.append((block_id, char_start, char_end))
            continue
        if source_text is not None and source_text.count(part) == 1:
            source_start = source_text.index(part)
            excluded_source_ranges.append((block_id, source_start, source_start + len(part)))
            continue
        # If the approved OCR text cannot be aligned exactly to the P04 source
        # Region, conservatively reserve the whole source block.
        excluded_source_ranges.append((block_id, 0, len(source_text or part)))
    semantic = _semantic_type(gold_text)
    requires_parent = semantic in {"list", "other"} or gold_text.startswith(
        ("其中", "因此", "上述", "这", "（", "(")
    )
    neighbors = (
        [
            _heading_neighbor(
                document_id=PDF_DOCUMENT_ID,
                document_version=PDF_DOCUMENT_VERSION,
                document_sha256=PDF_DOCUMENT_SHA256,
                section_number=section_number,
                heading_text=section.title,
                heading_unit_id=section.source_span.block_start or "",
                page_number=page_number,
            )
        ]
        if requires_parent
        else []
    )
    upstream_refs = [_upstream_ref("ds1_p05_ocr", ocr_record)]
    if section.review_status is ReviewStatus.APPROVED:
        upstream_refs.insert(0, _upstream_ref("ds1_p04_native_docx", section))
    if page_gold is not None:
        upstream_refs.append(_upstream_ref("ds1_p04_native_docx", page_gold))
    return (
        EvidenceRecord(
            record_id=evidence_id,
            evidence_id=evidence_id,
            candidate_source="approved_p05_source_projection",
            annotation_profile="formal_ds2_v1",
            course_id=PDF_COURSE_ID,
            source_span=SourceSpan(
                document_id=PDF_DOCUMENT_ID,
                document_version=PDF_DOCUMENT_VERSION,
                document_sha256=PDF_DOCUMENT_SHA256,
                section_path=_section_path(section_number),
                page_start=page_number,
                page_end=page_number,
                block_start=block_ids[0],
                block_end=block_ids[-1],
            ),
            source_type="ocr_derived",
            gold_text=gold_text,
            content_sha256=content_hash,
            normalized_content_sha256=_sha256_text(_normalized_text(gold_text)),
            source_units=units,
            text_assembly="ocr_regions_newline",
            upstream_record_refs=upstream_refs,
            semantic_unit_type=semantic,
            requires_parent=requires_parent,
            necessary_neighbor_text=[neighbor.text for neighbor in neighbors],
            necessary_neighbors=neighbors,
            bboxes=bboxes,
            ocr_confidence=None,
            ocr_provenance=EvidenceOCRProvenance(
                approved_ocr_record_id=ocr_record.record_id,
                approved_ocr_record_sha256=record_digest(ocr_record),
                ocr_input_document_id=ocr_record.ocr_input_document_id,
                ocr_input_document_version=ocr_record.ocr_input_document_version,
                ocr_input_document_sha256=ocr_record.ocr_input_document_sha256,
                ocr_input_page_number=ocr_record.ocr_input_page_number,
                source_page_number=page_number,
                source_region_ids=[str(region.region_id) for region, _, _, _, _ in selected],
            ),
            other_type_review_reason=(
                "Source-grounded OCR body statement does not fit a narrower frozen semantic label; "
                "course_owner must confirm or relabel it."
                if semantic == "other"
                else None
            ),
        ),
        excluded_source_ranges,
    )


def _pdf_table_record(
    table: TableGold,
    section: SectionGold,
    page_gold: PageGold,
) -> EvidenceRecord:
    table_region = next(
        region for region in page_gold.regions if region.block_id == table.source_span.block_start
    )
    if table_region.bbox is None or page_gold.page_width is None or page_gold.page_height is None:
        raise ValueError("Approved PDF TableGold has no bounded PageGold region")
    text = _table_text(table)
    piece = SourcePiece(
        source_unit_id=table.source_span.block_start or "pdf:p020:table:000",
        text=text,
        order=20_000.0,
        section_number="1.2.2",
        semantic_unit_type="table",
        assembly="table_tsv",
        table_gold=table,
    )
    return _record_from_piece(
        LocatedPiece(
            piece,
            (
                EvidenceBBox(
                    page_number=20,
                    bbox=table_region.bbox,
                    page_width=page_gold.page_width,
                    page_height=page_gold.page_height,
                    coordinate_space="pdf_points_top_left",
                    source_region_id=table_region.region_id,
                ),
            ),
        ),
        course_id=PDF_COURSE_ID,
        document_id=PDF_DOCUMENT_ID,
        document_version=PDF_DOCUMENT_VERSION,
        document_sha256=PDF_DOCUMENT_SHA256,
        section_number="1.2.2",
        heading_text=section.title,
        heading_unit_id=section.source_span.block_start or "",
        upstream=[
            _upstream_ref("ds1_p04_native_docx", section),
            _upstream_ref("ds1_p04_native_docx", table),
            _upstream_ref("ds1_p04_native_docx", page_gold),
        ],
    )


def _pdf_records(
    repository_root: Path,
    approved_p04: DS1ParsingDataset,
    approved_p05: DS1ParsingDataset,
) -> list[EvidenceRecord]:
    sections: dict[str, SectionGold] = {
        record.source_span.section_path[-1]: record
        for record in approved_p04.records
        if isinstance(record, SectionGold) and record.document_id == PDF_DOCUMENT_ID
    }
    sections.pop("1.4.1", None)
    _, _, page31_blocks = _poppler_page(repository_root, 31)
    _, _, page32_blocks = _poppler_page(repository_root, 32)
    start_block = next(block for block in page31_blocks if block.block_number == 15)
    title_block = next(block for block in page31_blocks if block.block_number == 16)
    end_block = next(block for block in page32_blocks if block.block_number == 25)
    if start_block.text.strip() != "1.3.7" or title_block.text.strip() != "AI＋电子支付":
        raise ValueError("replacement PDF Section 1.3.7 source heading changed")
    if next(block for block in page32_blocks if block.block_number == 26).text.strip() != "1.4":
        raise ValueError("replacement PDF Section 1.3.7 end boundary changed")
    sections["1.3.7"] = SectionGold(
        record_id="ds2-source-section-pdf-1-3-7",
        candidate_source="independent_poppler_boundary",
        document_id=PDF_DOCUMENT_ID,
        document_version=PDF_DOCUMENT_VERSION,
        section_id="gold-sec-doc-ai-general-education-excerpt-1-3-7",
        title="1.3.7　AI＋电子支付",
        level=3,
        parent_section_id="gold-sec-doc-ai-general-education-excerpt-1-3",
        source_span=SourceSpan(
            document_id=PDF_DOCUMENT_ID,
            document_version=PDF_DOCUMENT_VERSION,
            document_sha256=PDF_DOCUMENT_SHA256,
            section_path=["1", "1.3", "1.3.7"],
            page_start=31,
            page_end=32,
            block_start=start_block.source_unit_id,
            block_end=end_block.source_unit_id,
        ),
        first_text="1.3.7　AI＋电子支付",
        last_text=end_block.text,
    )
    if set(sections) != set(PDF_SECTIONS):
        raise ValueError("Approved P04 PDF Section set differs from frozen DS2 coverage")
    page_gold = {
        record.page_number: record
        for record in approved_p04.records
        if isinstance(record, PageGold) and record.document_id == PDF_DOCUMENT_ID
    }
    table = next(
        record
        for record in approved_p04.records
        if isinstance(record, TableGold) and record.source_span.document_id == PDF_DOCUMENT_ID
    )
    ocr_by_section: dict[str, list[EvidenceRecord]] = defaultdict(list)
    excluded_source_ranges: dict[str, list[tuple[int, int]]] = defaultdict(list)
    for page_number, section_number in OCR_PAGE_TO_SECTION.items():
        record, used_ranges = _ocr_evidence(
            page_number=page_number,
            section_number=section_number,
            section=sections[section_number],
            ocr_record=_ocr_record_for_page(approved_p05, page_number),
            page_gold=page_gold.get(page_number),
        )
        ocr_by_section[section_number].append(record)
        for block_id, char_start, char_end in used_ranges:
            excluded_source_ranges[block_id].append((char_start, char_end))

    records: list[EvidenceRecord] = []
    for section_number in PDF_SECTIONS:
        section = sections[section_number]
        mandatory_table = (
            _pdf_table_record(table, section, page_gold[20]) if section_number == "1.2.2" else None
        )
        native_needed = 5 - len(ocr_by_section[section_number])
        if mandatory_table is not None:
            native_needed -= 1
        candidates = [
            unit
            for unit in _pdf_complete_units(_blocks_for_section(repository_root, section))
            if not any(
                unit_source.char_start is not None
                and unit_source.char_end is not None
                and unit_source.char_start < excluded_end
                and excluded_start < unit_source.char_end
                for unit_source in unit.source_units
                for excluded_start, excluded_end in excluded_source_ranges.get(
                    unit_source.source_unit_id, []
                )
            )
        ]
        selected = [candidates[index] for index in _even_indexes(len(candidates), native_needed)]
        section_records: list[EvidenceRecord] = []
        for unit in selected:
            refs: list[UpstreamApprovedRecordRef] = []
            if section.review_status is ReviewStatus.APPROVED:
                refs.append(_upstream_ref("ds1_p04_native_docx", section))
            source_page_gold = page_gold.get(unit.blocks[0].page_number)
            if source_page_gold is not None:
                refs.append(_upstream_ref("ds1_p04_native_docx", source_page_gold))
            section_records.append(
                _pdf_native_record(
                    unit,
                    section_number=section_number,
                    section=section,
                    upstream=refs,
                )
            )
        if mandatory_table is not None:
            section_records.append(mandatory_table)
        section_records.extend(ocr_by_section[section_number])
        section_records.sort(
            key=lambda record: (
                record.source_span.page_start or 0,
                record.source_span.block_start or "",
                record.record_id,
            )
        )
        if len(section_records) != 5:
            raise AssertionError(
                f"PDF section {section_number} did not produce five Evidence units"
            )
        records.extend(section_records)
    if len(records) != 40:
        raise AssertionError(f"expected 40 PDF Evidence records, received {len(records)}")
    if Counter(record.source_type for record in records) != {"native_text": 30, "ocr_derived": 10}:
        raise AssertionError("PDF native/OCR-derived distribution differs from 30/10")
    return records


def _validate_batch(records: Sequence[EvidenceRecord]) -> tuple[dict[str, int], list[str]]:
    if len(records) != 120 or len({record.record_id for record in records}) != 120:
        raise ValueError("formal DS2 Candidate must contain 120 unique records")
    document_counts = Counter(record.source_span.document_id for record in records)
    if document_counts != {DOCX_DOCUMENT_ID: 80, PDF_DOCUMENT_ID: 40}:
        raise ValueError(f"formal DS2 document distribution differs: {document_counts}")
    section_counts = Counter(
        f"{record.source_span.document_id}:{record.source_span.section_path[-1]}"
        for record in records
    )
    if len(section_counts) != 24 or set(section_counts.values()) != {5}:
        raise ValueError(f"formal DS2 Section distribution differs: {section_counts}")
    full_table_gold_records = [
        record
        for record in records
        if len(record.source_units) == 1
        and re.fullmatch(
            r"table:\d+(?::rendered-visible-algorithm)?",
            record.source_units[0].source_unit_id,
        )
    ]
    if len(full_table_gold_records) != 10:
        raise ValueError("formal DS2 must contain exactly ten complete TableGold-derived records")
    if sum(record.semantic_unit_type == "table" for record in full_table_gold_records) not in {
        9,
        10,
    }:
        raise ValueError("formal DS2 TableGold-derived semantic distribution is invalid")
    ocr_pages = sorted(
        record.ocr_provenance.source_page_number
        for record in records
        if record.ocr_provenance is not None
    )
    if ocr_pages != sorted(OCR_PAGE_TO_SECTION):
        raise ValueError("formal DS2 OCR parent pages differ from the frozen ten-page selection")
    source_ranges = [
        (
            record.source_span.document_id,
            record.source_span.block_start,
            record.source_span.block_end,
            record.source_span.char_start,
            record.source_span.char_end,
        )
        for record in records
    ]
    if len(source_ranges) != len(set(source_ranges)):
        raise ValueError("formal DS2 contains duplicate exact source ranges")
    normalized: dict[str, list[str]] = defaultdict(list)
    for record in records:
        if record.normalized_content_sha256 is None:
            raise ValueError("formal DS2 record has no normalized content Hash")
        normalized[record.normalized_content_sha256].append(record.record_id)
    duplicate_ids = sorted(
        record_id for ids in normalized.values() if len(ids) > 1 for record_id in ids
    )
    return dict(sorted(section_counts.items())), duplicate_ids


def _review_image(
    record: EvidenceRecord,
    *,
    pdf_document: fitz.Document,
    docx_rendered: fitz.Document,
    output: Path,
) -> None:
    source = docx_rendered if record.source_span.document_id == DOCX_DOCUMENT_ID else pdf_document
    by_page: dict[int, list[EvidenceBBox]] = defaultdict(list)
    for bbox in record.bboxes:
        by_page[bbox.page_number].append(bbox)
    images: list[Image.Image] = []
    zoom = 1.35
    for page_number, boxes in sorted(by_page.items()):
        page = source[page_number - 1]
        pixmap = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
        draw = ImageDraw.Draw(image)
        for box in boxes:
            x0, y0, x1, y1 = box.bbox
            draw.rectangle(
                (x0 * zoom, y0 * zoom, x1 * zoom, y1 * zoom),
                outline=(220, 38, 38),
                width=4,
            )
        images.append(image)
    if not images:
        raise ValueError(f"review record has no renderable BBox: {record.record_id}")
    width = max(image.width for image in images)
    height = sum(image.height for image in images) + 16 * (len(images) - 1)
    montage = Image.new("RGB", (width, height), "white")
    y = 0
    for image in images:
        montage.paste(image, (0, y))
        y += image.height + 16
    output.parent.mkdir(parents=True, exist_ok=True)
    montage.save(output, format="PNG", optimize=False, compress_level=6)


def _record_card(record: EvidenceRecord, image_name: str) -> str:
    refs = "<br>".join(
        html.escape(f"{ref.dataset_component}: {ref.record_id} @ {ref.record_sha256}")
        for ref in record.upstream_record_refs
    )
    units = "<br>".join(
        html.escape(
            f"{unit.source_unit_id} chars={unit.char_start}:{unit.char_end} sha={unit.text_sha256}"
        )
        for unit in record.source_units
    )
    neighbors = (
        "<br>".join(
            html.escape(f"{neighbor.relation}: {neighbor.text}")
            for neighbor in record.necessary_neighbors
        )
        or "无"
    )
    payload = html.escape(
        json.dumps(record.model_dump(mode="json"), ensure_ascii=False, indent=2, sort_keys=True)
    )
    return f"""
<article id="{record.record_id}">
  <h2>{record.record_id}</h2>
  <p><b>课程/Section：</b>{html.escape(record.course_id or "")} / {html.escape(".".join(record.source_span.section_path))}<br>
  <b>来源/类型：</b>{record.source_type} / {record.semantic_unit_type}<br>
  <b>Span：</b>{html.escape(str(record.source_span.block_start))} → {html.escape(str(record.source_span.block_end))}, pages {record.source_span.page_start}-{record.source_span.page_end}<br>
  <b>组装：</b>{record.text_assembly}；<b>需要父上下文：</b>{str(record.requires_parent).lower()}</p>
  <img src="assets/{html.escape(image_name)}" alt="{record.record_id} BBox overlay">
  <h3>Gold 原文</h3><pre>{html.escape(record.gold_text)}</pre>
  <p><b>Source Units</b><br>{units}</p>
  <p><b>必要邻接</b><br>{neighbors}</p>
  <p><b>上游 Approved 记录</b><br>{refs}</p>
  <details><summary>完整 Gold JSON</summary><pre>{payload}</pre></details>
  <label><input type="radio" name="{record.record_id}" value="approve">通过</label>
  <label><input type="radio" name="{record.record_id}" value="return">退回</label>
  <input class="note" data-id="{record.record_id}" placeholder="审核备注">
</article>"""


def _group_html(
    *,
    candidate_sha256: str,
    group: DS2ReviewGroup,
    records: Sequence[EvidenceRecord],
    image_names: dict[str, str],
) -> str:
    cards = "\n".join(_record_card(record, image_names[record.record_id]) for record in records)
    return f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>DS2 {group.group_id}</title><style>
body{{font-family:system-ui,sans-serif;margin:0;background:#f4f6f8;color:#17202a}}header,main{{max-width:1120px;margin:auto;padding:20px}}
header{{background:#17202a;color:white;max-width:none}}header>div{{max-width:1120px;margin:auto}}article{{background:white;border:1px solid #cbd5e1;border-radius:10px;padding:18px;margin:18px 0}}
img{{max-width:100%;border:1px solid #94a3b8}}pre{{white-space:pre-wrap;word-break:break-word;background:#f8fafc;padding:12px;overflow:auto}}
.note{{width:60%;margin-left:12px}}nav a{{color:#93c5fd;margin-right:18px}}
</style></head><body><header><div><h1>{html.escape(group.title)}</h1>
<p>Candidate SHA-256: <code>{candidate_sha256}</code>；本组 30 条，只记录审核决定，不代表批准 Gold。</p>
<nav><a href="index.html">总览</a><a href="group-01.html">第1组</a><a href="group-02.html">第2组</a><a href="group-03.html">第3组</a><a href="group-04.html">第4组</a></nav></div></header>
<main>{cards}<button id="export">导出本组审核 JSON</button></main>
<script>document.getElementById('export').onclick=()=>{{const reviewed=[],returned=[],notes={{}};for(const id of {json.dumps(group.record_ids, ensure_ascii=False)}){{const value=document.querySelector(`input[name="${{id}}"]:checked`)?.value; if(value==='approve')reviewed.push(id);if(value==='return')returned.push(id);const note=document.querySelector(`.note[data-id="${{id}}"]`)?.value||'';if(note)notes[id]=note;}}const data={{schema_version:'courserag.ds2-review-group.v1',candidate_file_sha256:'{candidate_sha256}',group_id:'{group.group_id}',reviewed_record_ids:reviewed,returned_record_ids:returned,record_notes:notes}};const a=document.createElement('a');a.href=URL.createObjectURL(new Blob([JSON.stringify(data,null,2)],{{type:'application/json'}}));a.download='ds2_{group.group_id}_review.json';a.click();}};</script></body></html>"""


def _review_pack(
    repository_root: Path,
    records: Sequence[EvidenceRecord],
    groups: Sequence[DS2ReviewGroup],
    candidate_sha256: str,
) -> tuple[Path, str, dict[str, str]]:
    review_dir = repository_root / "storage_eval/ds2_p06_review" / candidate_sha256
    assets = review_dir / "assets"
    assets.mkdir(parents=True, exist_ok=True)
    image_names: dict[str, str] = {}
    image_hashes: dict[str, str] = {}
    with (
        fitz.open(repository_root / PRIMARY_PDF) as pdf_document,
        fitz.open(repository_root / DOCX_RENDER_RUNS[0]) as docx_rendered,
    ):
        for record in records:
            filename = f"{record.record_id}.png"
            output = assets / filename
            _review_image(
                record,
                pdf_document=pdf_document,
                docx_rendered=docx_rendered,
                output=output,
            )
            image_names[record.record_id] = filename
            image_hashes[record.record_id] = sha256_file(output)
    by_id = {record.record_id: record for record in records}
    for group in groups:
        group_records = [by_id[record_id] for record_id in group.record_ids]
        atomic_write_text(
            review_dir / f"{group.group_id}.html",
            _group_html(
                candidate_sha256=candidate_sha256,
                group=group,
                records=group_records,
                image_names=image_names,
            ),
        )
    links = "\n".join(
        f'<li><a href="{group.group_id}.html">{html.escape(group.title)}</a>（30 条）</li>'
        for group in groups
    )
    index = f"""<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><title>正式 DS2 Candidate 审核</title>
<style>body{{font-family:system-ui,sans-serif;max-width:900px;margin:40px auto;line-height:1.65}}code{{word-break:break-all}}li{{margin:14px 0}}</style></head><body>
<h1>P06 前正式 DS2 Candidate 审核</h1><p><b>Candidate SHA-256：</b><code>{candidate_sha256}</code></p>
<p>共 120 条；必须审核全部四组。重点检查逐字原文、语义完整性、Section/页码/源单元、BBox、上下文依赖、OCR 来源以及十张表。</p>
<ol>{links}</ol><p>四组合并后，只有全部 120 条完成且无退回项，才能使用精确 SHA-256 批准。</p></body></html>"""
    atomic_write_text(review_dir / "index.html", index)
    return review_dir, sha256_file(review_dir / "index.html"), image_hashes


def _manifest_source_artifacts(repository_root: Path) -> list[HashedArtifact]:
    return [
        _artifact(repository_root, PRIMARY_PDF, "application/pdf"),
        _artifact(
            repository_root,
            PRIMARY_DOCX,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ),
        _artifact(repository_root, P04_APPROVED, "application/json"),
        _artifact(repository_root, P05_APPROVED, "application/json"),
        _artifact(repository_root, P04_APPROVAL, "application/json"),
        _artifact(repository_root, P05_APPROVAL, "application/json"),
        _artifact(repository_root, P04_MANIFEST, "application/json"),
        _artifact(repository_root, P05_MANIFEST, "application/json"),
        _artifact(repository_root, RENDERER_PROFILE_PATH, "application/json"),
        _artifact(repository_root, RENDERER_FONT_LOCK_PATH, "application/json"),
        _artifact(repository_root, DOCX_RENDER_RUNS[0], "application/pdf"),
        _artifact(repository_root, DOCX_RENDER_RUNS[1], "application/pdf"),
    ]


def _update_candidate_governance(dataset_root: Path) -> None:
    manifest_path = dataset_root / "manifest.json"
    values = json.loads(manifest_path.read_text(encoding="utf-8"))
    if values.get("gold_status") not in {
        "ds1_p04_native_docx_approved",
        "ds1_p05_ocr_approved",
    }:
        raise ValueError("unexpected pre-DS2 Gold status")
    values["gold_components"] = {
        "ds0": "approved",
        "ds1_native_docx": "approved",
        "ds1_ocr": "approved",
        "ds2": "candidate_pending_course_owner",
    }
    phase = values.setdefault("phase_input_status", {})
    phase["p05"] = "completed_gate_passed"
    phase["p06"] = "candidate_pending_course_owner"
    atomic_write_json(manifest_path, values)


def build_ds2_candidate(repository_root: Path) -> dict[str, object]:
    repository_root = repository_root.resolve()
    dataset_root = repository_root / DATASET_ROOT
    if sha256_file(repository_root / PRIMARY_PDF) != PDF_DOCUMENT_SHA256:
        raise ValueError("primary PDF Hash differs from Approved DS0")
    if sha256_file(repository_root / PRIMARY_DOCX) != DOCX_DOCUMENT_SHA256:
        raise ValueError("primary DOCX Hash differs from Approved DS0")
    approved_p04 = DS1ParsingDataset.model_validate_json(
        (repository_root / P04_APPROVED).read_text(encoding="utf-8")
    )
    approved_p05 = DS1ParsingDataset.model_validate_json(
        (repository_root / P05_APPROVED).read_text(encoding="utf-8")
    )
    if any(record.review_status is not ReviewStatus.APPROVED for record in approved_p04.records):
        raise ValueError("formal DS2 requires fully Approved P04 DS1")
    if any(record.review_status is not ReviewStatus.APPROVED for record in approved_p05.records):
        raise ValueError("formal DS2 requires fully Approved P05 OCR Gold")
    p04_manifest = DS1CandidateManifest.model_validate_json(
        (repository_root / P04_MANIFEST).read_text(encoding="utf-8")
    )
    profile = RendererProfile.load(repository_root / RENDERER_PROFILE_PATH)

    docx_records, canonical_docx_sha256 = _docx_records(repository_root, approved_p04, p04_manifest)
    pdf_records = _pdf_records(repository_root, approved_p04, approved_p05)
    records = [*docx_records, *pdf_records]
    section_counts, duplicate_ids = _validate_batch(records)
    dataset = DS2EvidenceDataset(
        dataset_id="courserag-eval",
        dataset_version="v1",
        evidence=records,
    )
    candidate_path = repository_root / CANDIDATE_PATH
    r1_path = repository_root / R1_CANDIDATE_PATH
    r2_path = repository_root / R2_CANDIDATE_PATH
    r3_path = repository_root / R3_CANDIDATE_PATH
    predecessor_path = repository_root / PREDECESSOR_CANDIDATE_PATH
    if (
        not r1_path.is_file()
        or not r2_path.is_file()
        or not r3_path.is_file()
        or not predecessor_path.is_file()
    ):
        raise FileNotFoundError("pre-delivery DS2 r1/r2/r3/r4 audit artifacts are missing")
    r1_hash = sha256_file(r1_path)
    r2_hash = sha256_file(r2_path)
    r3_hash = sha256_file(r3_path)
    predecessor_hash = sha256_file(predecessor_path)
    candidate_text = (
        json.dumps(dataset.model_dump(mode="json"), ensure_ascii=False, indent=2, sort_keys=True)
        + "\n"
    )
    if candidate_path.exists() and candidate_path.read_text(encoding="utf-8") != candidate_text:
        raise ValueError("DS2 r5 Candidate is immutable; create r6 instead of overwriting it")
    atomic_write_text(candidate_path, candidate_text)
    candidate_sha256 = sha256_file(candidate_path)
    record_hashes = {record.record_id: record_digest(record) for record in records}

    groups = [
        DS2ReviewGroup(
            group_id=f"group-{index + 1:02d}",
            title=title,
            record_ids=[r.record_id for r in records[index * 30 : (index + 1) * 30]],
        )
        for index, title in enumerate(
            (
                "第 1 组：DOCX 前 6 个 Section",
                "第 2 组：DOCX 中间 6 个 Section",
                "第 3 组：DOCX 后 4 个 Section + PDF 前 2 个 Section",
                "第 4 组：PDF 后 6 个 Section",
            )
        )
    ]
    review_dir, review_index_hash, review_asset_hashes = _review_pack(
        repository_root, records, groups, candidate_sha256
    )
    type_counts = Counter(record.source_type for record in records)
    semantic_counts = Counter(record.semantic_unit_type for record in records)
    manifest = DS2CandidateManifest(
        dataset_id="courserag-eval",
        dataset_version="v1",
        candidate_revision=5,
        candidate_relative_path=_relative(repository_root, candidate_path),
        candidate_file_sha256=candidate_sha256,
        candidate_record_sha256=record_hashes,
        source_artifacts=_manifest_source_artifacts(repository_root),
        upstream_approved_file_sha256={
            "ds1_p04_native_docx": sha256_file(repository_root / P04_APPROVED),
            "ds1_p05_ocr": sha256_file(repository_root / P05_APPROVED),
        },
        record_counts={
            "total": 120,
            "docx": 80,
            "pdf": 40,
            "native_text": type_counts["native_text"],
            "ocr_derived": type_counts["ocr_derived"],
            "table": semantic_counts["table"],
        },
        section_record_counts=section_counts,
        semantic_type_counts=dict(sorted(semantic_counts.items())),
        review_groups=groups,
        review_pack_relative_path=_relative(repository_root, review_dir / "index.html"),
        review_pack_index_sha256=review_index_hash,
        review_asset_sha256=review_asset_hashes,
        renderer_profile={
            "provider": profile.provider,
            "renderer_version": profile.renderer_version,
            "font_manifest_sha256": profile.font_manifest_sha256,
            "profile_sha256": profile.profile_sha256,
        },
        canonical_rendered_pdf_sha256=canonical_docx_sha256,
        stable_id_profile_sha256=_canonical_hash(STABLE_ID_PROFILE),
        normalization_profile_sha256=_canonical_hash(NORMALIZATION_PROFILE),
        predecessor_candidate_file_sha256=predecessor_hash,
        duplicate_review_record_ids=duplicate_ids,
        constraints=[
            "Candidate only; exact course_owner approval of the full file SHA-256 is required.",
            "Only two independent semantic sources exist; derived OCR fixtures are provenance only.",
            "No P06 Evidence Builder, Chunker, Retriever, LLM Judge or system prediction defined Gold.",
            "All 120 records require human review across four fixed groups of 30.",
            "OCR confidence is intentionally null in Gold; prediction confidence belongs to run output.",
            "DOCX Table 20 uses the owner-approved 7x4 rendered-visible grid.",
            "PDF Section 1.4.1 was source-incomplete and was owner-approved for replacement by the source-complete 1.3.7 AI＋电子支付 Section.",
            "Dev/Test remain empty and Test remains unlocked.",
        ],
    )
    manifest_path = repository_root / CANDIDATE_MANIFEST_PATH
    atomic_write_json(manifest_path, manifest.model_dump(mode="json"))
    history = CandidateRevisionHistory(
        dataset_id="courserag-eval",
        dataset_version="v1",
        revisions=[
            CandidateRevisionArtifact(
                revision=1,
                candidate_relative_path=_relative(repository_root, r1_path),
                candidate_file_sha256=r1_hash,
                status="superseded",
                reason=(
                    "Pre-delivery automated audit found that native PDF line fragments were not "
                    "complete semantic units; r1 was never delivered for human review."
                ),
            ),
            CandidateRevisionArtifact(
                revision=2,
                candidate_relative_path=_relative(repository_root, r2_path),
                candidate_file_sha256=r2_hash,
                status="superseded",
                reason=(
                    "Pre-delivery audit found two OCR records joined an atomic heading/list to a "
                    "body sentence; r2 was never delivered for human review."
                ),
            ),
            CandidateRevisionArtifact(
                revision=3,
                candidate_relative_path=_relative(repository_root, r3_path),
                candidate_file_sha256=r3_hash,
                status="superseded",
                reason=(
                    "Pre-delivery visual audit found three native PDF object sequences that were "
                    "punctuated but not complete readable semantic units; r3 was never delivered "
                    "for human review."
                ),
            ),
            CandidateRevisionArtifact(
                revision=4,
                candidate_relative_path=_relative(repository_root, predecessor_path),
                candidate_file_sha256=predecessor_hash,
                status="superseded",
                reason=(
                    "Pre-delivery visual audit found DOCX table BBoxes covered matched text but not "
                    "the complete rendered table rules; r4 was never delivered for human review."
                ),
            ),
            CandidateRevisionArtifact(
                revision=5,
                candidate_relative_path=_relative(repository_root, candidate_path),
                candidate_file_sha256=candidate_sha256,
                status="pending_course_owner_review",
                reason=(
                    "Source-grounded r5 expands DOCX table BBoxes to fixed-renderer vector rules, "
                    "excludes exact audited layout anomalies without rewriting them, and applies "
                    "the owner-approved 1.3.7 replacement."
                ),
            ),
        ],
    )
    history_path = repository_root / REVISION_HISTORY_PATH
    if history_path.exists():
        existing = CandidateRevisionHistory.model_validate_json(
            history_path.read_text(encoding="utf-8")
        )
        allowed_pre_delivery_r4 = (
            len(existing.revisions) == 4
            and existing.revisions[0].revision == 1
            and existing.revisions[0].candidate_relative_path == _relative(repository_root, r1_path)
            and existing.revisions[0].candidate_file_sha256 == r1_hash
            and existing.revisions[0].status == "superseded"
            and existing.revisions[1].revision == 2
            and existing.revisions[1].candidate_relative_path == _relative(repository_root, r2_path)
            and existing.revisions[1].candidate_file_sha256 == r2_hash
            and existing.revisions[1].status == "superseded"
            and existing.revisions[2].revision == 3
            and existing.revisions[2].candidate_relative_path == _relative(repository_root, r3_path)
            and existing.revisions[2].candidate_file_sha256 == r3_hash
            and existing.revisions[2].status == "superseded"
            and existing.revisions[3].revision == 4
            and existing.revisions[3].candidate_relative_path
            == _relative(repository_root, predecessor_path)
            and existing.revisions[3].candidate_file_sha256 == predecessor_hash
            and existing.revisions[3].status == "pending_course_owner_review"
        )
        if existing != history and not allowed_pre_delivery_r4:
            raise ValueError("existing DS2 Candidate revision history differs from r5 lineage")
    atomic_write_json(history_path, history.model_dump(mode="json"))
    _update_candidate_governance(dataset_root)
    report = f"""# ED-PRE06 正式 DS2 Candidate 审核报告

- Candidate：`{_relative(repository_root, candidate_path)}`
- Candidate SHA-256：`{candidate_sha256}`
- 记录：120（DOCX 80 / PDF 40；PDF native 30 / OCR-derived 10）
- Section：DOCX 16 × 5；PDF 8 × 5
- 完整表格 Evidence：10
- 审核入口：`{_relative(repository_root, review_dir / "index.html")}`
- Manifest：`{_relative(repository_root, manifest_path)}`

## 审批边界

本批全部记录仍为 Candidate，`approval=null`。未创建 `approved/ds2`，未运行 P06 B1/B2，
Dev/Test 未写入且 Test 未锁定。只有回复
`批准正式 DS2 Candidate {candidate_sha256}` 才能执行精确 Hash 提升。

## 四组人工审核

1. DOCX 前 6 个 Section，共 30 条。
2. DOCX 中间 6 个 Section，共 30 条。
3. DOCX 后 4 个 Section与 PDF 前 2 个 Section，共 30 条。
4. PDF 后 6 个 Section，共 30 条。

逐条检查原文完整性、语义类型、Section/Span、页码与红框、上下文依赖、OCR 父页和
Region、十张表的矩形网格。规范化重复待复核记录数：{len(duplicate_ids)}。
"""
    atomic_write_text(repository_root / PHASE_REPORT_PATH, report)
    return {
        "candidate_file_sha256": candidate_sha256,
        "candidate_record_count": len(records),
        "candidate_manifest_sha256": sha256_file(manifest_path),
        "review_pack": _relative(repository_root, review_dir / "index.html"),
        "review_pack_sha256": review_index_hash,
        "review_groups": [f"{group.group_id}.html" for group in groups],
        "duplicate_review_record_count": len(duplicate_ids),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the source-grounded P06 DS2 Candidate.")
    parser.add_argument("--repository-root", type=Path, default=Path("."))
    args = parser.parse_args()
    result = build_ds2_candidate(args.repository_root)
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))


if __name__ == "__main__":
    main()
