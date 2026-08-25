"""Export and render the successful P18 Dev CP-B10 artifacts without model calls."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from docx import Document

from coursepilot.domain.exam import ExamArtifact
from coursepilot.domain.lesson import LessonArtifact
from coursepilot.domain.ppt import PPTArtifact
from coursepilot.exporters.exam_docx_exporter import ExamDocxExporter
from coursepilot.exporters.exam_versioned_exporter import ExamVersionedExporter
from coursepilot.exporters.lesson_docx_exporter import LessonDocxExporter
from coursepilot.exporters.lesson_versioned_exporter import VersionedLessonDocxExporter
from coursepilot.exporters.pptx import PPTXVersionedExporter
from coursepilot.exporters.pptx_exporter import PPTXExporter
from coursepilot.rendering.pptx import render_with_libreoffice
from coursepilot.schemas.exam_schema import ExamBlueprintContent
from coursepilot.schemas.lesson_schema import LessonDesignContent
from coursepilot.schemas.ppt_schema import SlideOutlineContent
from coursepilot.schemas.question_schema import QuestionItem
from evaluation.p16_closure import _profile


def run_exports(
    *,
    repository_root: Path,
    rows_path: Path,
    output_dir: Path,
    tracks: tuple[str, ...] = ("cp_b10",),
) -> dict[str, Any]:
    del repository_root
    rows = json.loads(rows_path.read_text(encoding="utf-8"))
    output_dir.mkdir(parents=True, exist_ok=True)
    results: list[dict[str, Any]] = []
    for row in rows:
        if row.get("track") not in tracks:
            continue
        record_id = str(row["record_id"])
        track = str(row["track"])
        artifact_dir = output_dir / track / record_id
        if row.get("status") != "succeeded" or not isinstance(row.get("artifact"), dict):
            results.append(
                {
                    "record_id": record_id,
                    "track": track,
                    "component": row.get("component"),
                    "passed": False,
                    "reason": "generation_failed",
                }
            )
            continue
        warnings: list[str]
        try:
            if row["component"] == "cp_ds1" and track == "cp_b0":
                lesson = LessonDesignContent.model_validate(row["artifact"])
                target = artifact_dir / "lesson.docx"
                LessonDocxExporter().export(lesson, target)
                Document(str(target))
                files = [str(target)]
                rendered = True
                warnings = []
            elif row["component"] == "cp_ds1":
                lesson_artifact = LessonArtifact.model_validate(row["artifact"])
                target = artifact_dir / "lesson.docx"
                VersionedLessonDocxExporter().export(lesson_artifact, target)
                Document(str(target))
                files = [str(target)]
                rendered = True
                warnings = []
            elif row["component"] == "cp_ds2" and track == "cp_b0":
                blueprint = ExamBlueprintContent.model_validate(row["artifact"]["blueprint"])
                questions = [
                    QuestionItem.model_validate(item) for item in row["artifact"]["questions"]
                ]
                exporter = ExamDocxExporter()
                targets = {
                    "student_exam": exporter.export_student_exam(
                        blueprint, questions, artifact_dir / "student_exam.docx"
                    ),
                    "answer_key": exporter.export_teacher_answer(
                        blueprint, questions, artifact_dir / "answer_key.docx"
                    ),
                    "detailed_explanation": exporter.export_explanation(
                        blueprint, questions, artifact_dir / "detailed_explanation.docx"
                    ),
                    "answer_sheet": exporter.export_answer_sheet(
                        blueprint, questions, artifact_dir / "answer_sheet.docx"
                    ),
                }
                for target in targets.values():
                    Document(str(target))
                files = [str(item) for item in targets.values()]
                rendered = True
                warnings = []
            elif row["component"] == "cp_ds2":
                exam_artifact = ExamArtifact.model_validate(row["artifact"])
                # Dev export verification exercises the already separate,
                # explicit Global Review approval boundary; it does not mutate
                # the generated Artifact or production defaults.
                exam_artifact = exam_artifact.model_copy(update={"global_review_approved": True})
                targets = ExamVersionedExporter().export(exam_artifact, artifact_dir)
                for target in targets.values():
                    Document(str(target))
                files = [str(item) for item in targets.values()]
                rendered = True
                warnings = []
            elif track == "cp_b0":
                outline = SlideOutlineContent.model_validate(row["artifact"])
                target = artifact_dir / "slides.pptx"
                PPTXExporter().export(outline, target)
                render_report = render_with_libreoffice(
                    target,
                    artifact_dir / "rendered",
                    expected_slide_count=len(outline.slides),
                )
                files = [str(target)]
                rendered = render_report.rendered
                warnings = render_report.warnings
            else:
                ppt_artifact = PPTArtifact.model_validate(row["artifact"])
                resource, profile, template_id = _profile(ppt_artifact.template_id)
                architecture = ppt_artifact.architecture.model_copy(
                    update={"template_id": template_id}
                )
                ppt_artifact = ppt_artifact.model_copy(
                    update={"template_id": template_id, "architecture": architecture}
                )
                target = artifact_dir / "slides.pptx"
                PPTXVersionedExporter().export(
                    ppt_artifact, target, template_path=resource, profile=profile
                )
                render_report = render_with_libreoffice(
                    target,
                    artifact_dir / "rendered",
                    expected_slide_count=len(ppt_artifact.slides),
                )
                files = [str(target)]
                rendered = render_report.rendered
                warnings = render_report.warnings
            results.append(
                {
                    "record_id": record_id,
                    "track": track,
                    "component": row["component"],
                    "files": files,
                    "rendered": rendered,
                    "warnings": warnings,
                    "passed": rendered,
                }
            )
        except Exception as exc:
            results.append(
                {
                    "record_id": record_id,
                    "track": track,
                    "component": row.get("component"),
                    "passed": False,
                    "reason": f"{type(exc).__name__}:{exc}",
                }
            )
    result_report = {
        "schema_version": "coursepilot.p18-export-smoke.v2",
        "test_access": tracks != ("cp_b10",),
        "external_provider_calls": 0,
        "expected_artifacts": len(results),
        "executed": len(results),
        "passed": sum(bool(item["passed"]) for item in results),
        "failed": sum(not bool(item["passed"]) for item in results),
        "results": results,
    }
    _atomic_json(output_dir / "report.json", result_report)
    return result_report


def _atomic_json(path: Path, payload: dict[str, Any]) -> None:
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--repository-root", type=Path, default=Path.cwd())
    parser.add_argument("--rows", type=Path, default=Path("storage_eval/p18/dev_real_r2/rows.json"))
    parser.add_argument("--output-dir", type=Path, default=Path("storage_eval/p18/dev_exports"))
    parser.add_argument(
        "--tracks",
        default="cp_b10",
        help="Comma-separated tracks to export (for example cp_b0,cp_b10)",
    )
    args = parser.parse_args()
    report = run_exports(
        repository_root=args.repository_root.resolve(),
        rows_path=args.rows.resolve(),
        output_dir=args.output_dir.resolve(),
        tracks=tuple(item.strip() for item in args.tracks.split(",") if item.strip()),
    )
    print(json.dumps({key: value for key, value in report.items() if key != "results"}))


if __name__ == "__main__":
    main()
