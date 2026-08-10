"""Formal P04-versus-B0 scoring on the exact human-approved DS1 batch.

This module consumes Approved Gold only.  It never writes predictions into the
dataset tree and never uses either parser output to create or modify Gold.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import unicodedata
from dataclasses import dataclass
from datetime import UTC, datetime
from functools import partial
from importlib.metadata import version
from pathlib import Path
from typing import cast

from docx import Document as OpenXmlDocument
from docx.oxml.ns import qn
from pydantic import JsonValue

from coursepilot.rag.parsers.docx_parser import DOCXParser
from coursepilot.rag.parsers.pdf_parser import PDFParser
from coursepilot.rag.types import ParsedDocument
from courserag.domain.document import (
    BlockIR,
    DocxPageAnchor,
    ParsedDocumentIR,
    SectionIR,
    TableRecord,
)
from courserag.evals.parsing_metrics import (
    HeadingObservation,
    binary_detection_metrics,
    exact_mapping_accuracy,
    heading_metrics,
    parsing_success_rate,
)
from courserag.evals.schemas import (
    DocxPaginationGold,
    DS1BatchApproval,
    DS1ParsingDataset,
    PageGold,
    SectionGold,
    TableGold,
)
from courserag.parsers.docx import StructuredDOCXParser
from courserag.parsers.pagination import DocxPaginationRenderer, RendererProfile, align_docx_blocks
from courserag.parsers.structure import enrich_document_structure
from evaluation.contracts import (
    DatasetSplit,
    FallbackPolicy,
    HashedArtifact,
    MetricResult,
    ReviewStatus,
    RunIntent,
    TestLock,
)
from evaluation.corpus_fixtures import sha256_file
from evaluation.io import atomic_write_json
from evaluation.manifest import RunDatasetRef, RunManifest
from evaluation.p04_pilot import (
    DATASET_ROOT,
    FIXTURE_MANIFEST,
    PRIMARY_PDF_RUNS,
    RENDERER_FONT_LOCK,
    RENDERER_PROFILE,
    STRESS_PDF_RUNS,
    WORK_PACKAGE,
    _git_state,
    _load_inputs,
    _parse_pdf,
    _requested_fonts,
)
from evaluation.runner import EvaluationRunner

APPROVED_DS1 = DATASET_ROOT / "approved/ds1/p04_native_docx.json"
DS1_APPROVAL = DATASET_ROOT / "provenance/ds1_p04_approval.json"
DATASET_MANIFEST = DATASET_ROOT / "manifest.json"
PILOT_SPLIT = DATASET_ROOT / "splits/pilot_ids.txt"
TEST_LOCK = DATASET_ROOT / "test.lock.json"

PDF_DOCUMENT_ID = "doc_ai_general_education_excerpt"
PRIMARY_DOCX_ID = "doc_ai_algorithms_systems"
STRESS_DOCX_ID = "doc_ai_algorithms_systems_structure_stress"
EVALUATED_DOCUMENT_IDS = (PDF_DOCUMENT_ID, PRIMARY_DOCX_ID, STRESS_DOCX_ID)

_DOCX_UNIT = re.compile(r"^docx:paragraph:(?P<index>\d+)$")
_GOLD_DOCX_UNIT = re.compile(r"^paragraph:(?P<index>\d+)$")
_GOLD_BOOKMARK_UNIT = re.compile(r"^bookmark:src_p_(?P<index>\d{6})$")
_GOLD_TABLE = re.compile(r"^table:(?P<index>\d+)$")


@dataclass(frozen=True)
class GoldTextSample:
    document_id: str
    text: str
    page_number: int | None = None


@dataclass(frozen=True)
class DocxSourceUnitReference:
    paragraph_index: int
    content_sha256: str


@dataclass(frozen=True)
class PaginationResolution:
    source_unit_id: str
    parser_paragraph_index: int
    source_reference_hash_match: bool
    parser_block_hash_match: bool
    block: BlockIR | None


@dataclass(frozen=True)
class ParseOutputs:
    p04: dict[str, ParsedDocumentIR | None]
    b0: dict[str, ParsedDocument | None]
    p04_success: dict[str, bool]
    b0_success: dict[str, bool]
    renderer_repeat_match: dict[str, bool]
    docx_source_unit_references: dict[str, dict[str, DocxSourceUnitReference]]


def _artifact(repository_root: Path, path: Path, media_type: str) -> HashedArtifact:
    resolved = (repository_root / path).resolve()
    if not resolved.is_file() or not resolved.is_relative_to(repository_root):
        raise ValueError(f"formal P04 input is missing or outside the repository: {path}")
    return HashedArtifact(
        path=path.as_posix(),
        sha256=sha256_file(resolved),
        size_bytes=resolved.stat().st_size,
        media_type=media_type,
    )


def _normalize(value: str) -> str:
    return " ".join(unicodedata.normalize("NFKC", value).casefold().split())


def _normalize_sample(value: str) -> str:
    return "".join(
        character
        for character in unicodedata.normalize("NFKC", value).casefold()
        if not character.isspace()
    )


def _docx_source_unit_references(content: bytes) -> dict[str, DocxSourceUnitReference]:
    """Resolve stable OOXML source units to parser paragraph ordinals.

    Structure-stress fixtures contain inserted section-break paragraphs. Their
    approved bookmark identities therefore cannot be treated as raw paragraph
    indexes. This mapping follows the same top-level paragraph ordering used by
    StructuredDOCXParser and never falls back to a text search.
    """

    document = OpenXmlDocument(io.BytesIO(content))
    references: dict[str, DocxSourceUnitReference] = {}
    for paragraph_index, paragraph in enumerate(document.paragraphs):
        reference = DocxSourceUnitReference(
            paragraph_index=paragraph_index,
            content_sha256=hashlib.sha256(paragraph.text.encode("utf-8")).hexdigest(),
        )
        references[f"paragraph:{paragraph_index}"] = reference
        for bookmark in paragraph._p.xpath(".//w:bookmarkStart"):
            name = bookmark.get(qn("w:name"))
            if not name:
                continue
            source_unit_id = f"bookmark:{name}"
            if source_unit_id in references:
                raise ValueError(f"duplicate DOCX bookmark source unit: {source_unit_id}")
            references[source_unit_id] = reference
    return references


def _metric_payload(metrics: dict[str, MetricResult]) -> dict[str, JsonValue]:
    return {name: metric.model_dump(mode="json") for name, metric in sorted(metrics.items())}


def _load_approved_gold(repository_root: Path) -> tuple[DS1ParsingDataset, DS1BatchApproval]:
    dataset_root = repository_root / DATASET_ROOT
    approved_path = repository_root / APPROVED_DS1
    approval_path = repository_root / DS1_APPROVAL
    approved = DS1ParsingDataset.model_validate_json(approved_path.read_text(encoding="utf-8"))
    approval = DS1BatchApproval.model_validate_json(approval_path.read_text(encoding="utf-8"))
    if sha256_file(approved_path) != approval.approved_file_sha256:
        raise ValueError("Approved DS1 file differs from its batch ApprovalRecord")
    if len(approved.records) != 55:
        raise ValueError("formal P04 evaluation requires exactly 55 Approved DS1 records")
    if any(
        record.review_status is not ReviewStatus.APPROVED or record.approval is None
        for record in approved.records
    ):
        raise ValueError("formal P04 evaluation accepts Approved DS1 records only")
    if {record.record_id for record in approved.records} != set(approval.approved_record_sha256):
        raise ValueError("batch ApprovalRecord does not cover every Approved DS1 record")
    manifest = json.loads((repository_root / DATASET_MANIFEST).read_text(encoding="utf-8"))
    components = manifest.get("gold_components")
    component_approved = (
        isinstance(components, dict) and components.get("ds1_native_docx") == "approved"
    )
    legacy_status_approved = manifest.get("gold_status") == "ds1_p04_native_docx_approved"
    if not component_approved and not legacy_status_approved:
        raise ValueError("dataset Manifest does not expose the approved P04 DS1 boundary")
    if manifest.get("phase_input_status", {}).get("p04") != "formal_eval_ready":
        raise ValueError("P04 phase input is not ready for formal evaluation")
    global_splits_populated = any(
        (dataset_root / "splits" / split_name).read_text(encoding="utf-8").strip()
        for split_name in ("dev_ids.txt", "test_ids.txt")
    )
    p08_retrieval_approved = (
        isinstance(components, dict) and components.get("ds5_retrieval") == "approved"
    )
    if global_splits_populated and not p08_retrieval_approved:
        raise ValueError("formal P04 evaluation found premature global Dev/Test data")
    lock = TestLock.model_validate_json((repository_root / TEST_LOCK).read_text(encoding="utf-8"))
    if lock.locked:
        raise ValueError("formal P04 stage evaluation must not lock Test")
    pilot_ids = set((repository_root / PILOT_SPLIT).read_text(encoding="utf-8").split())
    missing = {record.record_id for record in approved.records} - pilot_ids
    if missing:
        raise ValueError(f"Approved P04 DS1 records are missing from Pilot: {sorted(missing)}")
    return approved, approval


def _parse_p04_docx(
    *,
    path: Path,
    document_id: str,
    document_version_id: str,
    document_sha256: str,
    pdf_runs: tuple[Path, Path],
    repository_root: Path,
    renderer: DocxPaginationRenderer,
    low_confidence_threshold: float,
) -> tuple[ParsedDocumentIR, bool]:
    parsed = (
        StructuredDOCXParser()
        .parse(
            path.read_bytes(),
            document_id=document_id,
            document_version_id=document_version_id,
            document_sha256=document_sha256,
        )
        .document
    )
    snapshots = [
        renderer.build_snapshot(
            (repository_root / pdf_path).read_bytes(),
            requested_fonts=_requested_fonts(parsed),
        )
        for pdf_path in pdf_runs
    ]
    repeat_match = (
        snapshots[0].manifest.canonical_pdf_sha256 == snapshots[1].manifest.canonical_pdf_sha256
        and snapshots[0].manifest.page_manifest_sha256 == snapshots[1].manifest.page_manifest_sha256
        and snapshots[0].manifest.page_count == snapshots[1].manifest.page_count
    )
    aligned = align_docx_blocks(
        parsed,
        snapshots[0],
        low_confidence_threshold=low_confidence_threshold,
    )
    return enrich_document_structure(aligned), repeat_match


def _run_parsers(repository_root: Path) -> ParseOutputs:
    work_package, _, paths = _load_inputs(repository_root)
    bindings = {identity.document_id: identity for identity in work_package.identities}
    profile = RendererProfile.load(repository_root / RENDERER_PROFILE)
    renderer = DocxPaginationRenderer(profile, font_resolver=lambda font: None)
    p04: dict[str, ParsedDocumentIR | None] = {}
    b0: dict[str, ParsedDocument | None] = {}
    p04_success: dict[str, bool] = {}
    b0_success: dict[str, bool] = {}
    repeat_match: dict[str, bool] = {}
    source_unit_references: dict[str, dict[str, DocxSourceUnitReference]] = {}

    for document_id in EVALUATED_DOCUMENT_IDS:
        binding = bindings[document_id]
        if document_id != PDF_DOCUMENT_ID:
            source_unit_references[document_id] = _docx_source_unit_references(
                paths[document_id].read_bytes()
            )
        try:
            if document_id == PDF_DOCUMENT_ID:
                p04[document_id] = _parse_pdf(
                    paths[document_id],
                    binding.source_document_id,
                    binding.document_version_id,
                )
            else:
                pdf_runs = PRIMARY_PDF_RUNS if document_id == PRIMARY_DOCX_ID else STRESS_PDF_RUNS
                p04[document_id], repeat_match[document_id] = _parse_p04_docx(
                    path=paths[document_id],
                    document_id=binding.source_document_id,
                    document_version_id=binding.document_version_id,
                    document_sha256=binding.document_sha256,
                    pdf_runs=pdf_runs,
                    repository_root=repository_root,
                    renderer=renderer,
                    low_confidence_threshold=profile.low_alignment_confidence,
                )
            p04_success[document_id] = True
        except (OSError, RuntimeError, ValueError):
            p04[document_id] = None
            p04_success[document_id] = False
        try:
            parser = PDFParser() if document_id == PDF_DOCUMENT_ID else DOCXParser()
            b0[document_id] = parser.parse(paths[document_id])
            b0_success[document_id] = True
        except (OSError, RuntimeError, ValueError):
            b0[document_id] = None
            b0_success[document_id] = False
    return ParseOutputs(
        p04=p04,
        b0=b0,
        p04_success=p04_success,
        b0_success=b0_success,
        renderer_repeat_match=repeat_match,
        docx_source_unit_references=source_unit_references,
    )


def _p04_page_text(document: ParsedDocumentIR | None) -> dict[int, str]:
    if document is None:
        return {}
    return {
        page.physical_page_index: "\n".join(block.text for block in page.blocks)
        for page in document.pages
        if page.physical_page_index is not None
    }


def _b0_page_text(document: ParsedDocument | None) -> dict[int, str]:
    if document is None:
        return {}
    return {
        section.page: section.content for section in document.sections if section.page is not None
    }


def _p04_document_text(document: ParsedDocumentIR | None) -> str:
    if document is None:
        return ""
    return "\n".join(block.text for page in document.pages for block in page.blocks)


def _b0_document_text(document: ParsedDocument | None) -> str:
    if document is None:
        return ""
    return "\n".join(section.content for section in document.sections)


def _gold_text_samples(dataset: DS1ParsingDataset) -> list[GoldTextSample]:
    samples: list[GoldTextSample] = []
    for record in dataset.records:
        if isinstance(record, PageGold):
            samples.extend(
                GoldTextSample(record.document_id, region.text, record.page_number)
                for region in record.regions
                if region.text and _normalize_sample(region.text)
            )
        elif isinstance(record, DocxPaginationGold) and record.source_text:
            samples.append(GoldTextSample(record.document_id, record.source_text))
        elif isinstance(record, SectionGold):
            if record.first_text:
                samples.append(GoldTextSample(record.document_id, record.first_text))
            if record.last_text:
                samples.append(GoldTextSample(record.document_id, record.last_text))
        elif isinstance(record, TableGold):
            samples.extend(
                GoldTextSample(record.source_span.document_id, cell)
                for row in record.cells
                for cell in row
                if _normalize_sample(cell)
            )
    return samples


def _sample_recall(
    samples: list[GoldTextSample],
    document_text: dict[str, str],
    page_text: dict[str, dict[int, str]],
) -> MetricResult:
    recalled = 0
    for sample in samples:
        haystack = (
            page_text.get(sample.document_id, {}).get(sample.page_number, "")
            if sample.page_number is not None
            else document_text.get(sample.document_id, "")
        )
        recalled += bool(_normalize_sample(sample.text) in _normalize_sample(haystack))
    return MetricResult.ratio("normalized_sample_text_recall", recalled, len(samples))


def _find_section(document: ParsedDocumentIR | None, title: str) -> SectionIR | None:
    if document is None:
        return None
    expected = _normalize(title)
    return next(
        (section for section in document.sections if _normalize(section.title) == expected), None
    )


def _heading_results(
    section_gold: list[SectionGold],
    outputs: dict[str, ParsedDocumentIR | None],
) -> dict[str, MetricResult]:
    gold = [
        HeadingObservation(
            record.document_id,
            record.source_span.page_start,
            record.title,
            record.level,
        )
        for record in section_gold
    ]
    predicted: list[HeadingObservation] = []
    for record in section_gold:
        section = _find_section(outputs.get(record.document_id), record.title)
        if section is None:
            continue
        predicted.append(
            HeadingObservation(
                record.document_id,
                section.source_span.page_start,
                section.title,
                section.level,
            )
        )
    return heading_metrics(predicted, gold)


def _ordered_blocks(document: ParsedDocumentIR) -> tuple[list[BlockIR], dict[str, int]]:
    blocks = [block for page in document.pages for block in page.blocks]
    return blocks, {block.block_id: index for index, block in enumerate(blocks)}


def _paragraph_index(block: BlockIR) -> int | None:
    source_unit = block.source_span.source_unit or ""
    match = _DOCX_UNIT.fullmatch(source_unit)
    return int(match.group("index")) if match else None


def _text_anchor_index(
    blocks: list[BlockIR],
    *,
    text: str | None,
    page_start: int | None,
    page_end: int | None,
) -> int | None:
    if not text:
        return None
    expected = _normalize_sample(text)
    candidates: list[tuple[int, BlockIR]] = []
    for index, block in enumerate(blocks):
        page = block.source_span.page_start
        if page_start is not None and page is not None and page < page_start:
            continue
        if page_end is not None and page is not None and page > page_end:
            continue
        candidates.append((index, block))
    exact = [index for index, block in candidates if _normalize_sample(block.text) == expected]
    if exact:
        return exact[0]
    contained = [
        index
        for index, block in candidates
        if expected and expected in _normalize_sample(block.text)
    ]
    return contained[0] if contained else None


def _section_boundary_result(
    section_gold: list[SectionGold],
    outputs: dict[str, ParsedDocumentIR | None],
) -> MetricResult:
    correct = 0
    resolved = 0
    for record in section_gold:
        document = outputs.get(record.document_id)
        section = _find_section(document, record.title)
        if document is None or section is None or not section.block_ids:
            continue
        blocks, positions = _ordered_blocks(document)
        predicted_blocks = [block for block in blocks if block.block_id in set(section.block_ids)]
        if not predicted_blocks:
            continue
        expected_start: int | None
        expected_end: int | None
        actual_start: int
        actual_end: int
        if record.document_id == PRIMARY_DOCX_ID:
            start_match = _GOLD_DOCX_UNIT.fullmatch(record.source_span.block_start or "")
            end_match = _GOLD_DOCX_UNIT.fullmatch(record.source_span.block_end or "")
            paragraph_indexes = [
                index
                for block in predicted_blocks
                if (index := _paragraph_index(block)) is not None
            ]
            if start_match is None or end_match is None or not paragraph_indexes:
                continue
            expected_start = int(start_match.group("index"))
            expected_end = int(end_match.group("index"))
            actual_start, actual_end = min(paragraph_indexes), max(paragraph_indexes)
        else:
            expected_start = _text_anchor_index(
                blocks,
                text=record.first_text,
                page_start=record.source_span.page_start,
                page_end=record.source_span.page_start,
            )
            expected_end = _text_anchor_index(
                blocks,
                text=record.last_text,
                page_start=record.source_span.page_end,
                page_end=record.source_span.page_end,
            )
            actual_start = positions[section.block_ids[0]]
            actual_end = positions[section.block_ids[-1]]
            if expected_start is None or expected_end is None:
                continue
        resolved += 1
        correct += abs(actual_start - expected_start) <= 1 and abs(actual_end - expected_end) <= 1
    return MetricResult.ratio(
        "section_boundary_accuracy",
        correct,
        len(section_gold),
        details={"tolerance_source_units": 1, "resolved_gold_boundaries": resolved},
    )


def _bbox_iou(
    left: tuple[float, float, float, float], right: tuple[float, float, float, float]
) -> float:
    width = max(0.0, min(left[2], right[2]) - max(left[0], right[0]))
    height = max(0.0, min(left[3], right[3]) - max(left[1], right[1]))
    intersection = width * height
    left_area = max(0.0, left[2] - left[0]) * max(0.0, left[3] - left[1])
    right_area = max(0.0, right[2] - right[0]) * max(0.0, right[3] - right[1])
    union = left_area + right_area - intersection
    return intersection / union if union else 0.0


def _noise_results(
    page_gold: list[PageGold],
    document: ParsedDocumentIR | None,
) -> dict[str, MetricResult]:
    gold_ids = {
        f"{record.record_id}:{region.region_id}"
        for record in page_gold
        for region in record.noise_regions
    }
    if document is None:
        return binary_detection_metrics(set(), gold_ids, prefix="noise_detection")
    pages = {page.physical_page_index: page for page in document.pages}
    predicted_ids: set[str] = set()
    matched_blocks: set[str] = set()
    for record in page_gold:
        page = pages.get(record.page_number)
        if page is None:
            continue
        predicted_noise = [block for block in page.blocks if block.noise_labels]
        for region in record.noise_regions:
            expected = _normalize_sample(region.text or "")
            match = next(
                (
                    block
                    for block in predicted_noise
                    if block.block_id not in matched_blocks
                    and expected
                    and _normalize_sample(block.text) == expected
                    and (
                        region.bbox is None
                        or block.bbox is None
                        or _bbox_iou(block.bbox, region.bbox) >= 0.25
                    )
                ),
                None,
            )
            if match is not None:
                predicted_ids.add(f"{record.record_id}:{region.region_id}")
                matched_blocks.add(match.block_id)
        predicted_ids.update(
            f"{record.record_id}:extra:{block.block_id}"
            for block in predicted_noise
            if block.block_id not in matched_blocks
        )
    return binary_detection_metrics(predicted_ids, gold_ids, prefix="noise_detection")


def _table_grid(table: TableRecord) -> list[list[str]]:
    grid = [[""] * table.column_count for _ in range(table.row_count)]
    for cell in table.cells:
        grid[cell.row_index][cell.column_index] = cell.text
    return grid


def _find_table(document: ParsedDocumentIR | None, gold: TableGold) -> TableRecord | None:
    if document is None:
        return None
    if gold.source_span.document_id == PRIMARY_DOCX_ID:
        match = _GOLD_TABLE.fullmatch(gold.source_span.block_start or "")
        if match is None:
            return None
        source_unit = f"docx:table:{int(match.group('index'))}"
        return next(
            (table for table in document.tables if table.source_span.source_unit == source_unit),
            None,
        )
    return next(
        (
            table
            for table in document.tables
            if table.source_span.page_start == gold.source_span.page_start
        ),
        None,
    )


def _table_results(
    table_gold: list[TableGold],
    outputs: dict[str, ParsedDocumentIR | None],
) -> dict[str, MetricResult]:
    predicted_cells: set[tuple[str, int, int, str]] = set()
    gold_cells: set[tuple[str, int, int, str]] = set()
    structure_correct = 0
    for gold in table_gold:
        predicted = _find_table(outputs.get(gold.source_span.document_id), gold)
        for row_index, row in enumerate(gold.cells):
            for column_index, value in enumerate(row):
                gold_cells.add((gold.record_id, row_index, column_index, _normalize(value)))
        if predicted is None:
            continue
        grid = _table_grid(predicted)
        structure_correct += (
            predicted.row_count == gold.row_count and predicted.column_count == gold.column_count
        )
        for row_index, row in enumerate(grid):
            for column_index, value in enumerate(row):
                predicted_cells.add((gold.record_id, row_index, column_index, _normalize(value)))
    cell_metrics = binary_detection_metrics(
        {repr(item) for item in predicted_cells},
        {repr(item) for item in gold_cells},
        prefix="table_cell",
    )
    return {
        "table_row_column_structure_accuracy": MetricResult.ratio(
            "table_row_column_structure_accuracy", structure_correct, len(table_gold)
        ),
        **cell_metrics,
    }


def _docx_pagination_results(
    gold: list[DocxPaginationGold],
    outputs: ParseOutputs,
) -> dict[str, MetricResult]:
    physical_predictions: list[object | None] = []
    physical_gold: list[object] = []
    display_predictions: list[object | None] = []
    display_gold: list[object] = []
    section_predictions: list[object | None] = []
    section_gold: list[object] = []
    for record in gold:
        resolution = _pagination_resolution(record, outputs)
        anchor = resolution.block.page_anchor if resolution.block is not None else None
        physical_predictions.append(anchor.physical_page_index if anchor else None)
        physical_gold.append(record.expected_physical_page_index)
        if record.expected_display_page_label is not None:
            display_predictions.append(anchor.display_page_label if anchor else None)
            display_gold.append(record.expected_display_page_label)
        if record.expected_section_page_index is not None:
            section_predictions.append(anchor.section_page_index if anchor else None)
            section_gold.append(record.expected_section_page_index)
    return {
        "docx_physical_page_accuracy": exact_mapping_accuracy(
            physical_predictions,
            physical_gold,
            name="docx_physical_page_accuracy",
        ),
        "docx_display_page_label_accuracy": exact_mapping_accuracy(
            display_predictions,
            display_gold,
            name="docx_display_page_label_accuracy",
        ),
        "docx_section_page_accuracy": exact_mapping_accuracy(
            section_predictions,
            section_gold,
            name="docx_section_page_accuracy",
        ),
    }


def _pagination_resolution(
    record: DocxPaginationGold,
    outputs: ParseOutputs,
) -> PaginationResolution:
    source_unit_id = record.source_unit_id
    if source_unit_id is None or record.source_text_sha256 is None:
        raise ValueError(
            f"formal pagination Gold lacks a stable source identity: {record.record_id}"
        )
    references = outputs.docx_source_unit_references.get(record.document_id)
    if references is None or source_unit_id not in references:
        raise ValueError(
            f"formal pagination source identity is absent from OOXML: "
            f"{record.record_id} ({source_unit_id})"
        )
    reference = references[source_unit_id]
    paragraph_match = _GOLD_DOCX_UNIT.fullmatch(source_unit_id)
    bookmark_match = _GOLD_BOOKMARK_UNIT.fullmatch(source_unit_id)
    source_index_match = paragraph_match or bookmark_match
    if source_index_match is None:
        raise ValueError(
            f"formal pagination source identity has an unsupported form: {record.record_id}"
        )
    if int(source_index_match.group("index")) != record.paragraph_index:
        raise ValueError(
            f"formal pagination paragraph identity disagrees with Gold: {record.record_id}"
        )
    source_reference_hash_match = reference.content_sha256 == record.source_text_sha256
    if not source_reference_hash_match:
        raise ValueError(f"formal pagination source Hash disagrees with OOXML: {record.record_id}")

    document = outputs.p04.get(record.document_id)
    source_unit = f"docx:paragraph:{reference.paragraph_index}"
    block = None
    if document is not None:
        block = next(
            (
                candidate
                for page in document.pages
                for candidate in page.blocks
                if candidate.source_span.source_unit == source_unit
            ),
            None,
        )
    parser_block_hash_match = (
        block is not None and block.content_sha256 == record.source_text_sha256
    )
    if not parser_block_hash_match:
        block = None
    return PaginationResolution(
        source_unit_id=source_unit_id,
        parser_paragraph_index=reference.paragraph_index,
        source_reference_hash_match=source_reference_hash_match,
        parser_block_hash_match=parser_block_hash_match,
        block=block,
    )


def _pagination_anchor(
    record: DocxPaginationGold,
    outputs: ParseOutputs,
) -> DocxPageAnchor | None:
    resolution = _pagination_resolution(record, outputs)
    block = resolution.block
    return block.page_anchor if block is not None else None


def _pagination_diagnostics(
    dataset: DS1ParsingDataset,
    outputs: ParseOutputs,
) -> list[dict[str, JsonValue]]:
    diagnostics: list[dict[str, JsonValue]] = []
    for record in dataset.records:
        if not isinstance(record, DocxPaginationGold):
            continue
        resolution = _pagination_resolution(record, outputs)
        anchor = _pagination_anchor(record, outputs)
        physical = anchor.physical_page_index if anchor else None
        display = anchor.display_page_label if anchor else None
        section = anchor.section_page_index if anchor else None
        diagnostics.append(
            {
                "record_id": record.record_id,
                "document_id": record.document_id,
                "paragraph_index": record.paragraph_index,
                "source_unit_id": resolution.source_unit_id,
                "resolved_parser_paragraph_index": resolution.parser_paragraph_index,
                "source_reference_hash_match": resolution.source_reference_hash_match,
                "parser_block_hash_match": resolution.parser_block_hash_match,
                "expected_physical_page_index": record.expected_physical_page_index,
                "actual_physical_page_index": physical,
                "physical_page_match": physical == record.expected_physical_page_index,
                "expected_display_page_label": record.expected_display_page_label,
                "actual_display_page_label": display,
                "display_page_match": (
                    record.expected_display_page_label is None
                    or display == record.expected_display_page_label
                ),
                "expected_section_page_index": record.expected_section_page_index,
                "actual_section_page_index": section,
                "section_page_match": (
                    record.expected_section_page_index is None
                    or section == record.expected_section_page_index
                ),
                "alignment_confidence": anchor.alignment_confidence if anchor else None,
            }
        )
    return diagnostics


def _not_applicable(name: str, reason: str) -> MetricResult:
    return MetricResult(
        name=name,
        value=None,
        numerator=0,
        denominator=0,
        applicable=False,
        details={"reason": reason},
    )


def _build_metrics(
    dataset: DS1ParsingDataset,
    outputs: ParseOutputs,
) -> tuple[dict[str, MetricResult], dict[str, MetricResult]]:
    page_gold = [record for record in dataset.records if isinstance(record, PageGold)]
    pagination_gold = [
        record for record in dataset.records if isinstance(record, DocxPaginationGold)
    ]
    section_gold = [record for record in dataset.records if isinstance(record, SectionGold)]
    table_gold = [record for record in dataset.records if isinstance(record, TableGold)]
    samples = _gold_text_samples(dataset)

    p04_pages = {key: _p04_page_text(value) for key, value in outputs.p04.items()}
    b0_pages = {key: _b0_page_text(value) for key, value in outputs.b0.items()}
    p04_text = {key: _p04_document_text(value) for key, value in outputs.p04.items()}
    b0_text = {key: _b0_document_text(value) for key, value in outputs.b0.items()}
    p04_metrics: dict[str, MetricResult] = {
        "document_parse_success_rate": parsing_success_rate(
            [outputs.p04_success.get(document_id, False) for document_id in EVALUATED_DOCUMENT_IDS]
        ),
        "pdf_page_mapping_accuracy": exact_mapping_accuracy(
            [
                record.page_number
                if record.page_number in p04_pages.get(record.document_id, {})
                else None
                for record in page_gold
            ],
            [record.page_number for record in page_gold],
            name="pdf_page_mapping_accuracy",
        ),
        "normalized_sample_text_recall": _sample_recall(samples, p04_text, p04_pages),
    }
    p04_metrics.update(_heading_results(section_gold, outputs.p04))
    p04_metrics["section_boundary_accuracy"] = _section_boundary_result(section_gold, outputs.p04)
    p04_metrics.update(_noise_results(page_gold, outputs.p04.get(PDF_DOCUMENT_ID)))
    p04_metrics.update(_table_results(table_gold, outputs.p04))
    p04_metrics.update(_docx_pagination_results(pagination_gold, outputs))

    b0_metrics: dict[str, MetricResult] = {
        "document_parse_success_rate": parsing_success_rate(
            [outputs.b0_success.get(document_id, False) for document_id in EVALUATED_DOCUMENT_IDS]
        ),
        "pdf_page_mapping_accuracy": exact_mapping_accuracy(
            [
                record.page_number
                if record.page_number in b0_pages.get(record.document_id, {})
                else None
                for record in page_gold
            ],
            [record.page_number for record in page_gold],
            name="pdf_page_mapping_accuracy",
        ),
        "normalized_sample_text_recall": _sample_recall(samples, b0_text, b0_pages),
    }
    b0_metrics.update(
        heading_metrics(
            [],
            [
                HeadingObservation(
                    record.document_id,
                    record.source_span.page_start,
                    record.title,
                    record.level,
                )
                for record in section_gold
            ],
        )
    )
    b0_metrics["section_boundary_accuracy"] = MetricResult.ratio(
        "section_boundary_accuracy", 0, len(section_gold), details={"tolerance_source_units": 1}
    )
    b0_metrics.update(_noise_results(page_gold, None))
    b0_metrics.update(_table_results(table_gold, {}))
    for name in (
        "docx_physical_page_accuracy",
        "docx_display_page_label_accuracy",
        "docx_section_page_accuracy",
    ):
        b0_metrics[name] = _not_applicable(name, "legacy DOCX parser has no pagination contract")
    return p04_metrics, b0_metrics


def _value(metrics: dict[str, MetricResult], name: str) -> float | None:
    return metrics[name].value


def _regresses_over_two_points(
    p04: dict[str, MetricResult], b0: dict[str, MetricResult], name: str
) -> bool:
    p04_value = _value(p04, name)
    b0_value = _value(b0, name)
    return p04_value is not None and b0_value is not None and b0_value - p04_value > 0.0200000001


def _strict_gate(
    p04: dict[str, MetricResult],
    b0: dict[str, MetricResult],
    *,
    renderer_repeat_match: dict[str, bool],
) -> dict[str, JsonValue]:
    core = (
        "heading_detection_f1",
        "section_boundary_accuracy",
        "noise_detection_f1",
        "table_cell_f1",
    )
    strict_improvements = [
        name
        for name in core
        if _value(p04, name) is not None
        and _value(b0, name) is not None
        and (_value(p04, name) or 0.0) > (_value(b0, name) or 0.0)
    ]
    common = sorted(
        name for name in p04.keys() & b0.keys() if p04[name].applicable and b0[name].applicable
    )
    regressions = [name for name in common if _regresses_over_two_points(p04, b0, name)]
    checks: dict[str, bool] = {
        "p04_parse_success_100_percent": _value(p04, "document_parse_success_rate") == 1.0,
        "pdf_page_mapping_15_of_15": (
            p04["pdf_page_mapping_accuracy"].numerator == 15
            and p04["pdf_page_mapping_accuracy"].denominator == 15
        ),
        "docx_physical_mapping_10_of_10": (
            p04["docx_physical_page_accuracy"].numerator == 10
            and p04["docx_physical_page_accuracy"].denominator == 10
        ),
        "all_observable_docx_display_labels_correct": (
            _value(p04, "docx_display_page_label_accuracy") == 1.0
        ),
        "all_observable_docx_section_pages_correct": (
            _value(p04, "docx_section_page_accuracy") == 1.0
        ),
        "normalized_sample_text_recall_at_least_99_percent": (
            (_value(p04, "normalized_sample_text_recall") or 0.0) >= 0.99
        ),
        "at_least_two_core_metrics_strictly_better_than_b0": len(strict_improvements) >= 2,
        "no_common_metric_regression_over_two_points": not regressions,
        "renderer_canonical_repeat_for_both_docx_inputs": (
            set(renderer_repeat_match) == {PRIMARY_DOCX_ID, STRESS_DOCX_ID}
            and all(renderer_repeat_match.values())
        ),
    }
    return {
        "checks": cast(JsonValue, checks),
        "strictly_improved_core_metrics": cast(JsonValue, strict_improvements),
        "common_applicable_metrics": cast(JsonValue, common),
        "regressions_over_two_points": cast(JsonValue, regressions),
        "passed": all(checks.values()),
    }


def _case_summary(
    system: str,
    document_id: str,
    succeeded: bool,
    document: ParsedDocumentIR | ParsedDocument | None,
) -> dict[str, JsonValue]:
    if isinstance(document, ParsedDocumentIR):
        return {
            "system": system,
            "document_id": document_id,
            "succeeded": succeeded,
            "page_count": len(document.pages),
            "section_count": len(document.sections),
            "table_count": len(document.tables),
            "warning_count": len(document.warnings),
        }
    if isinstance(document, ParsedDocument):
        return {
            "system": system,
            "document_id": document_id,
            "succeeded": succeeded,
            "section_count": len(document.sections),
        }
    return {"system": system, "document_id": document_id, "succeeded": succeeded}


def run_p04_formal(
    *,
    repository_root: Path,
    output_dir: Path,
    run_id: str,
    git_commit: str | None = None,
    git_dirty: bool | None = None,
) -> Path:
    repository_root = repository_root.resolve()
    output_dir = output_dir.resolve()
    if not output_dir.is_relative_to((repository_root / "storage_eval").resolve()):
        raise ValueError("formal P04 output must stay under storage_eval")
    approved, approval = _load_approved_gold(repository_root)
    outputs = _run_parsers(repository_root)
    p04_metrics, b0_metrics = _build_metrics(approved, outputs)
    gate = _strict_gate(
        p04_metrics,
        b0_metrics,
        renderer_repeat_match=outputs.renderer_repeat_match,
    )

    if (git_commit is None) != (git_dirty is None):
        raise ValueError("git_commit and git_dirty must be supplied together")
    if git_commit is None or git_dirty is None:
        git_commit, git_dirty = _git_state(repository_root)
    dataset_root = (repository_root / DATASET_ROOT).resolve()
    input_paths = (
        APPROVED_DS1,
        DS1_APPROVAL,
        WORK_PACKAGE,
        FIXTURE_MANIFEST,
        RENDERER_PROFILE,
        RENDERER_FONT_LOCK,
        *PRIMARY_PDF_RUNS,
        *STRESS_PDF_RUNS,
    )
    manifest = RunManifest(
        run_id=run_id,
        created_at=datetime.now(UTC),
        dataset=RunDatasetRef(
            dataset_id="courserag-eval",
            dataset_version="v1",
            split=DatasetSplit.PILOT,
            manifest_sha256=sha256_file(repository_root / DATASET_MANIFEST),
            split_sha256=sha256_file(repository_root / PILOT_SPLIT),
        ),
        intent=RunIntent.EVALUATION,
        tuning_enabled=False,
        git_commit=git_commit,
        git_dirty=git_dirty,
        input_artifacts=[
            _artifact(
                repository_root,
                path,
                ("application/pdf" if path.suffix == ".pdf" else "application/json"),
            )
            for path in input_paths
        ],
        component_versions={
            "p04_parser": "1.0",
            "b0_pdf_parser": "legacy-pymupdf-text",
            "b0_docx_parser": "legacy-docx2txt",
            "pymupdf": version("pymupdf"),
            "python_docx": version("python-docx"),
            "docx2txt": version("docx2txt"),
        },
        configuration={
            "approved_ds1_sha256": approval.approved_file_sha256,
            "candidate_ds1_sha256": approval.candidate_file_sha256,
            "record_count": len(approved.records),
            "section_boundary_tolerance_source_units": 1,
            "normalized_sample_text_recall_threshold": 0.99,
            "maximum_common_metric_regression": 0.02,
            "minimum_strict_core_improvements": 2,
            "ocr_metrics": "not_applicable_deferred_to_p05",
            "gold_generated_from_system_output": False,
            "docx_source_unit_resolution": (
                "exact_ooxml_source_unit_to_top_level_paragraph_ordinal_no_text_fallback"
            ),
        },
        fallback_policy=FallbackPolicy.FAIL_RUN,
        random_seed=0,
        llm_as_judge=False,
    )
    output_dir.mkdir(parents=True, exist_ok=True)
    atomic_write_json(output_dir / "run_manifest.json", manifest.model_dump(mode="json"))
    runner = EvaluationRunner(
        manifest=manifest,
        checkpoint_path=output_dir / "checkpoint.json",
        partial_report_path=output_dir / "partial_report.json",
        final_report_path=output_dir / "report.json",
        dataset_root=dataset_root,
    )
    for system, documents, successes in (
        ("p04", outputs.p04, outputs.p04_success),
        ("b0", outputs.b0, outputs.b0_success),
    ):
        for document_id in EVALUATED_DOCUMENT_IDS:
            case_fn = partial(
                _case_summary,
                system,
                document_id,
                successes.get(document_id, False),
                documents.get(document_id),
            )
            runner.run_case(
                f"{system}-{document_id}",
                sha256_file(repository_root / APPROVED_DS1),
                case_fn,
            )
    report: dict[str, JsonValue] = {
        "evaluation_scope": "p04_stage_pilot_full_approved_ds1_not_formal_test",
        "approved_ds1_sha256": approval.approved_file_sha256,
        "candidate_ds1_sha256": approval.candidate_file_sha256,
        "approved_record_count": len(approved.records),
        "p04_metrics": _metric_payload(p04_metrics),
        "b0_metrics": _metric_payload(b0_metrics),
        "strict_gate": gate,
        "docx_pagination_diagnostics": cast(JsonValue, _pagination_diagnostics(approved, outputs)),
        "ocr_metric_status": "not_applicable_deferred_to_p05",
        "gold_generated_from_system_output": False,
        "llm_as_judge_used": False,
        "fallback_used": False,
        "b0_runtime_modified": False,
        "dev_test_populated": False,
        "test_locked": False,
    }
    runner.complete(report)
    return output_dir / "report.json"


def main() -> None:
    parser = argparse.ArgumentParser(description="Run formal P04/B0 scoring on Approved DS1.")
    parser.add_argument("--repository-root", type=Path, default=Path.cwd())
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--run-id", required=True)
    parser.add_argument("--git-commit")
    parser.add_argument("--git-dirty", choices=("true", "false"))
    args = parser.parse_args()
    report_path = run_p04_formal(
        repository_root=args.repository_root,
        output_dir=args.output_dir,
        run_id=args.run_id,
        git_commit=args.git_commit,
        git_dirty=None if args.git_dirty is None else args.git_dirty == "true",
    )
    print(report_path)


if __name__ == "__main__":
    main()
