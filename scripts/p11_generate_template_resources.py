"""Generate deterministic P11 logical definitions and editable DOCX skeletons."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_ROOT = ROOT / "resources" / "templates"
EXPORTERS = TEMPLATE_ROOT / "exporters"

BLUE = RGBColor(37, 71, 106)
SLATE = RGBColor(67, 80, 95)
LIGHT = "E8EEF4"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = SLATE
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.3
    for name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 13)):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.color.rgb = BLUE
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
    styles["Title"].paragraph_format.space_after = Pt(12)
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(6)
    styles["Heading 2"].paragraph_format.space_before = Pt(10)
    styles["Heading 2"].paragraph_format.space_after = Pt(4)
    header = section.header.paragraphs[0]
    header.text = "COURSEPILOT · EDITABLE FOUNDATION TEMPLATE"
    header.style = styles["Normal"]
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(113, 128, 145)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("CoursePilot · v1")
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(113, 128, 145)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def add_metadata_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=len(rows), cols=2)
    table.autofit = False
    table.columns[0].width = Cm(4.0)
    table.columns[1].width = Cm(12.2)
    for index, (label, value) in enumerate(rows):
        left, right = table.rows[index].cells
        left.width, right.width = Cm(4.0), Cm(12.2)
        left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        left.text, right.text = label, value
        shade_cell(left, LIGHT)
        left.paragraphs[0].runs[0].font.bold = True
        for cell in (left, right):
            cell.paragraphs[0].paragraph_format.space_after = Pt(3)
            cell.paragraphs[0].paragraph_format.space_before = Pt(3)


def build_lesson_docx() -> Path:
    document = Document()
    configure_document(document)
    document.add_heading("课程教学设计", 0)
    subtitle = document.add_paragraph("可编辑通用骨架 · 标准 / 研讨 / 实践模板共用")
    subtitle.runs[0].font.color.rgb = RGBColor(113, 128, 145)
    add_metadata_table(
        document,
        [("课程名称", "{{ course_name }}"), ("主题", "{{ topic }}"), ("课时", "{{ duration }}")],
    )
    for title, body in (
        ("一、学习目标", "{{ learning_objectives }}"),
        ("二、证据与重难点", "{{ evidence_summary }}"),
        ("三、教学流程", "{{ teaching_sequence }}"),
        ("四、评价与反思", "{{ assessment_and_reflection }}"),
    ):
        document.add_heading(title, level=1)
        document.add_paragraph(body)
    path = EXPORTERS / "lesson_default_v1.docx"
    document.save(path)
    return path


def build_exam_docx() -> Path:
    document = Document()
    configure_document(document)
    document.add_heading("课程试卷 / 作业", 0)
    subtitle = document.add_paragraph("可编辑通用骨架 · 作业 / 单元测验 / 期中期末共用")
    subtitle.runs[0].font.color.rgb = RGBColor(113, 128, 145)
    add_metadata_table(
        document,
        [("课程名称", "{{ course_name }}"), ("范围", "{{ scope }}"), ("总分", "{{ total_score }}")],
    )
    document.add_heading("命题说明", level=1)
    document.add_paragraph("{{ blueprint_summary }}")
    document.add_heading("试题", level=1)
    document.add_paragraph("{{ question_blocks }}")
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_heading("参考答案与评分标准", level=1)
    document.add_paragraph("{{ answer_key_and_rubric }}")
    path = EXPORTERS / "exam_default_v1.docx"
    document.save(path)
    return path


def definitions() -> list[dict[str, object]]:
    specs = (
        (
            "lesson_standard_university_v1",
            "lesson",
            "lesson_default_docx",
            {"mode": "standard_university"},
        ),
        ("lesson_seminar_v1", "lesson", "lesson_default_docx", {"mode": "seminar"}),
        ("lesson_lab_practice_v1", "lesson", "lesson_default_docx", {"mode": "lab_practice"}),
        ("exam_chapter_assignment_v1", "exam", "exam_default_docx", {"mode": "chapter_assignment"}),
        ("exam_unit_quiz_v1", "exam", "exam_default_docx", {"mode": "unit_quiz"}),
        ("exam_midterm_final_v1", "exam", "exam_default_docx", {"mode": "midterm_final"}),
        (
            "ppt_standard_lecture_v1",
            "ppt",
            "ppt_standard_lecture_pptx",
            {"mode": "standard_lecture"},
        ),
        (
            "ppt_concept_explanation_v1",
            "ppt",
            "ppt_concept_explanation_pptx",
            {"mode": "concept_explanation"},
        ),
        ("ppt_case_seminar_v1", "ppt", "ppt_case_seminar_pptx", {"mode": "case_seminar"}),
    )
    result = []
    for template_id, artifact_type, resource, parameters in specs:
        result.append(
            {
                "artifact_type": artifact_type,
                "default_parameters": parameters,
                "exporter_profile": f"{artifact_type}_exporter",
                "generator_prompt_profile": f"{artifact_type}_generator",
                "input_schema_role": f"{artifact_type}_task_input",
                "model_route_profiles": ["planner_main", "generator_main", "content_repair_main"],
                "output_schema_role": f"{artifact_type}_artifact_output",
                "physical_resource_roles": [resource],
                "planner_prompt_profile": f"{artifact_type}_planner",
                "repair_profile": f"{artifact_type}_repair",
                "template_id": template_id,
                "validator_profile": f"{artifact_type}_validator",
                "version": "1.0.0",
            }
        )
    return result


def write_definitions() -> list[dict[str, str]]:
    entries = []
    for definition in definitions():
        artifact_type = str(definition["artifact_type"])
        path = TEMPLATE_ROOT / artifact_type / f"{definition['template_id']}.yaml"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(
            json.dumps(definition, ensure_ascii=False, sort_keys=True, indent=2) + "\n",
            encoding="utf-8",
        )
        entries.append(
            {
                "path": path.relative_to(TEMPLATE_ROOT).as_posix(),
                "sha256": sha256(path),
                "template_id": str(definition["template_id"]),
            }
        )
    return entries


def write_registry(definition_entries: list[dict[str, str]]) -> None:
    resource_specs = {
        "lesson_default_docx": "exporters/lesson_default_v1.docx",
        "exam_default_docx": "exporters/exam_default_v1.docx",
        "ppt_standard_lecture_pptx": "exporters/ppt_standard_lecture_v1.pptx",
        "ppt_concept_explanation_pptx": "exporters/ppt_concept_explanation_v1.pptx",
        "ppt_case_seminar_pptx": "exporters/ppt_case_seminar_v1.pptx",
    }
    payload = {
        "definitions": definition_entries,
        "resources": {
            role: {"path": path, "sha256": sha256(TEMPLATE_ROOT / path)}
            for role, path in resource_specs.items()
        },
        "schema_version": "coursepilot.template-registry.v1",
    }
    (TEMPLATE_ROOT / "registry_v1.yaml").write_text(
        json.dumps(payload, ensure_ascii=False, sort_keys=True, indent=2) + "\n",
        encoding="utf-8",
    )


def main() -> None:
    EXPORTERS.mkdir(parents=True, exist_ok=True)
    build_lesson_docx()
    build_exam_docx()
    entries = write_definitions()
    if all(
        (EXPORTERS / name).exists()
        for name in (
            "ppt_standard_lecture_v1.pptx",
            "ppt_concept_explanation_v1.pptx",
            "ppt_case_seminar_v1.pptx",
        )
    ):
        write_registry(entries)


if __name__ == "__main__":
    main()
